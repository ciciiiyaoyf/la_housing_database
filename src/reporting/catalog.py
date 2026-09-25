"""Catalog, DOCX variable audit, and raw repository completeness outputs."""

from __future__ import annotations

from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
import re
from uuid import uuid4

import duckdb
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import yaml

from src.cleaning.ahs_table_creator import SOURCE_ID as AHS_TABLE_CREATOR_SOURCE_ID
from src.cleaning.sipp import SOURCE_IDS as SIPP_SOURCE_IDS
from src.common.config import ProjectConfig, SourceConfig
from src.housing_finance_monthly.metadata import (
    PRODUCT_SOURCE_IDS,
    register_monthly_finance_contracts,
)
from src.reporting.system_inventory import build_system_inventory


HUD_USPS_SOURCE_ID = "hud_usps_ncwm_quarterly"
HUD_USPS_VARIABLE_IDS = (
    "usps_total_residential_addresses",
    "usps_short_term_vacant_residential_addresses",
    "usps_long_term_vacant_residential_addresses",
    "usps_vacant_residential_addresses",
    "usps_no_stat_residential_addresses",
    "usps_active_residential_addresses",
    "usps_residential_vacancy_rate",
    "usps_residential_no_stat_rate",
)
ACS_RENTAL_AVAILABILITY_SOURCE_IDS = (
    "census_acs_5year_b25003",
    "census_acs_5year_b25004",
)
ACS_RENTAL_AVAILABILITY_VARIABLE_IDS = (
    "acs_renter_occupied_units",
    "acs_renter_occupied_units_moe",
    "acs_for_rent_vacant_units",
    "acs_for_rent_vacant_units_moe",
    "acs_rental_availability_rate",
)
RENTAL_INVENTORY_STATUS = "Not implemented"
HUD_USPS_SOURCE_MANIFEST_PATH = "metadata/hud_usps_source_manifest.csv"
HUD_USPS_HISTORICAL_QUARANTINE_MANIFEST_PATH = (
    "metadata/hud_usps_quarantine_manifest.csv"
)
HUD_USPS_RESTRICTED_TABLES = (
    "data/restricted/hud_usps/active/ca_county_quarterly.parquet",
    "data/restricted/hud_usps/active/la_tract_2020_quarterly.parquet",
    "data/restricted/hud_usps/active/"
    "la_postal_zip_county_part_quarterly.parquet",
    "data/restricted/hud_usps/active/long_observations_with_lineage.parquet",
)
HUD_USPS_ARTIFACT_ROLE_COUNTS = {
    "ncwm_standardized_2020": 73,
    "tract_zip_crosswalk": 1,
    "county_fips_reference_2020": 1,
}


def _req(
    category: str,
    variable_name: str,
    *,
    sheet_source: str = "",
    sheet_frequency: str = "",
    sheet_la_availability: str = "",
    sheet_notes: str = "",
    why_hard: str = "",
    best_public_workaround: str = "",
    sheet_variable_name: str = "",
) -> dict[str, str]:
    return {
        "docx_category": category,
        "variable_name": variable_name,
        "sheet_variable_name": sheet_variable_name or variable_name,
        "sheet_source": sheet_source,
        "sheet_frequency": sheet_frequency,
        "sheet_la_availability": sheet_la_availability,
        "sheet_notes": sheet_notes,
        "why_hard": why_hard,
        "best_public_workaround": best_public_workaround,
    }


DOCX_VARIABLES = [
    _req(
        "Difficult Variables",
        "Aggregate Home Equity",
        why_hard="No public state/county balance-sheet series; Fed Flow of Funds is national only",
        best_public_workaround="Use the Version A benchmark: published ACS B25082 aggregate owner values less the matched-universe, California-calibrated first-mortgage debt benchmark; add published mortgage-free owner value in full for all-owner equity.",
    ),
    _req(
        "Difficult Variables",
        "Aggregate Mortgage Debt (outstanding balances)",
        why_hard="ACS reports payments, not balances; no public CA/LA balance series",
        best_public_workaround="Use the Version A benchmark: direct ACS B25089/B25081 levels, weighted ACS PUMS first-mortgage P&I parameters, an exact-date joint NMDB rate-term distribution, and a compatible California UPB control. Retain PMMS only as a labeled sensitivity fallback.",
    ),
    _req(
        "Difficult Variables",
        "Monthly Mortgage Payment",
        why_hard="Public data gives rates, not realized payments by geography",
        best_public_workaround="Construct using local price series (Redfin/Zillow), prevailing mortgage rates, assumed LTV/term; optionally validate against ACS owner costs.",
    ),
    _req(
        "Difficult Variables",
        "Homebuying Costs - Downpayment",
        sheet_variable_name="Homebuying Costs – Downpayment",
        why_hard="Not reported publicly by geography",
        best_public_workaround="Assume percent of median sale price (scenario-based: 10%, 20%, FHA-style); treat as policy-sensitive parameter.",
    ),
    _req(
        "Difficult Variables",
        "Homebuying Costs - Closing Costs",
        sheet_variable_name="Homebuying Costs – Closing Costs",
        why_hard="No public geographic series",
        best_public_workaround="Apply standardized percent of sale price, state-specific if available.",
    ),
    _req(
        "Difficult Variables",
        "Homebuying Costs - Property Insurance",
        sheet_variable_name="Homebuying Costs – Property Insurance",
        why_hard="NAIC only publishes state-level averages",
        best_public_workaround="Use CA state series.",
    ),
    _req(
        "Difficult Variables",
        "Rental price increase in non-rent-controlled locations",
        sheet_variable_name="Rental price increase in non–rent-controlled locations",
        why_hard="Public rent indices do not flag rent-control status",
        best_public_workaround="Use Zillow ZORI with the methodology-listed ACS B25032/B25127 renter structure and year-built inputs. Preserve lower/upper structural exposure bounds because published ACS bins cannot recover an exact rolling 15-year cutoff; construct an exact point exposure only after a complete dated property-level policy and exemption crosswalk is available.",
    ),
    _req(
        "Difficult Variables",
        "Eviction trends (high-frequency, consistent)",
        why_hard="County coverage uneven; time consistency varies",
        best_public_workaround="Use Eviction Lab ETS only where the official release contains the geography; otherwise retain official annual modeled estimates at native annual frequency without interpolation.",
    ),
    _req(
        "Difficult Variables",
        "New Apartments Completed (county time series)",
        why_hard="Clean completions data not consistently public at county level",
        best_public_workaround="Use separately labeled Central and Slow distributed-lag scenarios from 5+ unit permits, with the approved bucket weights and complete-lag missingness rules.",
    ),
    _req(
        "Difficult Variables",
        "Multifamily Under Construction",
        why_hard="No public stock measure at county/state level",
        best_public_workaround="Build separately labeled Central and Slow stock-flow scenarios from current permits, scenario completions, and every surviving pre-output permit cohort; never initialize the opening stock to zero.",
    ),
    _req(
        "Difficult Variables",
        "Single-Family Rental Construction",
        why_hard='No public "build-for-rent" flag in Census data',
        best_public_workaround="Proxy using SFR permits in high-investor / rental-growth tracts or small MF permits.",
    ),
    _req(
        "Difficult Variables",
        "Net Change in Occupied Apartments (quarterly)",
        why_hard="ACS is annual; no quarterly public alternative",
        best_public_workaround="Annual ACS change in occupied renter units; optionally interpolate or use CPS/HVS for state-level timing.",
    ),
    _req(
        "Difficult Variables",
        "Price-to-Income Ratio (quarterly/local)",
        why_hard="Income is annual; prices are monthly",
        best_public_workaround="Construct ratio using annual ACS income with interpolated or stepwise series.",
    ),
    _req(
        "Difficult Variables",
        "Natural Population Change (quarterly)",
        why_hard="Birth/death data is annual at county level",
        best_public_workaround="Annual CDC WONDER series; interpolate if quarterly timing is needed.",
    ),
    _req(
        "Difficult Variables",
        "Net Immigration / Migration (county)",
        why_hard="IRS data is annual and lagged",
        best_public_workaround="Annual IRS inflow minus outflow; treat as slow-moving VAR state variable.",
    ),
    _req(
        "Difficult Variables",
        "Sheltered vs Unsheltered Homelessness (high frequency)",
        why_hard="PIT count is annual (January only)",
        best_public_workaround="Annual HUD PIT; no credible quarterly public alternative.",
    ),
    _req("AVAILABLE VARIABLES - Homeowners", "Home Price Index", sheet_source="FHFA HPI", sheet_frequency="Quarterly (also monthly)", sheet_notes="Clean state & large-county/MSA coverage"),
    _req("AVAILABLE VARIABLES - Homeowners", "Median Sale Price (SF & all homes)", sheet_source="Redfin Data Center", sheet_frequency="Monthly (weekly raw)", sheet_notes="Level prices ideal for affordability & payment calculations"),
    _req("AVAILABLE VARIABLES - Homeowners", "Existing Homes for Sale (Inventory)", sheet_source="Redfin Data Center", sheet_frequency="Monthly", sheet_notes="Direct stock measure"),
    _req("AVAILABLE VARIABLES - Homeowners", "Months of Supply", sheet_source="Redfin Data Center", sheet_frequency="Monthly", sheet_notes="Internally consistent with inventory & sales"),
    _req("AVAILABLE VARIABLES - Homeowners", "New Listings", sheet_source="Redfin Data Center", sheet_frequency="Monthly", sheet_notes="Useful leading indicator"),
    _req("AVAILABLE VARIABLES - Homeowners", "Homes Sold / Sales Volume", sheet_source="Redfin Data Center", sheet_frequency="Monthly", sheet_notes="Transaction-side activity"),
    _req("AVAILABLE VARIABLES - Homeowners", "Mortgage Rates (30Y FRM)", sheet_source="Freddie Mac PMMS", sheet_frequency="Weekly / Monthly", sheet_notes="Commonly used national rate in regional VARs"),
    _req("AVAILABLE VARIABLES - Homeowners", "Single-Family Housing Starts (permits)", sheet_source="Census BPS (1-unit permits)", sheet_frequency="Monthly (state)", sheet_notes="Monthly at CA level; LA county is annual only"),
    _req("AVAILABLE VARIABLES - Renters", "Market Rent Index (ZORI)", sheet_source="Zillow Research", sheet_frequency="Monthly", sheet_la_availability="Yes (county)", sheet_notes="Best public rent index"),
    _req("AVAILABLE VARIABLES - Renters", "Rent Growth (YoY / MoM)", sheet_source="Zillow ZORI", sheet_frequency="Monthly", sheet_la_availability="Yes", sheet_notes="Computed directly from ZORI"),
    _req("AVAILABLE VARIABLES - Renters", "Rental Inventory", sheet_source="Zillow Observed Rent Index / Listings", sheet_frequency="Monthly", sheet_la_availability="Yes", sheet_notes="Stock-side rental measure"),
    _req("AVAILABLE VARIABLES - Renters", "Rental Vacancy (HVS / CPS)", sheet_source="Census Housing Vacancy Survey", sheet_frequency="Quarterly", sheet_la_availability="No", sheet_notes="Quarterly vacancy rate (state only)"),
    _req("AVAILABLE VARIABLES - Renters", "Multifamily Permits (5+ units)", sheet_source="Census BPS", sheet_frequency="Monthly (state)", sheet_la_availability="Annual county", sheet_notes="Best public proxy for MF starts"),
    _req("AVAILABLE VARIABLES - Overall Housing Market", "Price-to-Rent Ratio", sheet_source="Constructed from FHFA / ZORI", sheet_frequency="Monthly / Quarterly", sheet_notes="Very clean once inputs chosen"),
    _req("AVAILABLE VARIABLES - Overall Housing Market", "Housing Turnover", sheet_source="Redfin (sales / inventory)", sheet_frequency="Monthly", sheet_notes="Useful liquidity metric"),
    _req("AVAILABLE VARIABLES - Overall Housing Market", "Market Time (Days on Market)", sheet_source="Redfin Data Center", sheet_frequency="Monthly", sheet_notes="Strong cyclical indicator"),
    _req("AVAILABLE VARIABLES - Overall Housing Market", "Share of Homes with Price Drops", sheet_source="Redfin Data Center", sheet_frequency="Monthly", sheet_notes="Forward-looking price pressure"),
    _req("Population, Labor, Macro Controls", "Payroll Employment", sheet_source="BLS CES", sheet_frequency="Monthly", sheet_la_availability="Yes (MSA)", sheet_notes="LA metro rather than county"),
    _req("Population, Labor, Macro Controls", "Unemployment Rate", sheet_source="BLS LAUS", sheet_frequency="Monthly", sheet_la_availability="Yes(county)", sheet_notes="Clean county coverage"),
    _req("Population, Labor, Macro Controls", "Labor Force", sheet_source="BLS LAUS", sheet_frequency="Monthly", sheet_la_availability="Yes"),
    _req("Population, Labor, Macro Controls", "Personal Income (proxy)", sheet_source="BEA", sheet_frequency="Quarterly", sheet_la_availability="county quarterly limited", sheet_notes="State is clean; county is noisier"),
    _req("Population, Labor, Macro Controls", "CPI / Inflation", sheet_source="BLS CPI-U", sheet_frequency="Monthly", sheet_la_availability="LA metro", sheet_notes="LA CPI-U is metro-level"),
]


def _impl(
    status: str,
    source_id: str,
    variable_id: str,
    table: str,
    column: str,
    kind: str,
    limitation: str,
) -> dict[str, str]:
    return {
        "current_status": status,
        "implemented_source_id": source_id,
        "implemented_variable_id": variable_id,
        "implemented_table": table,
        "implemented_column": column,
        "observed_constructed_proxy_or_scenario": kind,
        "limitation": limitation,
    }


IMPLEMENTATION = {
    "Aggregate Home Equity": _impl(
        "Partially implemented",
        "census_acs5_finance|census_acs5_pums_housing_ca|census_tract_to_puma_2010|census_tract_to_puma_2020|fhfa_nmdb_joint_distribution|fhfa_nmdb_ca_upb|freddie_mac_pmms|bls_cpi_u_us_city_average|census_sipp_2025_state_wealth_2024",
        "acs_implied_owner_occupied_home_equity_mortgaged_pmms_20y_scenario_benchmark|acs_implied_owner_occupied_home_equity_mortgaged_pmms_25y_scenario_benchmark|acs_implied_owner_occupied_home_equity_mortgaged_pmms_30y_scenario_benchmark|acs_implied_owner_occupied_home_equity_all_pmms_20y_scenario_benchmark|acs_implied_owner_occupied_home_equity_all_pmms_25y_scenario_benchmark|acs_implied_owner_occupied_home_equity_all_pmms_30y_scenario_benchmark|acs_implied_owner_occupied_home_equity_mortgaged_benchmark|acs_implied_owner_occupied_home_equity_all_benchmark|acs_implied_owner_occupied_home_equity_mortgaged_real_2024_benchmark|acs_implied_owner_occupied_home_equity_all_real_2024_benchmark|sipp_ca_aggregate_positive_home_equity_sensitivity",
        "analytics.housing_finance_benchmark|analytics.state_annual_controls",
        "acs_implied_owner_occupied_home_equity_all_pmms_20y_scenario_benchmark|acs_implied_owner_occupied_home_equity_all_pmms_25y_scenario_benchmark|acs_implied_owner_occupied_home_equity_all_pmms_30y_scenario_benchmark|sipp_ca_aggregate_positive_home_equity_sensitivity",
        "scenario sensitivity|published-table sensitivity|schema only",
        "The methodology-authorized raw PMMS 20/25/30-year debt sensitivities produce matching mortgaged-owner and all-owner equity scenarios for all 58 California counties. They are explicit uncalibrated scenarios. The official 2025 SIPP state wealth table now also produces one separately labeled 2024 California positive-home-equity sensitivity of $4,913,341,020,000 from rounded published inputs. It is conditional on positive equity, truncates the conditional mean to net worth below the 99th percentile, and has no valid covariance-based aggregate standard error. It is not the preferred matched-universe county equity target. The calibrated nominal and real equity targets remain null until an exact-date, owner-occupied first-lien California UPB control is available. The older B25106 midpoint proxy remains only as a separately labeled legacy modeled scenario.",
    ),
    "Aggregate Mortgage Debt (outstanding balances)": _impl(
        "Partially implemented",
        "census_acs5_finance|census_acs5_pums_housing_ca|census_tract_to_puma_2010|census_tract_to_puma_2020|fhfa_nmdb_joint_distribution|fhfa_nmdb_ca_upb|freddie_mac_pmms|census_sipp_2025_state_debt_2024|census_ahs_2023_ca_table14a",
        "acs_implied_first_mortgage_balance_raw_pmms_20y_sensitivity_benchmark|acs_implied_first_mortgage_balance_raw_pmms_25y_sensitivity_benchmark|acs_implied_first_mortgage_balance_raw_pmms_30y_sensitivity_benchmark|acs_implied_first_mortgage_balance_raw_benchmark|acs_implied_first_mortgage_balance_cal_benchmark|acs_first_mortgage_debt_per_mortgaged_owner_unit_benchmark|acs_first_mortgage_debt_per_owner_unit_benchmark|acs_first_mortgage_debt_per_capita_benchmark|acs_first_mortgage_debt_to_owner_income_benchmark|sipp_ca_aggregate_primary_residence_debt_sensitivity|ahs_ca_reported_case_primary_mortgage_outstanding_principal_proxy",
        "analytics.housing_finance_benchmark|analytics.state_annual_controls",
        "acs_implied_first_mortgage_balance_cal_benchmark|sipp_ca_aggregate_primary_residence_debt_sensitivity|ahs_ca_reported_case_primary_mortgage_outstanding_principal_proxy",
        "scenario sensitivity|published-table sensitivity|reported-case public proxy|schema only",
        "Version A publishes explicit raw PMMS 20/25/30-year debt sensitivities for all 58 California counties while retaining the composite raw envelope. The official 2025 SIPP state debt table adds a broader 2024 California Home Debt sensitivity of $1,627,458,504,000. The official 2023 California AHS Table 14A adds a reported-case primary-mortgage proxy of $1,413,679,720,000 from 3,381,200 reported-bin households times the $418,100 conditional mean. AHS excludes 1,203,800 regular-mortgage households with no reported balance (26.2540 percent), does not make a missing-at-random extrapolation, may not match legal first-lien status, and has no covariance-based aggregate uncertainty measure. Neither public sensitivity matches the current Version A target universe and 2022-07-01 reference date exactly. Calibrated debt and normalized ratios remain null; the older B25106 proxy remains a separate legacy construction.",
    ),
    "Monthly Mortgage Payment": _impl("Scenario implemented", "", "mortgage_payment_3_5pct_down|mortgage_payment_10pct_down|mortgage_payment_20pct_down", "analytics.county_monthly_panel", "mortgage_payment_3_5pct_down|mortgage_payment_10pct_down|mortgage_payment_20pct_down", "scenario assumption", "Principal-and-interest scenarios only; excludes taxes, insurance, HOA, maintenance, and PMI."),
    "Homebuying Costs - Downpayment": _impl("Scenario implemented", "", "down_payment_3_5pct|down_payment_10pct|down_payment_20pct", "analytics.county_monthly_panel", "down_payment_3_5pct|down_payment_10pct|down_payment_20pct", "scenario assumption", "Scenario percentages of Redfin county median sale price."),
    "Homebuying Costs - Closing Costs": _impl("Scenario implemented", "", "closing_cost_2pct|closing_cost_3pct|closing_cost_5pct", "analytics.county_monthly_panel", "closing_cost_2pct|closing_cost_3pct|closing_cost_5pct", "scenario assumption", "Scenario percentages of Redfin county median sale price; not observed closing-cost data."),
    "Homebuying Costs - Property Insurance": _impl("Implemented", "ca_doi_residential_property_part_i|ca_dir_california_cpi", "ca_homeowners_written_premium|ca_homeowners_written_exposures|ca_homeowners_avg_written_premium|ca_cpi_u_all_items_annual|ca_homeowners_avg_written_premium_real|ca_homeowners_avg_written_premium_yoy|ca_homeowners_avg_written_premium_monthly_control", "analytics.fact_housing_observation|analytics.state_annual_controls|analytics.state_monthly_controls", "ca_homeowners_avg_written_premium|ca_cpi_u_all_items_annual|ca_homeowners_avg_written_premium_real|ca_homeowners_avg_written_premium_monthly_control", "observed annual|constructed annual|carried annual control", "Official California DOI Homeowners Multi-Peril annual average premium is deflated with official California DIR annual CPI-U as nominal_premium_year * CPI_2023 / CPI_year. The 16 real values cover 2008-2023 and the 2023 real value equals nominal by construction. CPI and real premium remain annual. The separately authorized nominal calendar-year monthly control is flagged carried_annual_control and mixed_frequency_control with is_interpolated=false; it is neither observed monthly nor Los Angeles County insurance."),
    "Rental price increase in non-rent-controlled locations": _impl(
        "Partially implemented",
        "zillow_county_zori_smoothed_sa|zillow_zip_zori_smoothed_sa|census_acs_5year_b25032|census_acs_5year_b25127",
        "zori_smoothed_sa|acs_renter_occupied_1_unit_share|acs_renter_occupied_2_to_4_units_share|acs_renter_occupied_5_plus_units_share|acs_renter_occupied_mobile_other_share|acs_renter_recent_construction_share_lower_bound|acs_renter_recent_construction_share_upper_bound|noncontrolled_exposure_share_structural_proxy_lower_bound|noncontrolled_exposure_share_structural_proxy_upper_bound|exposure_adjusted_noncontrolled_rent_growth_structural_proxy_lower_bound|exposure_adjusted_noncontrolled_rent_growth_structural_proxy_upper_bound",
        "analytics.county_monthly_panel|analytics.zip_monthly_panel|analytics.county_annual_controls|analytics.zcta_annual_controls",
        "noncontrolled_exposure_share_structural_proxy_lower_bound|noncontrolled_exposure_share_structural_proxy_upper_bound|exposure_adjusted_noncontrolled_rent_growth_structural_proxy_lower_bound|exposure_adjusted_noncontrolled_rent_growth_structural_proxy_upper_bound",
        "observed rent index|constructed structural proxy bounds|annual exposure-scaled sensitivity",
        "The methodology-listed official ACS B25032/B25127 inputs and eight same-vintage annual structure/age/exposure constructions are implemented at county and Census-ZCTA geography; the structural exposure bounds are null before 2020 and are not legal eligibility determinations. Five 2020-2024 annual Los Angeles County sensitivities multiply the exact December county ZORI year-over-year change by the same-ending-year structural lower and upper exposure bounds without monthly carry-forward. These are not observed rent growth inside legally noncontrolled units. Exact noncontrolled_exposure_share and exact exposure-adjusted growth remain unavailable until a complete dated property-level policy/exemption crosswalk exists. The official-source scaffold is not applied, and Census ZCTAs are not USPS postal ZIPs.",
    ),
    "Eviction trends (high-frequency, consistent)": _impl("Partially implemented", "eviction_lab_county_estimates", "eviction_lab_filings_estimate_annual|eviction_lab_filings_rate_estimate_annual|eviction_lab_households_threatened_estimate_annual|eviction_lab_households_threatened_rate_estimate_annual", "analytics.fact_housing_observation|analytics.county_annual_controls", "eviction_lab_filings_rate_estimate_annual|eviction_lab_households_threatened_rate_estimate_annual", "modeled annual", "The official ETS Version 3.0 release was checked on 2026-08-16 and contains neither Los Angeles nor any California geography; the official observed/validated county analysis files also contain no LA County rows. The implementation therefore retains official Eviction Lab annual county modeled estimates for 2000-2018, with source-observation flags and 95% credible interval bounds preserved. No annual value is interpolated or expanded to monthly frequency, and the series is not direct local court filing data."),
    "New Apartments Completed (county time series)": _impl("Proxy implemented", "census_bps_county", "new_apartments_completed_proxy_central|new_apartments_completed_proxy_slow", "analytics.county_monthly_panel|analytics.renter_market_monthly", "new_apartments_completed_proxy_central|new_apartments_completed_proxy_slow", "modeled Central and Slow scenarios", "Central distributes 20% of each 5+ permit cohort equally over lags 12-17, 55% over lags 18-23, and 25% over lags 24-30; Slow distributes 20% over lags 18-23, 50% over lags 24-29, and 30% over lags 30-36. Central has 288 and Slow 282 non-null LA County months. Every required lag must be present, missing permits are never zero-filled, and there is no unsuffixed best estimate. These are panel-derived scenario fields exposed in the county panel and renter view; they are intentionally not materialized as native rows in analytics.fact_housing_observation."),
    "Multifamily Under Construction": _impl("Proxy implemented", "census_bps_county", "multifamily_under_construction_proxy_central|multifamily_under_construction_proxy_slow", "analytics.county_monthly_panel|analytics.renter_market_monthly", "multifamily_under_construction_proxy_central|multifamily_under_construction_proxy_slow", "modeled Central and Slow scenarios", "Each scenario follows prior stock + current 5+ permits - scenario completions. Opening stock includes every permit cohort with positive uncompleted weight in a complete 30-month Central or 36-month Slow active window rather than starting at zero. Central has 289 and Slow 283 non-null LA County months; missing active-window inputs propagate to null and no unsuffixed best estimate is emitted. These are panel-derived scenario fields exposed in the county panel and renter view; they are intentionally not materialized as native rows in analytics.fact_housing_observation."),
    "Single-Family Rental Construction": _impl("Not implemented", "census_bps_county|census_bps_state", "permits_1_unit", "analytics.fact_housing_observation|analytics.county_monthly_panel", "", "closed / not pursued", "The approved decision keeps this target closed. Generic one-unit permits do not identify rental tenure or investor intent, and the repository does not invent an investor/rental-growth tract classifier. Reopen only if a qualifying official public source and explicit method are approved."),
    "Net Change in Occupied Apartments (quarterly)": _impl("Partially implemented", "census_acs_5year_b25003", "net_change_occupied_apartments_annual|zcta_renter_occupied_units_vintage_change_annual", "analytics.county_annual_controls|analytics.zcta_annual_controls", "net_change_occupied_apartments_annual", "proxy", "Fourteen consecutive-vintage annual county differences are implemented for 2011-2024, and 3,562 exact-ZCTA differences are implemented for 2012-2024. Both are panel-only constructions from direct acs_renter_occupied_units inputs and are not fact observations. Adjacent ACS five-year estimates overlap for four survey years, so neither result is a one-year flow; no quarterly timing is created, and Census ZCTAs are not postal ZIPs."),
    "Price-to-Income Ratio (quarterly/local)": _impl("Partially implemented", "census_acs_5year_b19013|census_acs_5year_b25077", "county_annual_price_to_income_acs|zcta_annual_price_to_income_acs", "analytics.county_annual_controls|analytics.zcta_annual_controls", "county_annual_price_to_income_acs", "constructed", "Fifteen annual LA County ACS ratios are implemented for 2010-2024. Matched-vintage ratios are also implemented for the 274 configured Census ZCTAs for 2011-2024, with six source-null ZCTA-vintages preserved and 3,830 non-null ratios. Both ratios are panel-only constructions rather than fact observations. No quarterly timing is invented, and Census ZCTAs are not postal ZIPs."),
    "Natural Population Change (quarterly)": _impl("Implemented with quarterly allocation", "cdc_wonder_natality_county|cdc_wonder_ucd_1999_2020_la_county|cdc_wonder_ucd_2018_2024_county", "cdc_wonder_live_births_annual|cdc_wonder_deaths_annual|natural_population_change_annual|natural_population_change_quarterly_allocated", "analytics.fact_housing_observation|analytics.county_annual_controls|analytics.county_quarterly_controls", "natural_population_change_quarterly_allocated", "proxy", "Annual CDC WONDER births and deaths are loaded as observed annual county facts; annual natural population change is births minus deaths. The quarterly series evenly allocates the annual total to four quarters and is flagged as imputed/interpolated timing, not observed quarterly vital events."),
    "Net Immigration / Migration (county)": _impl(
        "Implemented",
        "irs_soi_county_migration|census_pep_county_population_v2019|census_pep_county_population_v2025|census_acs_5year_b05002|census_acs_5year_b07001",
        "irs_soi_net_migration_exemptions|irs_soi_net_migration_returns|irs_soi_net_migration_agi|census_pep_county_population_annual|net_exemption_rate|acs_foreign_born_population|acs_foreign_born_population_moe|acs_moved_from_abroad_past_year|acs_moved_from_abroad_past_year_moe",
        "analytics.county_annual_controls|analytics.zcta_annual_controls|analytics.fact_housing_observation",
        "irs_soi_net_migration_exemptions|census_pep_county_population_annual|net_exemption_rate|acs_foreign_born_population|acs_moved_from_abroad_past_year",
        "constructed IRS net values|observed ACS survey supplements",
        "Annual IRS SOI in-migration minus out-migration is implemented, with exemptions as the preferred tax-return people proxy and returns/AGI retained as companions. net_exemption_rate divides each interval's net exemptions by official Census PEP resident population for that same ending year and multiplies by 1,000; it never uses prior-year or ACS population and remains annual. The methodology recommends ACS foreign-born and residence-one-year-ago county supplements and lists the Census ACS source family; after user approval, those concepts are operationalized with official B05002 and B07001. The foreign-born field is a population stock and the moved-from-abroad field is gross current-resident mobility; both remain separate annual five-year survey estimates with their published 90-percent margins of error. Neither is net migration or a replacement for IRS SOI; five-year vintages overlap and are not expanded to monthly frequency. Census-ZCTA rows are a separately labeled repository extension, never postal ZIPs. The IRS 2022-2023 interval remains explicitly flagged as the provider's enhanced-matching/new-series methodology break.",
    ),
    "Sheltered vs Unsheltered Homelessness (high frequency)": _impl("Partially implemented", "hud_pit_coc_2007_2025", "hud_pit_overall_homeless_annual|hud_pit_sheltered_homeless_annual|hud_pit_unsheltered_homeless_annual", "analytics.fact_housing_observation|analytics.coc_annual_controls", "hud_pit_overall_homeless_annual|hud_pit_sheltered_homeless_annual|hud_pit_unsheltered_homeless_annual", "observed annual control", "The official HUD PIT annual 2007-2025 CoC source and explicit CA-600/606/607/612 project composite are implemented. The source is annual point-in-time only, the composite is not Los Angeles County, and row-level sheltered-only, partial-unsheltered, and no-count records remain analytically incomplete. Explicit geography-year rows remain in CoC annual controls even when all three analytical values are null; the requested high-frequency target is still unavailable."),
    "Home Price Index": _impl("Implemented", "fhfa_hpi_state_purchase_only|fhfa_hpi_metro_purchase_only|fhfa_hpi_county_annual", "fhfa_hpi_ca_purchase_only_nsa|fhfa_hpi_ca_purchase_only_sa|fhfa_hpi_la_metropolitan_division_purchase_only_nsa|fhfa_hpi_la_metropolitan_division_purchase_only_sa|fhfa_hpi_la_county_all_transactions_nsa_annual", "analytics.state_quarterly_panel|analytics.metro_quarterly_controls|analytics.county_annual_controls|analytics.fact_housing_observation", "fhfa_hpi_ca_purchase_only_nsa|fhfa_hpi_la_county_all_transactions_nsa_annual", "observed", "Aligned with the sheet's FHFA HPI source. Zillow ZHVI remains available separately but is not counted here as the sheet HPI source."),
    "Median Sale Price (SF & all homes)": _impl("Implemented", "redfin_county_monthly|redfin_zip_monthly", "median_sale_price|redfin_zip_median_sale_price_nsa_rolling_3m", "analytics.county_monthly_panel|analytics.zip_monthly_panel", "median_sale_price", "observed", "ZIP Redfin values are overlapping rolling-three-month NSA Key Metrics."),
    "Existing Homes for Sale (Inventory)": _impl("Implemented", "redfin_county_monthly", "active_inventory", "analytics.county_monthly_panel", "active_inventory", "observed", "Aligned with the sheet's Redfin Data Center stock measure at county level. Redfin ZIP active listings are retained separately but not used to satisfy this inventory row."),
    "Months of Supply": _impl("Implemented", "redfin_county_monthly", "months_of_supply", "analytics.county_monthly_panel", "months_of_supply", "observed", "Aligned with the sheet's Redfin Data Center monthly source at county level."),
    "New Listings": _impl("Implemented", "redfin_county_monthly|redfin_zip_monthly", "new_listings|redfin_zip_new_listings_rolling_3m", "analytics.county_monthly_panel|analytics.zip_monthly_panel", "new_listings", "observed", "ZIP methodology is overlapping rolling-three-month."),
    "Homes Sold / Sales Volume": _impl("Implemented", "redfin_county_monthly|redfin_zip_monthly", "homes_sold|redfin_zip_homes_sold_rolling_3m", "analytics.county_monthly_panel|analytics.zip_monthly_panel", "homes_sold", "observed", "ZIP methodology is overlapping rolling-three-month."),
    "Mortgage Rates (30Y FRM)": _impl("Implemented", "freddie_mac_pmms", "mortgage_rate_30y_weekly|mortgage_rate_30y_monthly_average|mortgage_rate_30y_monthly_median|mortgage_rate_30y_month_end|mortgage_rate_30y_observation_count", "analytics.fact_housing_observation|analytics.county_monthly_panel", "mortgage_rate_30y_monthly_average", "observed|constructed", "National PMMS, not LA-specific."),
    "Single-Family Housing Starts (permits)": _impl("Proxy implemented", "census_bps_county|census_bps_state|census_bps_county_annual_final|census_pep_county_housing_units_v2019|census_pep_county_housing_units_v2025", "permits_1_unit|permits_1_unit_annual_final|census_pep_county_housing_units_annual|permits_1_unit_per_1000_housing_units", "analytics.fact_housing_observation|analytics.county_monthly_panel|analytics.county_annual_controls", "permits_1_unit|permits_1_unit_per_1000_housing_units", "authorization proxy|constructed annual rate", "Census BPS permits are authorizations, not starts or completions. The annual normalized rate divides final annual one-unit permits by the same-year Census PEP July 1 housing-unit estimate and multiplies by 1,000. It is retained annually and never constructed from monthly permits or an ACS/lagged denominator."),
    "Market Rent Index (ZORI)": _impl("Implemented", "zillow_county_zori_smoothed_sa|zillow_zip_zori_smoothed_sa", "zori_smoothed_sa", "analytics.county_monthly_panel|analytics.zip_monthly_panel", "zori", "observed", "Active ZORI is smoothed, seasonally adjusted, all homes plus multifamily."),
    "Rent Growth (YoY / MoM)": _impl("Implemented", "", "zori_growth_mom|zori_growth_yoy", "analytics.county_monthly_panel", "zori_growth_mom|zori_growth_yoy", "constructed", "Calculated from seasonally adjusted ZORI with initial missing months preserved."),
    "Rental Inventory": _impl(
        RENTAL_INVENTORY_STATUS,
        "",
        "",
        "",
        "",
        "unavailable target",
        "The annual ACS rental-availability proxy is for-rent vacant units "
        "divided by renter-occupied plus for-rent vacant units, using matched "
        "B25003 and B25004 estimates. Component 90% margins of error are "
        "retained, and unreliable county ratios are suppressed. Direct ACS "
        "counts and margins of error are retained at county and ZCTA geography; "
        "ZCTAs are not postal ZIPs. The eight active HUD-USPS variables are "
        "native-quarterly only and remain a separate all-residential measure. The "
        "vacancy proxy is (short-term vacant + long-term vacant) / total "
        "residential addresses, while no-stat remains a separate diagnostic. "
        "Direct provider components are observed at standardized 2020-tract "
        "geography; California county aggregates are constructed, and Los "
        "Angeles County postal-ZIP county-parts are modeled with the fixed "
        "2026 Q1 crosswalk. Detailed tract and modeled ZIP-part derivatives "
        "remain access-controlled under the HUD-USPS sublicense. This is not "
        "true rental-listing inventory; ZORI is not inventory (it is a rent "
        "index), and Redfin for-sale active listings are not a valid substitute. "
        "The governing methodology names no approved source for the true monthly "
        "rental-listing series, so the blocker is a methodology/source-identification "
        "decision rather than a credential or coding step. "
        "The former derived package file was deleted; its historical "
        "quarantine manifest is superseded and is not an active input.",
    ),
    "Rental Vacancy (HVS / CPS)": _impl("Implemented", "census_hvs_state_rental_vacancy|census_hvs_metro_rental_vacancy_2005_2014|census_hvs_metro_rental_vacancy_2015_2025|census_hvs_metro_rental_vacancy_2026", "hvs_rental_vacancy_rate_ca|hvs_rental_vacancy_rate_la_metro", "analytics.state_quarterly_panel|analytics.metro_quarterly_controls", "hvs_rental_vacancy_rate_ca|hvs_rental_vacancy_rate_la_metro", "observed", "State/metro quarterly only; not LA County."),
    "Multifamily Permits (5+ units)": _impl("Implemented", "census_bps_county|census_bps_state", "permits_5_plus_units", "analytics.fact_housing_observation|analytics.county_monthly_panel", "permits_5_plus_units", "observed", "Supply-pipeline authorization measure; not starts or completions."),
    "Price-to-Rent Ratio": _impl("Implemented", "zillow_county_zhvi|zillow_zip_zhvi|zillow_county_zori_smoothed_sa|zillow_zip_zori_smoothed_sa", "price_to_rent_ratio_zillow|price_to_rent_ratio|implied_gross_rent_yield", "analytics.county_monthly_panel|analytics.zip_monthly_panel", "price_to_rent_ratio_zillow|implied_gross_rent_yield", "constructed", "The latest methodology governs: matched-month positive Zillow ZHVI divided by active smoothed-seasonally-adjusted ZORI is the canonical county and postal-ZIP ratio. The inverse gross-yield measure remains explicitly labeled gross."),
    "Housing Turnover": _impl("Implemented", "redfin_county_monthly", "sales_to_inventory_ratio|inventory_to_sales_ratio", "analytics.county_monthly_panel", "sales_to_inventory_ratio|inventory_to_sales_ratio", "constructed", "The listed Redfin sales-to-inventory and inventory-to-sales liquidity proxies are implemented. They are not true sales divided by total housing stock and retain that qualification."),
    "Market Time (Days on Market)": _impl("Implemented", "redfin_county_monthly|redfin_zip_monthly", "days_on_market|redfin_zip_median_days_on_market_rolling_3m", "analytics.county_monthly_panel|analytics.zip_monthly_panel", "days_on_market|median_days_on_market", "observed", "ZIP methodology is overlapping rolling-three-month."),
    "Share of Homes with Price Drops": _impl("Implemented", "redfin_county_monthly", "price_drop_share", "analytics.county_monthly_panel", "price_drop_share", "observed", "Aligned with the sheet's Redfin Data Center monthly source at county level."),
    "Payroll Employment": _impl("Implemented", "bls_ces_la_metro", "ces_total_nonfarm_employment_sa", "analytics.metro_monthly_controls", "ces_total_nonfarm_employment_sa", "observed", "LA metro series, not county."),
    "Unemployment Rate": _impl("Implemented", "bls_laus_la_county", "unemployment_rate", "analytics.county_monthly_panel", "unemployment_rate", "observed", "LA County LAUS NSA."),
    "Labor Force": _impl("Implemented", "bls_laus_la_county", "labor_force|employment|unemployment", "analytics.county_monthly_panel", "labor_force", "observed", "LA County LAUS NSA."),
    "Personal Income (proxy)": _impl("Implemented", "bea_state_personal_income|bea_county_personal_income", "bea_ca_personal_income_q_sa|bea_la_county_personal_income_annual", "analytics.state_quarterly_panel|analytics.county_annual_controls", "bea_ca_personal_income_q_sa|bea_la_county_personal_income_annual", "observed", "Aligned with BEA as the listed source. California quarterly and Los Angeles County annual series are loaded at their native frequencies; the county annual values are never relabeled as observed quarterly data."),
    "CPI / Inflation": _impl("Implemented", "bls_cpi_la_metro", "cpi_u_all_items_la_nsa", "analytics.metro_monthly_controls", "cpi_u_all_items_la_nsa", "observed", "LA metro CPI-U, not county."),
}


# Version A defines nine source IDs, but ``freddie_mac_pmms`` is already in the
# general source registry. These are the eight additional IDs needed for one
# consolidated, duplicate-free data catalog.
VERSION_A_UNIQUE_SOURCE_IDS = (
    "census_acs5_finance",
    "census_acs5_pums_housing_ca",
    "census_tract_to_puma_2010",
    "census_tract_to_puma_2020",
    "cfpb_hmda_modified_lar",
    "fhfa_nmdb_joint_distribution",
    "fhfa_nmdb_ca_upb",
    "bls_cpi_u_us_city_average",
)

VERSION_A_SOURCE_PRODUCTS = {
    "census_acs5_finance": "ACS 5-Year Housing Finance Benchmark Tables",
    "census_acs5_pums_housing_ca": (
        "ACS 5-Year PUMS California Housing Records"
    ),
    "census_tract_to_puma_2010": (
        "2010 Census Tract-to-PUMA Relationship"
    ),
    "census_tract_to_puma_2020": (
        "2020 Census Tract-to-PUMA Relationship"
    ),
    "cfpb_hmda_modified_lar": "HMDA Modified Loan/Application Register",
    "fhfa_nmdb_joint_distribution": (
        "NMDB Joint Outstanding Mortgage Rate-Term Distribution"
    ),
    "fhfa_nmdb_ca_upb": "NMDB California Unpaid Principal Balance Control",
    "bls_cpi_u_us_city_average": (
        "CPI-U U.S. City Average All Items Annual Average"
    ),
}

VERSION_A_SOURCE_VARIABLES = {
    "census_acs5_finance": (
        "acs_total_population_benchmark",
        "acs_aggregate_selected_monthly_owner_costs_mortgaged_benchmark",
        "acs_owner_occupied_units_benchmark",
        "acs_mortgaged_owner_units_benchmark",
        "acs_aggregate_owner_value_all_benchmark",
        "acs_aggregate_owner_value_mortgaged_benchmark",
        "acs_aggregate_owner_value_no_mortgage_benchmark",
        "acs_aggregate_household_income_owner_benchmark",
        "acs_aggregate_household_income_mortgaged_owner_benchmark",
        "acs_aggregate_household_income_no_mortgage_owner_benchmark",
    ),
    "census_acs5_pums_housing_ca": (
        "pums_first_mortgage_pi_share_benchmark",
        "pums_first_mortgage_pi_share_applied_benchmark",
        "pums_junior_lien_payment_per_mortgaged_unit_benchmark",
    ),
    "census_tract_to_puma_2010": ("tract_to_puma_relationship_2010",),
    "census_tract_to_puma_2020": ("tract_to_puma_relationship_2020",),
    "cfpb_hmda_modified_lar": (
        "hmda_median_downpayment_rate_benchmark",
        "hmda_property_value_weighted_downpayment_rate_benchmark",
        "hmda_aggregate_downpayment_flow_benchmark",
    ),
    "fhfa_nmdb_joint_distribution": (
        "nmdb_joint_outstanding_rate_remaining_term_distribution",
    ),
    "fhfa_nmdb_ca_upb": ("nmdb_california_upb_control",),
    "bls_cpi_u_us_city_average": (
        "cpi_u_us_city_average_annual_average_benchmark",
    ),
}


def _source_lifecycle_status(source: SourceConfig) -> str:
    """Return the single provenance-aware source status used by all catalogs."""

    provenance = str(source.provenance_status or "").strip().lower()
    if provenance.startswith("quarantined"):
        return "Quarantined"
    return "Active" if source.enabled else "Inactive"


def _version_a_source_settings(config: ProjectConfig) -> dict[str, dict]:
    path = config.root / "config" / "version_a_benchmark.yaml"
    if not path.exists():
        return {}
    with path.open("r", encoding="utf-8") as handle:
        document = yaml.safe_load(handle) or {}
    sources = document.get("sources", {})
    if not isinstance(sources, dict):
        raise ValueError("version_a_benchmark.yaml sources must be a mapping")
    missing = sorted(set(VERSION_A_UNIQUE_SOURCE_IDS) - set(sources))
    if missing:
        raise ValueError(
            "Version A catalog source definitions are missing: "
            f"{missing}"
        )
    return sources


def _latest_version_a_manifest_rows(config: ProjectConfig) -> pd.DataFrame:
    path = config.root / "metadata" / "version_a_raw_file_manifest.csv"
    if not path.exists():
        return pd.DataFrame()
    manifest = pd.read_csv(path, dtype=str)
    if manifest.empty:
        return manifest
    required = {"source_id", "retrieved_at"}
    if not required.issubset(manifest.columns):
        raise ValueError(
            "Version A raw manifest is missing columns: "
            f"{sorted(required - set(manifest.columns))}"
        )
    return (
        manifest.sort_values(["source_id", "retrieved_at"])
        .drop_duplicates("source_id", keep="last")
        .set_index("source_id", drop=False)
    )


def _clean_text(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    return str(value)


def _version_a_catalog_records(config: ProjectConfig) -> list[dict[str, object]]:
    settings = _version_a_source_settings(config)
    if not settings:
        return []
    manifest = _latest_version_a_manifest_rows(config)
    benchmark = _read_parquet(
        config.root / "data" / "final" / "housing_finance_benchmark.parquet"
    )
    records: list[dict[str, object]] = []
    for source_id in VERSION_A_UNIQUE_SOURCE_IDS:
        source = settings[source_id]
        manifest_row = (
            manifest.loc[source_id].to_dict()
            if not manifest.empty and source_id in manifest.index
            else {}
        )
        raw_relative_path = _clean_text(manifest_row.get("relative_path"))
        raw_path = config.root / raw_relative_path if raw_relative_path else None
        archived = bool(raw_path and raw_path.is_file())
        if benchmark.empty or "source_id" not in benchmark:
            source_rows = pd.DataFrame()
        else:
            source_pattern = rf"(?:^|\|){re.escape(source_id)}(?:\||$)"
            source_rows = benchmark.loc[
                benchmark["source_id"]
                .astype("string")
                .str.contains(source_pattern, regex=True, na=False)
            ].copy()
        date_column = next(
            (
                column
                for column in (
                    "source_period_end",
                    "reference_date",
                    "month",
                )
                if column in source_rows
            ),
            None,
        )
        dates = (
            pd.to_datetime(source_rows[date_column], errors="coerce")
            if date_column
            else pd.Series(dtype="datetime64[ns]")
        )
        acquisition = _clean_text(source.get("acquisition"))
        manual = (
            _clean_text(source.get("implementation_status"))
            == "manual_required"
        )
        download_url = next(
            (
                _clean_text(source.get(key))
                for key in (
                    "download_url",
                    "download_url_template",
                    "api_url_template",
                )
                if _clean_text(source.get(key))
            ),
            "",
        )
        source_page_url = _clean_text(source.get("source_page_url")) or download_url
        source_filename = _clean_text(manifest_row.get("source_filename"))
        file_format = Path(source_filename).suffix.lstrip(".")
        if not file_format:
            file_format = "provider_export" if manual else "provider_file"
        update_command = (
            "python -m src.benchmark.pipeline --stage download --source "
            f"{source_id}"
            + (
                " --start-year YYYY --local-file PATH_TO_OFFICIAL_FILE"
                " --release-date YYYY-MM-DD"
                if manual
                else " --force-download"
            )
        )
        geography = _clean_text(source.get("geography"))
        notes = _clean_text(source.get("notes"))
        records.append(
            {
                "source_id": source_id,
                "provider": _clean_text(source.get("provider")),
                "dataset_name": VERSION_A_SOURCE_PRODUCTS[source_id],
                "source_page_url": source_page_url,
                "download_url": download_url,
                "access_method": acquisition,
                "active_status": "Version A benchmark",
                "provenance_status": (
                    "archived_official_source"
                    if archived
                    else "registered_official_source_no_archived_vintage"
                ),
                "manual_download_required": manual,
                "manual_reason": notes if manual else "",
                "file_format": file_format,
                "raw_file_path": raw_relative_path,
                "original_filename": _clean_text(
                    manifest_row.get("original_filename")
                ),
                "retrieved_at": _clean_text(manifest_row.get("retrieved_at")),
                "provider_vintage": _clean_text(
                    manifest_row.get("release_date")
                )
                or _clean_text(manifest_row.get("data_vintage")),
                "checksum_sha256": _clean_text(
                    manifest_row.get("checksum_sha256")
                ),
                "native_frequency": _clean_text(source.get("native_frequency")),
                "geography": geography,
                "variable_ids": VERSION_A_SOURCE_VARIABLES[source_id],
                "start_date": dates.min() if dates.notna().any() else None,
                "end_date": dates.max() if dates.notna().any() else None,
                "row_count": len(source_rows),
                "destination_clean_table": "data/benchmark/clean",
                "destination_analytical_table": (
                    "analytics.housing_finance_benchmark"
                ),
                "update_command": update_command,
                "known_limitations": notes,
                "license_or_use_notes": "Official public source; cite provider terms.",
                "observed_constructed_proxy_or_scenario": (
                    "geography crosswalk"
                    if source_id.startswith("census_tract_to_puma_")
                    else "published parameter input"
                    if source_id.startswith("fhfa_nmdb_")
                    else "observed source input"
                ),
                "notes": (
                    "Version A source archived with immutable provenance."
                    if archived
                    else "Version A source registered; no archived vintage is active."
                ),
            }
        )
    return records


PANEL_PATHS = {
    "analytics.fact_housing_observation": "fact_housing_observation.parquet",
    "analytics.county_monthly_panel": "county_monthly_panel.parquet",
    "analytics.zip_monthly_panel": "zip_monthly_panel.parquet",
    "analytics.state_quarterly_panel": "state_quarterly_panel.parquet",
    "analytics.metro_monthly_controls": "metro_monthly_controls.parquet",
    "analytics.metro_quarterly_controls": "metro_quarterly_controls.parquet",
    "analytics.county_annual_controls": "county_annual_controls.parquet",
    "analytics.zcta_annual_controls": "zcta_annual_controls.parquet",
    "analytics.county_quarterly_controls": "county_quarterly_controls.parquet",
    "analytics.state_annual_controls": "state_annual_controls.parquet",
    "analytics.state_monthly_controls": "state_monthly_controls.parquet",
    "analytics.coc_annual_controls": "coc_annual_controls.parquet",
}


def _read_parquet(path: Path) -> pd.DataFrame:
    return pd.read_parquet(path) if path.exists() else pd.DataFrame()


def _load_fact(config: ProjectConfig) -> pd.DataFrame:
    return _read_parquet(config.root / "data" / "final" / "fact_housing_observation.parquet")


def _fact_stats(fact: pd.DataFrame, variable_ids: str) -> dict[str, object]:
    ids = [item for item in str(variable_ids).split("|") if item]
    if fact.empty or not ids:
        return {"start_date": None, "end_date": None, "row_count": 0, "missingness_summary": ""}
    work = fact.loc[fact["variable_id"].isin(ids)].copy()
    if work.empty:
        return {"start_date": None, "end_date": None, "row_count": 0, "missingness_summary": ""}
    missing = int(work["value"].isna().sum())
    return {
        "start_date": work["observation_date"].min(),
        "end_date": work["observation_date"].max(),
        "row_count": len(work),
        "missingness_summary": f"{missing} missing of {len(work)} fact rows",
    }


def _panel_stats(
    config: ProjectConfig,
    tables: str,
    columns: str,
    *,
    non_null_only: bool = False,
) -> dict[str, object]:
    table_names = [item for item in str(tables).split("|") if item in PANEL_PATHS]
    column_names = [item for item in str(columns).split("|") if item]
    parts: list[pd.DataFrame] = []
    for table in table_names:
        panel = _read_parquet(config.root / "data" / "final" / PANEL_PATHS[table])
        if panel.empty:
            continue
        present = [column for column in column_names if column in panel.columns]
        if not present:
            continue
        date_col = next(
            (
                candidate
                for candidate in ("observation_date", "date")
                if candidate in panel.columns
            ),
            None,
        )
        if date_col is None:
            continue
        melted = panel[[date_col, *present]].melt(
            id_vars=[date_col],
            value_vars=present,
            var_name="column",
            value_name="value",
        )
        parts.append(melted.rename(columns={date_col: "observation_date"}))
    if not parts:
        return {"start_date": None, "end_date": None, "row_count": 0, "missingness_summary": ""}
    work = pd.concat(parts, ignore_index=True)
    if non_null_only:
        work = work.loc[work["value"].notna()].copy()
        if work.empty:
            return {
                "start_date": None,
                "end_date": None,
                "row_count": 0,
                "missingness_summary": "",
            }
    missing = int(work["value"].isna().sum())
    return {
        "start_date": work["observation_date"].min(),
        "end_date": work["observation_date"].max(),
        "row_count": len(work),
        "missingness_summary": f"{missing} missing of {len(work)} panel cells",
    }


def _benchmark_stats(
    config: ProjectConfig,
    tables: str,
    variable_ids: str,
) -> dict[str, object]:
    if "analytics.housing_finance_benchmark" not in str(tables).split("|"):
        return {
            "start_date": None,
            "end_date": None,
            "row_count": 0,
            "missingness_summary": "",
        }
    frame = _read_parquet(
        config.root / "data" / "final" / "housing_finance_benchmark.parquet"
    )
    ids = [item for item in str(variable_ids).split("|") if item]
    if frame.empty or not ids or "variable_id" not in frame:
        return {
            "start_date": None,
            "end_date": None,
            "row_count": 0,
            "missingness_summary": "",
        }
    work = frame.loc[frame["variable_id"].isin(ids)].copy()
    if work.empty:
        return {
            "start_date": None,
            "end_date": None,
            "row_count": 0,
            "missingness_summary": "",
        }
    date_column = next(
        (
            column
            for column in ("source_period_end", "reference_date", "month")
            if column in work
        ),
        None,
    )
    dates = (
        pd.to_datetime(work[date_column], errors="coerce")
        if date_column
        else pd.Series(dtype="datetime64[ns]")
    )
    missing = int(work["value"].isna().sum()) if "value" in work else 0
    return {
        "start_date": dates.min() if dates.notna().any() else None,
        "end_date": dates.max() if dates.notna().any() else None,
        "row_count": len(work),
        "missingness_summary": (
            f"{missing} missing of {len(work)} Version A benchmark rows"
        ),
    }


def _best_stats(config: ProjectConfig, fact: pd.DataFrame, implementation: dict[str, str]) -> dict[str, object]:
    non_fact_derived_columns = {
        "net_change_occupied_apartments_annual",
        "county_annual_price_to_income_acs",
        "zcta_annual_price_to_income_acs",
        "zcta_renter_occupied_units_vintage_change_annual",
        "noncontrolled_exposure_share_structural_proxy_lower_bound",
        "noncontrolled_exposure_share_structural_proxy_upper_bound",
    }
    implementation_columns = {
        item
        for item in str(implementation["implemented_column"]).split("|")
        if item
    }
    if implementation_columns & non_fact_derived_columns:
        return _panel_stats(
            config,
            implementation["implemented_table"],
            implementation["implemented_column"],
            non_null_only=True,
        )
    stats = _fact_stats(fact, implementation["implemented_variable_id"])
    if stats["row_count"]:
        return stats
    stats = _benchmark_stats(
        config,
        implementation["implemented_table"],
        implementation["implemented_variable_id"],
    )
    if stats["row_count"]:
        return stats
    return _panel_stats(
        config,
        implementation["implemented_table"],
        implementation["implemented_column"],
        non_null_only=True,
    )


def _raw_manifest(config: ProjectConfig) -> pd.DataFrame:
    path = config.root / "metadata" / "raw_file_manifest.csv"
    return pd.read_csv(path, dtype=str) if path.exists() else pd.DataFrame()


def _hud_usps_source_manifest(config: ProjectConfig) -> pd.DataFrame:
    """Return the official active HUD-USPS artifact inventory."""

    path = config.root / HUD_USPS_SOURCE_MANIFEST_PATH
    if not path.exists():
        return pd.DataFrame()
    manifest = pd.read_csv(path, dtype=str)
    if manifest.empty:
        return manifest
    if "source_id" not in manifest:
        return pd.DataFrame()
    return manifest.loc[
        manifest["source_id"].eq(HUD_USPS_SOURCE_ID)
    ].copy()


def _hud_usps_expected_quarters() -> set[str]:
    return {
        f"{year}-Q{quarter}"
        for year in range(2008, 2027)
        for quarter in range(1, 5)
        if (year, quarter) <= (2026, 1)
    }


def _hud_usps_manifest_checks(config: ProjectConfig) -> dict[str, bool]:
    """Validate the 75-row official inventory without reading licensed payloads."""

    manifest_path = config.root / HUD_USPS_SOURCE_MANIFEST_PATH
    manifest = _hud_usps_source_manifest(config)
    required = {
        "artifact_role",
        "archive_period",
        "provider_vintage",
        "source_filename",
        "relative_path",
        "retrieved_at",
        "checksum_sha256",
    }
    columns_present = required.issubset(manifest.columns)
    if not columns_present:
        return {
            "raw_snapshot_present": False,
            "manifest_present": False,
            "checksum_present": False,
            "retrieval_timestamp_present": False,
        }

    role_counts = manifest["artifact_role"].value_counts().to_dict()
    role_inventory_complete = role_counts == HUD_USPS_ARTIFACT_ROLE_COUNTS
    ncwm = manifest.loc[
        manifest["artifact_role"].eq("ncwm_standardized_2020")
    ]
    quarter_inventory_complete = (
        set(ncwm["archive_period"].dropna())
        == _hud_usps_expected_quarters()
    )
    paths = [
        config.root / relative_path
        for relative_path in manifest["relative_path"].fillna("")
        if relative_path
    ]
    raw_snapshot_present = (
        len(paths) == sum(HUD_USPS_ARTIFACT_ROLE_COUNTS.values())
        and len(set(paths)) == len(paths)
        and all(path.is_file() for path in paths)
    )
    checksums = manifest["checksum_sha256"].fillna("")
    checksum_present = checksums.str.fullmatch(r"[0-9a-fA-F]{64}").all()
    retrieval_timestamp_present = pd.to_datetime(
        manifest["retrieved_at"],
        utc=True,
        errors="coerce",
        format="mixed",
    ).notna().all()
    return {
        "raw_snapshot_present": bool(raw_snapshot_present),
        "manifest_present": bool(
            manifest_path.is_file()
            and len(manifest) == sum(HUD_USPS_ARTIFACT_ROLE_COUNTS.values())
            and role_inventory_complete
            and quarter_inventory_complete
        ),
        "checksum_present": bool(checksum_present),
        "retrieval_timestamp_present": bool(retrieval_timestamp_present),
    }


def _hud_usps_catalog_manifest_row(
    config: ProjectConfig,
) -> dict[str, object]:
    """Summarize the active source manifest without substituting the old package."""

    manifest = _hud_usps_source_manifest(config)
    if manifest.empty or "artifact_role" not in manifest:
        return {}
    ncwm = manifest.loc[
        manifest["artifact_role"].eq("ncwm_standardized_2020")
    ].copy()
    if ncwm.empty:
        return {}
    latest = ncwm.sort_values("archive_period").iloc[-1].to_dict()
    retrieved = pd.to_datetime(
        manifest["retrieved_at"],
        utc=True,
        errors="coerce",
        format="mixed",
    )
    if retrieved.notna().any():
        latest["retrieved_at"] = retrieved.max().isoformat()
    latest["provider_vintage"] = "2008Q1 through 2026Q1; crosswalk 2026Q1"
    latest["source_manifest_path"] = HUD_USPS_SOURCE_MANIFEST_PATH
    latest["source_manifest_artifact_count"] = len(manifest)
    return latest


def _latest_manifest_rows(config: ProjectConfig) -> pd.DataFrame:
    manifest = _raw_manifest(config)
    if manifest.empty:
        return pd.DataFrame()
    return (
        manifest.sort_values(["source_id", "retrieved_at"])
        .drop_duplicates("source_id", keep="last")
        .set_index("source_id", drop=False)
    )


def _raw_path(config: ProjectConfig, manifest_row: dict[str, object]) -> Path | None:
    relative_path = str(manifest_row.get("relative_path") or "")
    if not relative_path or relative_path.lower() == "nan":
        return None
    return config.root / relative_path


def _clean_table_names(config: ProjectConfig) -> set[str]:
    clean_dir = config.root / "data" / "clean"
    return {path.stem for path in clean_dir.glob("*.parquet")} if clean_dir.exists() else set()


def _source_clean_name(source: SourceConfig) -> str:
    return source.destination_clean_table or source.source_id


def _source_observation_path(config: ProjectConfig, source: SourceConfig) -> Path:
    if source.source_id == "redfin_zip_monthly":
        return (
            config.root
            / "data"
            / "interim"
            / "redfin_zip_monthly_observations.parquet"
        )
    return config.root / "data" / "clean" / f"{source.source_id}.parquet"


def _source_stats(config: ProjectConfig, source: SourceConfig) -> dict[str, object]:
    frame = _read_parquet(_source_observation_path(config, source))
    if frame.empty or "observation_date" not in frame.columns:
        return {
            "start_date": None,
            "end_date": None,
            "row_count": 0,
            "missingness_summary": "",
        }
    if "variable_id" in frame.columns and source.key_variables:
        frame = frame.loc[frame["variable_id"].isin(source.key_variables)].copy()
    if frame.empty:
        return {
            "start_date": None,
            "end_date": None,
            "row_count": 0,
            "missingness_summary": "",
        }
    missing = int(frame["value"].isna().sum()) if "value" in frame.columns else 0
    return {
        "start_date": frame["observation_date"].min(),
        "end_date": frame["observation_date"].max(),
        "row_count": len(frame),
        "missingness_summary": f"{missing} missing of {len(frame)} clean rows",
    }


def _docx_audit(config: ProjectConfig, fact: pd.DataFrame) -> pd.DataFrame:
    manifest = _latest_manifest_rows(config)
    clean_tables = _clean_table_names(config)
    source_by_id = {source.source_id: source for source in config.sources}
    rows = []
    for requirement in DOCX_VARIABLES:
        name = requirement["variable_name"]
        category = requirement["docx_category"]
        implementation = IMPLEMENTATION[name]
        stats = _best_stats(config, fact, implementation)
        source_ids = [
            item
            for item in implementation["implemented_source_id"].split("|")
            if item
        ]
        clean_names = [
            _source_clean_name(source_by_id[item])
            for item in source_ids
            if item in source_by_id
        ]
        rows.append(
            {
                "docx_category": category,
                "variable_name": name,
                "sheet_variable_name": requirement["sheet_variable_name"],
                "expected_source_from_docx": requirement["sheet_source"],
                "expected_frequency_from_docx": requirement["sheet_frequency"],
                "expected_la_availability_from_docx": requirement[
                    "sheet_la_availability"
                ],
                "expected_notes_from_docx": requirement["sheet_notes"],
                "why_hard_from_docx": requirement["why_hard"],
                "best_public_workaround_from_docx": requirement[
                    "best_public_workaround"
                ],
                **implementation,
                "native_frequency": "",
                "analytical_frequency": "",
                "geography_level": "",
                "geography_name": "",
                "raw_source_present": bool(source_ids)
                and all(item in manifest.index for item in source_ids),
                "clean_table_present": bool(clean_names)
                and any(item in clean_tables for item in clean_names),
                "fact_table_present": bool(implementation["implemented_variable_id"])
                and _fact_stats(fact, implementation["implemented_variable_id"])["row_count"] > 0,
                "panel_or_view_present": bool(implementation["implemented_table"]),
                "start_date": stats["start_date"],
                "end_date": stats["end_date"],
                "row_count": stats["row_count"],
                "missingness_summary": stats["missingness_summary"],
                "recommended_next_action": (
                    "Ready for data-collection use"
                    if implementation["current_status"]
                    in {"Implemented", "Scenario implemented", "Proxy implemented"}
                    else "Review documented limitation before modeling"
                ),
            }
        )
    return pd.DataFrame(rows)


def _data_catalog(config: ProjectConfig, fact: pd.DataFrame) -> pd.DataFrame:
    manifest = _latest_manifest_rows(config)
    rows = []
    for source in config.sources:
        manifest_row = (
            manifest.loc[source.source_id].to_dict()
            if source.source_id in manifest.index
            else {}
        )
        if source.source_id == HUD_USPS_SOURCE_ID:
            manifest_row = _hud_usps_catalog_manifest_row(config)
        variables = list(source.key_variables)
        stats = _source_stats(config, source)
        lifecycle_status = _source_lifecycle_status(source)
        update_command = (
            "BLOCKED: resolve the source's documented provenance or access "
            "issue before running its download workflow."
            if lifecycle_status == "Quarantined"
            else (
                f"python run_pipeline.py --stage download --source "
                f"{source.source_id} --force-download"
                if source.auto_download
                else f"python run_pipeline.py --stage download --source "
                f"{source.source_id} --local-file PATH_TO_OFFICIAL_FILE"
            )
        )
        catalog_file_row = manifest_row
        raw_path = _raw_path(config, catalog_file_row)
        rows.append(
            {
                "Source": source.dataset_name,
                "Frequency": source.native_frequency,
                "Geography": "|".join(source.geography_levels),
                "Key variables": "|".join(variables),
                "source_id": source.source_id,
                "provider": source.provider,
                "product": source.dataset_name,
                "official_source_page": source.source_page_url,
                "exact_download_url_or_api": source.download_url,
                "access_method": source.access_method,
                "active_status": lifecycle_status,
                "provenance_status": source.provenance_status,
                "observed_constructed_proxy_or_scenario": (
                    "observed tract components|constructed county aggregates|"
                    "modeled postal ZIP county-part proxy"
                    if source.source_id == HUD_USPS_SOURCE_ID
                    else
                    "constructed proxy (quarantined)"
                    if lifecycle_status == "Quarantined"
                    else "constructed published-table sensitivity"
                    if source.source_id in SIPP_SOURCE_IDS
                    else "observed published estimates|constructed reported-case proxy"
                    if source.source_id == AHS_TABLE_CREATOR_SOURCE_ID
                    else "observed"
                ),
                "manual_download_required": not source.auto_download,
                "manual_reason": source.manual_reason,
                "raw_file_format": (
                    "json|txt (75 immutable artifacts indexed by CSV manifest)"
                    if source.source_id == HUD_USPS_SOURCE_ID
                    else source.file_format
                ),
                "raw_file_path": catalog_file_row.get("relative_path"),
                "source_manifest_path": catalog_file_row.get(
                    "source_manifest_path"
                ),
                "original_filename": catalog_file_row.get("original_filename")
                or catalog_file_row.get("source_filename"),
                "retrieval_date": manifest_row.get("retrieved_at"),
                "provider_vintage": manifest_row.get("provider_vintage")
                or manifest_row.get("provider_last_updated")
                or manifest_row.get("source_last_updated")
                or manifest_row.get("data_vintage")
                or "",
                "checksum_sha256": catalog_file_row.get("checksum_sha256")
                or "",
                "native_frequency": source.native_frequency,
                "analytical_frequency": source.native_frequency,
                "geography_level": "|".join(source.geography_levels),
                "geography_name": "|".join(source.geography_levels),
                "geography_id": "",
                "start_date": stats["start_date"],
                "end_date": stats["end_date"],
                "row_count": stats["row_count"],
                "variable_ids": "|".join(variables),
                "seasonal_adjustment": "",
                "smoothing": "",
                "units": "",
                "destination_clean_table": (
                    ""
                    if lifecycle_status == "Quarantined"
                    else "clean.hud_usps_ncwm_quarterly|"
                    "data/restricted/hud_usps/active/"
                    "long_observations_with_lineage.parquet"
                    if source.source_id == HUD_USPS_SOURCE_ID
                    else _source_clean_name(source)
                ),
                "destination_analytical_table": (
                    ""
                    if lifecycle_status == "Quarantined"
                    else "analytics.fact_housing_observation|"
                    + "|".join(HUD_USPS_RESTRICTED_TABLES[:3])
                    if source.source_id == HUD_USPS_SOURCE_ID
                    else source.destination_analytical_table
                    or "analytics.fact_housing_observation"
                ),
                "update_command": update_command,
                "missing_value_policy": "preserve null",
                "revision_policy": source.revision_notes,
                "known_limitations": source.methodology_notes,
                "license_or_use_notes": source.license_or_use_notes,
                "last_verified": config.verified_at,
                "notes": (
                    "Official 75-artifact manifest verified at "
                    f"{HUD_USPS_SOURCE_MANIFEST_PATH}; active licensed detail "
                    "is restricted. The former package file was deleted; its "
                    "historical record remains only in "
                    f"{HUD_USPS_HISTORICAL_QUARANTINE_MANIFEST_PATH} as a "
                    "historical, superseded record."
                    if source.source_id == HUD_USPS_SOURCE_ID
                    else "Intentional quarantine; excluded from active ingestion"
                    if lifecycle_status == "Quarantined"
                    else "Raw file exists"
                    if raw_path and raw_path.exists()
                    else "No local raw file"
                ),
            }
        )
    existing_ids = {str(row["source_id"]) for row in rows}
    for source in _version_a_catalog_records(config):
        if source["source_id"] in existing_ids:
            continue
        variable_ids = tuple(source["variable_ids"])
        rows.append(
            {
                "Source": source["dataset_name"],
                "Frequency": source["native_frequency"],
                "Geography": source["geography"],
                "Key variables": "|".join(variable_ids),
                "source_id": source["source_id"],
                "provider": source["provider"],
                "product": source["dataset_name"],
                "official_source_page": source["source_page_url"],
                "exact_download_url_or_api": source["download_url"],
                "access_method": source["access_method"],
                "active_status": source["active_status"],
                "provenance_status": source["provenance_status"],
                "observed_constructed_proxy_or_scenario": source[
                    "observed_constructed_proxy_or_scenario"
                ],
                "manual_download_required": source[
                    "manual_download_required"
                ],
                "manual_reason": source["manual_reason"],
                "raw_file_format": source["file_format"],
                "raw_file_path": source["raw_file_path"],
                "original_filename": source["original_filename"],
                "retrieval_date": source["retrieved_at"],
                "provider_vintage": source["provider_vintage"],
                "checksum_sha256": source["checksum_sha256"],
                "native_frequency": source["native_frequency"],
                "analytical_frequency": source["native_frequency"],
                "geography_level": source["geography"],
                "geography_name": source["geography"],
                "geography_id": "",
                "start_date": source["start_date"],
                "end_date": source["end_date"],
                "row_count": source["row_count"],
                "variable_ids": "|".join(variable_ids),
                "seasonal_adjustment": "not applicable",
                "smoothing": "none",
                "units": "source-specific",
                "destination_clean_table": source[
                    "destination_clean_table"
                ],
                "destination_analytical_table": source[
                    "destination_analytical_table"
                ],
                "update_command": source["update_command"],
                "missing_value_policy": "preserve null",
                "revision_policy": "Archive every immutable source vintage.",
                "known_limitations": source["known_limitations"],
                "license_or_use_notes": source["license_or_use_notes"],
                "last_verified": config.verified_at,
                "notes": source["notes"],
            }
        )
        existing_ids.add(str(source["source_id"]))
    first_columns = ["Source", "Frequency", "Geography", "Key variables"]
    catalog = pd.DataFrame(rows)
    return catalog[first_columns + [col for col in catalog.columns if col not in first_columns]]


def _raw_repository_completeness(config: ProjectConfig) -> pd.DataFrame:
    manifest = _latest_manifest_rows(config)
    registry_ids = {source.source_id for source in config.sources}
    rows = []
    for source in config.sources:
        manifest_row = (
            manifest.loc[source.source_id].to_dict()
            if source.source_id in manifest.index
            else {}
        )
        raw_path = _raw_path(config, manifest_row)
        checks = {
            "raw_snapshot_present": bool(raw_path and raw_path.exists()),
            "manifest_present": source.source_id in manifest.index,
            "registry_present": source.source_id in registry_ids,
            "checksum_present": bool(manifest_row.get("checksum_sha256")),
            "retrieval_timestamp_present": bool(manifest_row.get("retrieved_at")),
            "source_url_present": bool(source.download_url or source.source_page_url),
            "frequency_present": bool(source.native_frequency),
            "geography_present": bool(source.geography_levels),
            "key_variables_present": bool(source.key_variables),
        }
        if source.source_id == HUD_USPS_SOURCE_ID:
            checks.update(_hud_usps_manifest_checks(config))
        issues = [name for name, ok in checks.items() if not ok]
        lifecycle_status = _source_lifecycle_status(source)
        if lifecycle_status == "Active":
            completeness_status = "PASS" if not issues else "FAIL"
        elif lifecycle_status == "Quarantined":
            completeness_status = "QUARANTINED"
        else:
            completeness_status = "INACTIVE"
        rows.append(
            {
                "source_id": source.source_id,
                "provider": source.provider,
                "dataset_name": source.dataset_name,
                "active_status": lifecycle_status,
                "provenance_status": source.provenance_status,
                **checks,
                "completeness_status": completeness_status,
                "issue": "|".join(issues),
            }
        )
    return pd.DataFrame(rows)


def _proxy_methodology_assumptions() -> pd.DataFrame:
    rows = [
        {
            "methodology_id": "version_a_first_mortgage_debt_benchmark_v1",
            "variable_id": "acs_implied_first_mortgage_balance_cal_benchmark",
            "assumption": "Authoritative aggregate mortgage-debt method",
            "value": "Use direct ACS B25089/B25081 source-period levels, weighted PUMS first-mortgage P&I parameters, an exact-date joint NMDB outstanding rate-term distribution, and an exact-date universe-compatible California UPB control.",
            "rationale": "This is the consolidated Version A standard. PMMS may be retained only as a labeled sensitivity fallback; calibrated values remain null when the required California control is unavailable.",
        },
        {
            "methodology_id": "version_a_owner_equity_benchmark_v1",
            "variable_id": "acs_implied_owner_occupied_home_equity_all_benchmark",
            "assumption": "Authoritative aggregate home-equity method",
            "value": "Subtract Version A California-calibrated first-mortgage debt from published ACS B25082 mortgaged-owner aggregate value, then add published mortgage-free owner value in full for the all-owner measure.",
            "rationale": "The value and debt universes remain matched. No Zillow level, median-times-count construction, or B25106 payment-bin balance substitutes for the Version A benchmark.",
        },
        {
            "methodology_id": "sipp_ca_primary_residence_debt_sensitivity_v1",
            "variable_id": "sipp_ca_aggregate_primary_residence_debt_sensitivity",
            "assumption": "Published-table state sensitivity formula",
            "value": "California households (14,420 thousand) * published Home Debt holding share (32.6 percent) * published conditional mean Home Debt ($346,200) = $1,627,458,504,000 for 2024.",
            "rationale": "This 2025 SIPP public table product is separately labeled because the inputs are rounded, the conditional mean is limited to households with total debt below the 99th percentile, and Home Debt is broader than the exact owner-occupied closed-end first-lien calibration universe. No county allocation or covariance-based aggregate standard error is created.",
        },
        {
            "methodology_id": "sipp_ca_positive_home_equity_sensitivity_v1",
            "variable_id": "sipp_ca_aggregate_positive_home_equity_sensitivity",
            "assumption": "Published-table state sensitivity formula",
            "value": "California households (14,420 thousand) * published Equity in Own Home holding share (51.0 percent) * published conditional mean Equity in Own Home ($668,100) = $4,913,341,020,000 for 2024.",
            "rationale": "This 2025 SIPP public table product is separately labeled because it covers positive-equity holders, uses rounded inputs, and the conditional mean is limited to households with net worth below the 99th percentile. It is not a matched-universe county equity benchmark, and no covariance-based aggregate standard error is created.",
        },
        {
            "methodology_id": "ahs_table14a_reported_case_outstanding_principal_proxy_v1",
            "variable_id": "ahs_ca_reported_case_primary_mortgage_outstanding_principal_proxy",
            "assumption": "AHS reported-case public calibration proxy",
            "value": "Sum of fourteen published 2023 California reported outstanding-principal-bin household estimates (3,381.2 thousand) * 1,000 * published conditional mean outstanding principal ($418,100) = $1,413,679,720,000.",
            "rationale": "The Four Variable Data Status Report authorizes AHS as the primary public proxy. The construction excludes 1,203.8 thousand regular-mortgage households with unreported balances, does not assign them the reported-case mean, and emits no all-household extrapolation or covariance-based aggregate uncertainty. AHS primary mortgage may not equal legal first lien, and the 2023 statewide proxy does not silently populate county targets calibrated to the current Version A universe and 2022-07-01 reference date.",
        },
        {
            "methodology_id": "acs_rental_availability_proxy_v1",
            "variable_id": "acs_rental_availability_rate",
            "assumption": "Matched annual ACS rental-availability proxy",
            "value": "For each matched 2010-2024 ACS five-year county vintage, B25004 for-rent vacant units / (B25003 renter-occupied units + B25004 for-rent vacant units). Both component 90% margins of error are retained; the rate is suppressed when either component MOE is greater than or equal to its estimate. Annual values are not interpolated or expanded to monthly frequency.",
            "rationale": "This preserves a transparent annual rental-availability proxy with matched ACS universes and explicit reliability screening. It is not an active rental-listing stock, and it remains separate from the quarterly HUD-USPS all-residential proxy.",
        },
        {
            "methodology_id": "acs_noncontrolled_structural_exposure_bounds_v1",
            "variable_id": "acs_renter_occupied_1_unit_share|acs_renter_occupied_2_to_4_units_share|acs_renter_occupied_5_plus_units_share|acs_renter_occupied_mobile_other_share",
            "assumption": "Same-vintage B25032 renter structure shares",
            "value": "Divide the mutually exclusive B25032 renter-occupied structure groups by the same-vintage B25032 renter total. Combine detached and attached for one unit; combine two and three-to-four for two-to-four; combine all published five-plus cells; combine mobile home and boat-RV-other.",
            "rationale": "These are transparent annual descriptive structure shares at county or Census-ZCTA geography. They do not by themselves classify rent-control coverage, and they are never interpolated or expanded to months.",
        },
        {
            "methodology_id": "acs_noncontrolled_structural_exposure_bounds_v1",
            "variable_id": "acs_renter_recent_construction_share_lower_bound|acs_renter_recent_construction_share_upper_bound",
            "assumption": "Vintage-normalized B25127 recent-construction bounds",
            "value": "Lower = normalized definitely-recent renter total / B25127 renter total. Upper = (definitely-recent renter total + cutoff-mixed renter total) / B25127 renter total. Regimes are 2000+ and 1980-1999 for 2010-2014; 2010+ and 2000-2009 for 2015-2020; and 2020+ and 2000-2019 for 2021-2024.",
            "rationale": "Published B25127 year-built bins and cell sequences change across vintages and cannot recover an exact rolling 15-year cutoff. Resolving cells by label and retaining a lower/upper interval prevents false precision.",
        },
        {
            "methodology_id": "acs_noncontrolled_structural_exposure_bounds_v1",
            "variable_id": "noncontrolled_exposure_share_structural_proxy_lower_bound",
            "assumption": "2020 publication floor and lower structural proxy",
            "value": "For ACS observations ending in 2020 or later, publish normalized B25127 definitely-recent renter total / B25127 renter total; preserve null before 2020.",
            "rationale": "The 2020 floor corresponds to the AB 1482 effective period but is not proof that any unit is legally exempt or noncontrolled.",
        },
        {
            "methodology_id": "acs_noncontrolled_structural_exposure_bounds_v1",
            "variable_id": "noncontrolled_exposure_share_structural_proxy_upper_bound",
            "assumption": "Broad recent-age and one-unit union upper proxy",
            "value": "For 2020 onward, min(1, (B25127 definitely-recent total + B25127 cutoff-mixed total + B25032 one-unit total - B25127 definitely-recent one-unit - B25127 cutoff-mixed one-unit) / B25127 renter total).",
            "rationale": "Subtracting the two B25127 one-unit components prevents double counting. Treating every remaining one-unit rental as potentially exempt makes this deliberately broad because ACS does not observe ownership and tenant-notice conditions.",
        },
        {
            "methodology_id": "acs_noncontrolled_structural_exposure_bounds_v1",
            "variable_id": "noncontrolled_exposure_share",
            "assumption": "Exact point estimate remains absent",
            "value": "config/rent_control_policy_crosswalk.csv is an official-source scaffold marked scaffold_only_not_applied. It is incomplete across incorporated municipalities and has no property-level join to the annual ACS aggregates.",
            "rationale": "A defensible point estimate requires dated state and local rules, jurisdiction, property age, ownership and notice status, affordable-housing and other exemptions, plus an approved aggregation rule. COUNTY_06037 cannot represent an unincorporated-only ordinance, and Census ZCTAs are not USPS postal ZIPs.",
        },
        {
            "methodology_id": "acs_renter_occupied_vintage_change_v1",
            "variable_id": "net_change_occupied_apartments_annual",
            "assumption": "Consecutive ACS vintage difference",
            "value": "Subtract the immediately prior ending-year B25003 renter-occupied estimate only when both county vintages exist and their ending years are consecutive. Do not allocate or forward-fill the result to quarters.",
            "rationale": "Adjacent ACS five-year estimates overlap for four survey years. The result is a transparent vintage-to-vintage change proxy, not an observed one-year or quarterly flow.",
        },
        {
            "methodology_id": "acs_county_price_to_income_v1",
            "variable_id": "county_annual_price_to_income_acs",
            "assumption": "Matched-vintage Los Angeles County price-to-income ratio",
            "value": "For each 2010-2024 ACS ending vintage, divide B25077 median home value by B19013 median household income matched on Los Angeles County and the exact ending year. Preserve a null when either source estimate is null or invalid, and retain the construction only in county annual controls.",
            "rationale": "The construction uses matched official ACS five-year facts and does not create a fact observation or monthly/quarterly timing. Structural QA recomputes every eligible ratio from its two panel inputs.",
        },
        {
            "methodology_id": "acs_zcta_price_to_income_v1",
            "variable_id": "zcta_annual_price_to_income_acs",
            "assumption": "Matched-vintage Census-ZCTA price-to-income ratio",
            "value": "For each configured Census ZCTA and 2011-2024 ACS ending vintage, divide B25077 median home value by B19013 median household income from the same ZCTA and vintage. Preserve a null when either source estimate is null or invalid, and do not expand the annual estimate to quarters.",
            "rationale": "The construction uses only matched official ACS five-year facts. A Census ZCTA is not a USPS postal ZIP, and published input gaps remain visible rather than being filled.",
        },
        {
            "methodology_id": "acs_zcta_renter_occupied_vintage_change_v1",
            "variable_id": "zcta_renter_occupied_units_vintage_change_annual",
            "assumption": "Consecutive ACS B25003 Census-ZCTA vintage difference",
            "value": "Within each exact configured Census ZCTA, subtract the immediately prior ending-year B25003 renter-occupied estimate only when both published five-year vintages exist and their ending years are consecutive. Do not allocate, interpolate, or forward-fill the result to quarters.",
            "rationale": "Adjacent ACS five-year estimates overlap for four survey years. The result is a vintage-to-vintage change, not an observed one-year or quarterly flow, and the Census ZCTA is never relabeled as a postal ZIP.",
        },
        {
            "methodology_id": "hud_usps_residential_vacancy_proxy_v2",
            "variable_id": "usps_residential_vacancy_rate",
            "assumption": "All-residential vacancy-proxy formula",
            "value": "(short-term vacant residential addresses + long-term vacant residential addresses) / total residential addresses; rates remain ratios in [0, 1].",
            "rationale": "No-stat residential addresses remain a separate count and rate diagnostic and never enter the vacancy numerator. This is an all-residential address-status proxy, not rental-listing inventory.",
        },
        {
            "methodology_id": "hud_usps_residential_vacancy_proxy_v2",
            "variable_id": "usps_residential_vacancy_rate",
            "assumption": "Official active provenance",
            "value": "provenance_status=official_raw_archive_rebuilt; 75 immutable official artifacts are indexed by metadata/hud_usps_source_manifest.csv, including 73 contiguous 2008 Q1-2026 Q1 NCWM responses, the 2026 Q1 tract-to-postal-ZIP crosswalk, and the 2020 Census California county reference.",
            "rationale": "Every active observation is rebuilt from the verified official archive. The former derived package file was deleted; only its historical, superseded quarantine manifest remains, and it is not an active input.",
        },
        {
            "methodology_id": "hud_usps_residential_vacancy_proxy_v2",
            "variable_id": "usps_residential_vacancy_rate",
            "assumption": "Geography-specific construction",
            "value": "Provider component fields are direct at standardized 2020-tract geography; California county values are constructed from complete tract components; Los Angeles County postal-ZIP county-part values are modeled with the fixed 2026 Q1 crosswalk.",
            "rationale": "Modeled ZIP rows describe only the Los Angeles County portion of USPS postal ZIPs. They are neither whole postal ZIPs nor Census ZCTAs, and no tract percentages are averaged.",
        },
        {
            "methodology_id": "hud_usps_residential_vacancy_proxy_v2",
            "variable_id": "usps_residential_no_stat_rate",
            "assumption": "Quarterly-only restricted diagnostic",
            "value": "No-stat / total remains a separate native-quarterly diagnostic. The eight active variables are stored in the common fact and restricted native-quarterly tables; they are not expanded into observed months.",
            "rationale": "Licensed tract and modeled postal-ZIP county-part observations and derivatives remain access-controlled under data/restricted/hud_usps/ and outside public exports under the HUD-USPS sublicense.",
        },
        {
            "methodology_id": "acs_implied_mortgage_debt_proxy_v1",
            "variable_id": "acs_implied_aggregate_mortgage_debt_proxy",
            "assumption": "Legacy ACS B25106 table role",
            "value": "Legacy scenario only: B25106 owner income x housing-cost-burden bins infer monthly owner housing costs; B25106 is not an observed mortgage balance and is not the authoritative Version A method.",
            "rationale": "Retained in facts for reproducibility and labeled legacy; do not select it in place of the Version A _benchmark variables.",
        },
        {
            "methodology_id": "acs_implied_mortgage_debt_proxy_v1",
            "variable_id": "acs_implied_aggregate_mortgage_debt_proxy",
            "assumption": "Income bin midpoints",
            "value": "10000; 27500; 42500; 62500; 112500",
            "rationale": "Closed bins use midpoint values. The open-ended $75k+ bin uses $112,500, or 1.5 times the lower bound, as a transparent conservative top-bin assumption.",
        },
        {
            "methodology_id": "acs_implied_mortgage_debt_proxy_v1",
            "variable_id": "acs_implied_aggregate_mortgage_debt_proxy",
            "assumption": "Housing-cost-burden midpoints",
            "value": "0.10 for <20%; 0.25 for 20-29%; 0.40 for 30%+",
            "rationale": "Closed or semi-closed categories use practical midpoints; 30%+ uses 40% to avoid implying unlimited burden while retaining high-burden signal.",
        },
        {
            "methodology_id": "acs_implied_mortgage_debt_proxy_v1",
            "variable_id": "acs_implied_aggregate_mortgage_debt_proxy",
            "assumption": "Debt-service share of ACS owner housing cost",
            "value": "0.80",
            "rationale": "ACS owner housing costs include non-debt items such as taxes, insurance, utilities, and fees; 80% assigns most but not all inferred owner cost to mortgage debt service.",
        },
        {
            "methodology_id": "acs_implied_mortgage_debt_proxy_v1",
            "variable_id": "acs_implied_aggregate_mortgage_debt_proxy",
            "assumption": "Mortgage-rate assumption",
            "value": "Mean Freddie Mac PMMS monthly-average 30-year fixed rate over the ACS five-year estimate window.",
            "rationale": "The ACS 5-year estimate pools observations over five years; using the same window avoids using only one endpoint-year rate.",
        },
        {
            "methodology_id": "acs_implied_mortgage_debt_proxy_v1",
            "variable_id": "acs_implied_aggregate_mortgage_debt_proxy",
            "assumption": "Remaining amortization term",
            "value": "25 years",
            "rationale": "Outstanding balances are on existing mortgages, not all new 30-year loans; 25 years is a transparent middle-ground remaining-term assumption.",
        },
        {
            "methodology_id": "acs_implied_mortgage_debt_proxy_v1",
            "variable_id": "acs_implied_aggregate_mortgage_debt_proxy",
            "assumption": "Allocation of B25081 mortgaged units",
            "value": "B25081 owner-with-mortgage count is allocated across positive-income B25106 bins proportional to B25106 owner unit counts.",
            "rationale": "B25106 bins are not mortgage-status-specific. This forces aggregate mortgaged units to match the directly observed B25081 count while preserving the B25106 income/burden distribution.",
        },
        {
            "methodology_id": "acs_implied_mortgage_debt_proxy_v1",
            "variable_id": "acs_implied_aggregate_mortgage_debt_proxy",
            "assumption": "Zero/negative income handling",
            "value": "B25106 zero/negative-income owner units are preserved as an input but excluded from payment-to-balance inference.",
            "rationale": "A positive income midpoint is required to infer monthly owner housing cost from cost-burden percentages.",
        },
        {
            "methodology_id": "ca_homeowners_premium_real_2023_cpi_v1",
            "variable_id": "ca_homeowners_avg_written_premium_real",
            "assumption": "California CPI deflator and base year",
            "value": "nominal California DOI average written premium in year y multiplied by California DIR annual CPI-U for 2023 and divided by the same annual CPI-U for year y.",
            "rationale": "This implements the workbook-approved constant-2023-dollar convention on exact calendar-year matches. The 2023 real value must equal the 2023 nominal value; missing or nonpositive inputs remain null, and no subannual CPI or real-premium observations are fabricated.",
        },
        {
            "methodology_id": "bps_5_plus_distributed_completion_lag_v1",
            "variable_id": "new_apartments_completed_proxy_central",
            "assumption": "Central completion timing",
            "value": "Allocate 20% of each 5+ permit cohort equally across lags 12-17, 55% equally across lags 18-23, and 25% equally across lags 24-30.",
            "rationale": "The weights are the exact approved Central scenario and sum to one. All lags 12-30 must be present; a missing source month propagates to a null proxy month rather than being treated as zero.",
        },
        {
            "methodology_id": "bps_5_plus_distributed_completion_lag_v1",
            "variable_id": "new_apartments_completed_proxy_slow",
            "assumption": "Slow completion timing",
            "value": "Allocate 20% of each 5+ permit cohort equally across lags 18-23, 50% equally across lags 24-29, and 30% equally across lags 30-36.",
            "rationale": "The weights are the exact approved Slow scenario and sum to one. All lags 18-36 must be present; a missing source month propagates to a null proxy month rather than being treated as zero.",
        },
        {
            "methodology_id": "bps_5_plus_stock_flow_pipeline_v1",
            "variable_id": "multifamily_under_construction_proxy_central",
            "assumption": "Central pipeline recurrence and initialization",
            "value": "Month-end stock equals prior stock plus current 5+ permits minus Central modeled completions; initialize from every cohort with positive uncompleted weight in the complete current-through-age-29 permit window.",
            "rationale": "The approved 30-month active-cohort initialization avoids an arbitrary zero opening stock. Missing active-window inputs produce null until a complete window permits reinitialization.",
        },
        {
            "methodology_id": "bps_5_plus_stock_flow_pipeline_v1",
            "variable_id": "multifamily_under_construction_proxy_slow",
            "assumption": "Slow pipeline recurrence and initialization",
            "value": "Month-end stock equals prior stock plus current 5+ permits minus Slow modeled completions; initialize from every cohort with positive uncompleted weight in the complete current-through-age-35 permit window.",
            "rationale": "The approved 36-month active-cohort initialization avoids an arbitrary zero opening stock. Missing active-window inputs produce null until a complete window permits reinitialization.",
        },
        {
            "methodology_id": "irs_soi_net_exemptions_per_1000_pep_v1",
            "variable_id": "census_pep_county_population_annual",
            "assumption": "PEP population vintage segments",
            "value": "Use Vintage 2019 for 2010-2019 and Vintage 2025 for 2020-2025, retaining July 1 source-year estimates as annual December 31 analytic endpoints.",
            "rationale": "The segments are nonoverlapping and preserve the 2020 vintage boundary. PEP annual values are not interpolated or expanded to a higher frequency.",
        },
        {
            "methodology_id": "irs_soi_net_exemptions_per_1000_pep_v1",
            "variable_id": "net_exemption_rate",
            "assumption": "Annual population denominator",
            "value": "IRS SOI net migration exemptions for each filing-year interval divided by Census PEP resident population for the same ending year, multiplied by 1,000.",
            "rationale": "The workbook expressly selects the interval ending year. Prior-year and ACS population denominators are not substituted, and the annual rate is not divided or carried into observed months.",
        },
        {
            "methodology_id": "bps_1_unit_per_1000_pep_housing_units_v1",
            "variable_id": "census_pep_county_housing_units_annual",
            "assumption": "PEP housing-unit vintage segments",
            "value": "Use Vintage 2019 for 2010-2019 and Vintage 2025 for 2020-2025, retaining July 1 source-year estimates as annual December 31 analytic endpoints.",
            "rationale": "The segments are nonoverlapping and preserve the 2020 vintage boundary. PEP annual values are not interpolated or expanded to a higher frequency.",
        },
        {
            "methodology_id": "bps_1_unit_per_1000_pep_housing_units_v1",
            "variable_id": "permits_1_unit_per_1000_housing_units",
            "assumption": "Annual housing-stock denominator",
            "value": "Final annual Census BPS one-unit authorizations divided by the same-year Census PEP housing-unit estimate, multiplied by 1,000.",
            "rationale": "The approved rate uses the final annual BPS product and same-year PEP housing stock. Monthly permits, ACS housing units, and lagged denominators are not substituted; the output remains annual.",
        },
        {
            "methodology_id": "eviction_lab_county_estimates_v1",
            "variable_id": "eviction_lab_filings_estimate_annual",
            "assumption": "LA eviction-source choice",
            "value": "Retain official Eviction Lab annual county modeled estimates for Los Angeles County for 2000-2018. The official ETS Version 3.0 release was checked on 2026-08-16 and contains neither Los Angeles nor any California geography; the observed/validated county analysis files also provide no LA County rows.",
            "rationale": "This preserves the expected source family without silently substituting another geography. Annual values are not interpolated or expanded to monthly observations, and the limitation remains explicit for modeling.",
        },
        {
            "methodology_id": "eviction_lab_county_rate_v1",
            "variable_id": "eviction_lab_filings_rate_estimate_annual",
            "assumption": "Annual rate denominator",
            "value": "filings_estimate divided by Eviction Lab renting_hh, multiplied by 100.",
            "rationale": "A rate normalizes the trend for renter-household base changes while using only fields published in the same official Eviction Lab file.",
        },
        {
            "methodology_id": "eviction_lab_county_rate_v1",
            "variable_id": "eviction_lab_households_threatened_rate_estimate_annual",
            "assumption": "Annual rate denominator",
            "value": "hh_threat_estimate divided by Eviction Lab renting_hh, multiplied by 100.",
            "rationale": "A rate normalizes households-threatened estimates for renter-household base changes while using only fields published in the same official Eviction Lab file.",
        },
    ]
    frame = pd.DataFrame(rows)
    legacy = frame["variable_id"].eq(
        "acs_implied_aggregate_mortgage_debt_proxy"
    )
    frame.loc[legacy, "methodology_id"] = (
        "acs_implied_mortgage_debt_proxy_legacy_v0"
    )
    frame["methodology_status"] = "active"
    frame.loc[legacy, "methodology_status"] = "legacy"
    frame.loc[
        frame["variable_id"].eq("noncontrolled_exposure_share"),
        "methodology_status",
    ] = "blocked"
    return frame


def _source_tracker_summary(catalog: pd.DataFrame) -> pd.DataFrame:
    """Compact, human-readable source tracker sheet.

    The full data catalog intentionally keeps every reproducibility field. This
    sheet is the first place a mentor/reviewer should look when skimming what
    data exists, where it came from, and how current it is.
    """
    columns = [
        "Source",
        "provider",
        "active_status",
        "Frequency",
        "Geography",
        "Key variables",
        "start_date",
        "end_date",
        "row_count",
        "observed_constructed_proxy_or_scenario",
        "retrieval_date",
        "raw_file_path",
        "update_command",
        "known_limitations",
    ]
    labels = {
        "provider": "Provider",
        "active_status": "Status",
        "start_date": "Start date",
        "end_date": "End date",
        "row_count": "Rows",
        "observed_constructed_proxy_or_scenario": "Observed / proxy / scenario",
        "retrieval_date": "Retrieved at",
        "raw_file_path": "Raw file",
        "update_command": "Refresh command",
        "known_limitations": "Limitations",
    }
    available = [column for column in columns if column in catalog.columns]
    return catalog.loc[:, available].rename(columns=labels)


def _source_tracker_overview(
    catalog: pd.DataFrame,
    audit: pd.DataFrame,
    dictionary: pd.DataFrame,
    coverage: pd.DataFrame,
    manifest: pd.DataFrame,
    raw_completeness: pd.DataFrame,
    gap: pd.DataFrame,
    assumptions_frame: pd.DataFrame,
) -> pd.DataFrame:
    """Workbook landing page explaining how to read the source tracker."""
    active_sources = int(catalog["active_status"].eq("Active").sum())
    completeness_pass = int(raw_completeness["completeness_status"].eq("PASS").sum())
    completeness_fail = int(raw_completeness["completeness_status"].eq("FAIL").sum())
    earliest = catalog["start_date"].dropna().min()
    latest = catalog["end_date"].dropna().max()
    return pd.DataFrame(
        [
            {
                "Item": "Purpose",
                "Value": (
                    "Human-readable tracker for current Los Angeles housing "
                    "database requested variables, active sources, coverage, "
                    "raw-file completeness, gaps, and proxy assumptions."
                ),
            },
            {
                "Item": "Generated by",
                "Value": "python run_pipeline.py --stage catalog",
            },
            {
                "Item": "Do not hand edit",
                "Value": (
                    "This workbook is regenerated from config, source manifests, "
                    "DuckDB coverage, and pipeline metadata. Change the pipeline "
                    "inputs/code, then regenerate."
                ),
            },
            {
                "Item": "How to skim",
                "Value": (
                    "Start with Source Summary for provider/source status, "
                    "DOCX Variable Audit for requested-variable status, "
                    "Variable Dictionary for active variable definitions, "
                    "Coverage Summary for date/geography coverage, Data Catalog "
                    "for full source metadata, Raw Completeness for archival "
                    "checks, Missing Deferred for gaps, and Proxy Assumptions "
                    "for constructed variables."
                ),
            },
            {
                "Item": "Tracker scope",
                "Value": (
                    "This workbook tracks metadata, source coverage, variable "
                    "definitions, and implementation status. It does not duplicate "
                    "every observation row; the observations live in DuckDB, "
                    "Parquet, and CSV exports."
                ),
            },
            {
                "Item": "Eviction Trends location",
                "Value": (
                    "Requested variable row: DOCX Variable Audit. Source row: "
                    "Data Catalog source_id eviction_lab_county_estimates. "
                    "Implemented variable coverage: Coverage Summary rows whose "
                    "variable_id starts with eviction_lab_."
                ),
            },
            {"Item": "Catalog source rows", "Value": len(catalog)},
            {"Item": "DOCX variable audit rows", "Value": len(audit)},
            {"Item": "Active variable dictionary rows", "Value": len(dictionary)},
            {"Item": "Coverage summary rows", "Value": len(coverage)},
            {"Item": "Raw manifest rows", "Value": len(manifest)},
            {"Item": "Active sources", "Value": active_sources},
            {"Item": "Raw completeness pass", "Value": completeness_pass},
            {"Item": "Raw completeness fail", "Value": completeness_fail},
            {"Item": "Deferred / missing rows", "Value": len(gap)},
            {"Item": "Proxy assumption rows", "Value": len(assumptions_frame)},
            {"Item": "Earliest catalog date", "Value": earliest},
            {"Item": "Latest catalog date", "Value": latest},
        ]
    )


def _excel_table_name(sheet_name: str) -> str:
    name = re.sub(r"[^A-Za-z0-9_]", "_", sheet_name)
    name = re.sub(r"_+", "_", name).strip("_")
    if not name or not re.match(r"[A-Za-z_]", name):
        name = f"Table_{name}"
    return f"{name[:240]}Table"


def _format_tracker_sheet(
    worksheet,
    *,
    tab_color: str,
    table_style: str = "TableStyleMedium2",
) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_properties.tabColor = tab_color
    worksheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )
    thin_border = Border(
        bottom=Side(style="thin", color="B7C9D6"),
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border
    worksheet.row_dimensions[1].height = 34

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    if worksheet.max_row > 1 and worksheet.max_column > 1:
        ref = (
            f"A1:{get_column_letter(worksheet.max_column)}"
            f"{worksheet.max_row}"
        )
        table = Table(displayName=_excel_table_name(worksheet.title), ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name=table_style,
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)
        worksheet.auto_filter.ref = ref

    long_text_markers = {
        "url",
        "command",
        "limitation",
        "notes",
        "rationale",
        "value",
        "variables",
        "path",
        "filename",
        "issue",
        "detail",
        "workaround",
        "source",
    }
    for column_index in range(1, worksheet.max_column + 1):
        letter = get_column_letter(column_index)
        header = str(worksheet.cell(row=1, column=column_index).value or "")
        sample_values = [
            worksheet.cell(row=row, column=column_index).value
            for row in range(1, min(worksheet.max_row, 80) + 1)
        ]
        max_length = max(
            [len(str(value)) for value in sample_values if value is not None]
            or [len(header)]
        )
        cap = 62 if any(marker in header.lower() for marker in long_text_markers) else 28
        worksheet.column_dimensions[letter].width = max(12, min(cap, max_length + 2))


def _format_source_tracker_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    tab_colors = {
        "Overview": "17365D",
        "Source Summary": "2F75B5",
        "DOCX Variable Audit": "4472C4",
        "Variable Dictionary": "9E480E",
        "Coverage Summary": "00B0F0",
        "Data Catalog": "5B9BD5",
        "Raw File Inventory": "548235",
        "Raw Completeness": "70AD47",
        "Missing Deferred": "FFC000",
        "Proxy Assumptions": "A64D79",
        "Update Commands": "8064A2",
    }
    for worksheet in workbook.worksheets:
        _format_tracker_sheet(
            worksheet,
            tab_color=tab_colors.get(worksheet.title, "808080"),
        )
    workbook.active = 0
    workbook.save(path)


def _write_source_tracker_workbook(
    path: Path,
    catalog: pd.DataFrame,
    audit: pd.DataFrame,
    dictionary: pd.DataFrame,
    manifest: pd.DataFrame,
    coverage: pd.DataFrame,
    raw_completeness: pd.DataFrame,
    gap: pd.DataFrame,
    assumptions_frame: pd.DataFrame,
) -> None:
    overview = _source_tracker_overview(
        catalog,
        audit,
        dictionary,
        coverage,
        manifest,
        raw_completeness,
        gap,
        assumptions_frame,
    )
    source_summary = _source_tracker_summary(catalog)
    with _atomic_workbook_target(path) as workbook_path:
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            overview.to_excel(writer, sheet_name="Overview", index=False)
            source_summary.to_excel(writer, sheet_name="Source Summary", index=False)
            audit.to_excel(writer, sheet_name="DOCX Variable Audit", index=False)
            dictionary.to_excel(writer, sheet_name="Variable Dictionary", index=False)
            coverage.to_excel(writer, sheet_name="Coverage Summary", index=False)
            catalog.to_excel(writer, sheet_name="Data Catalog", index=False)
            manifest.to_excel(writer, sheet_name="Raw File Inventory", index=False)
            raw_completeness.to_excel(writer, sheet_name="Raw Completeness", index=False)
            gap.to_excel(writer, sheet_name="Missing Deferred", index=False)
            assumptions_frame.to_excel(
                writer,
                sheet_name="Proxy Assumptions",
                index=False,
            )
            catalog[["source_id", "update_command"]].to_excel(
                writer,
                sheet_name="Update Commands",
                index=False,
            )
        _format_source_tracker_workbook(workbook_path)


@contextmanager
def _atomic_workbook_target(path: Path):
    """Yield a same-directory temporary path, then atomically replace the target."""
    temp = path.with_name(f".{path.stem}.{uuid4().hex}.tmp.xlsx")
    try:
        yield temp
        temp.replace(path)
    finally:
        temp.unlink(missing_ok=True)


def _write_workbooks(
    config: ProjectConfig,
    catalog: pd.DataFrame,
    audit: pd.DataFrame,
    gap: pd.DataFrame,
    raw_completeness: pd.DataFrame,
    assumptions_frame: pd.DataFrame,
    dictionary: pd.DataFrame | None = None,
) -> None:
    metadata_dir = config.root / "metadata"
    manifest = _raw_manifest(config)
    dictionary = (
        dictionary.copy()
        if dictionary is not None
        else pd.DataFrame(config.variables)
    )
    requirements = pd.DataFrame(DOCX_VARIABLES)
    geography_matrix = (
        catalog.groupby(["Frequency", "Geography"])
        .size()
        .rename("source_count")
        .reset_index()
    )
    coverage_path = metadata_dir / "source_coverage.csv"
    coverage = pd.read_csv(coverage_path) if coverage_path.exists() else pd.DataFrame()
    assumptions = audit.loc[
        audit["observed_constructed_proxy_or_scenario"].isin(
            {"proxy", "scenario assumption", "deferred", "unavailable"}
        ),
        [
            "variable_name",
            "observed_constructed_proxy_or_scenario",
            "limitation",
            "recommended_next_action",
        ],
    ]
    with _atomic_workbook_target(metadata_dir / "data_catalog.xlsx") as workbook_path:
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            catalog.to_excel(writer, sheet_name="Data Catalog", index=False)
            dictionary.to_excel(writer, sheet_name="Variable Dictionary", index=False)
            requirements.to_excel(writer, sheet_name="DOCX Requirements", index=False)
            audit.to_excel(writer, sheet_name="DOCX Variable Audit", index=False)
            assumptions_frame.to_excel(
                writer, sheet_name="Proxy Assumptions", index=False
            )
            manifest.to_excel(writer, sheet_name="Raw File Inventory", index=False)
            geography_matrix.to_excel(
                writer, sheet_name="Geography-Frequency Matrix", index=False
            )
            coverage.to_excel(writer, sheet_name="Coverage Summary", index=False)
            audit.loc[audit["docx_category"].eq("Difficult Variables")].to_excel(
                writer,
                sheet_name="Difficult Variables",
                index=False,
            )
            gap.to_excel(writer, sheet_name="Missing Deferred", index=False)
            catalog[["source_id", "update_command"]].to_excel(
                writer,
                sheet_name="Update Commands",
                index=False,
            )
            assumptions.to_excel(
                writer, sheet_name="Assumptions Questions", index=False
            )
    with _atomic_workbook_target(
        metadata_dir / "docx_variable_audit.xlsx"
    ) as workbook_path:
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            requirements.to_excel(writer, sheet_name="DOCX Requirements", index=False)
            audit.to_excel(writer, sheet_name="DOCX Variable Audit", index=False)
            assumptions_frame.to_excel(
                writer, sheet_name="Proxy Assumptions", index=False
            )
    _write_gap_workbook_atomic(
        metadata_dir / "data_gap_analysis.xlsx",
        gap,
        "catalog-stage",
    )
    _write_source_tracker_workbook(
        metadata_dir / "LA_Housing_Source_Tracker.xlsx",
        catalog,
        audit,
        dictionary,
        manifest,
        coverage,
        raw_completeness,
        gap,
        assumptions_frame,
    )


def _write_gap_workbook_atomic(
    path: Path,
    gap: pd.DataFrame,
    build_version: str,
) -> None:
    """Write the generated gap workbook without truncating the prior copy first."""
    temp = path.with_name(f".{path.stem}.{uuid4().hex}.tmp.xlsx")
    try:
        with pd.ExcelWriter(temp, engine="openpyxl") as writer:
            gap.to_excel(writer, sheet_name="Gap Analysis", index=False)
        _format_workbook(temp, build_version)
        temp.replace(path)
    finally:
        temp.unlink(missing_ok=True)


def _sync_raw_completeness_to_duckdb(
    config: ProjectConfig,
    raw_completeness: pd.DataFrame,
) -> None:
    database_path = config.root / config.pipeline["database_path"]
    if not database_path.exists():
        return
    with duckdb.connect(str(database_path)) as connection:
        connection.execute("CREATE SCHEMA IF NOT EXISTS quality")
        connection.register("raw_completeness_frame", raw_completeness)
        connection.execute(
            "CREATE OR REPLACE TABLE quality.raw_repository_completeness AS "
            "SELECT * FROM raw_completeness_frame"
        )


def _write_catalog_preserving_monthly_finance_products(
    catalog: pd.DataFrame,
    path: Path,
) -> None:
    """Replace general catalog rows without discarding completed B/C metadata.

    The general reporting builders regenerate their portion of the catalog from
    configuration.  Completed monthly-finance products are registered by a
    separate pipeline and carry concrete checksums, cutoff vintages, row counts,
    and output paths that cannot be reconstructed by the reporting builder.
    Keep those six owned rows long enough for
    ``register_monthly_finance_contracts`` to refresh their definitions while
    preserving the completed-build fields.
    """

    generated = catalog.loc[
        ~catalog["source_id"].isin(PRODUCT_SOURCE_IDS)
    ].copy()
    preserved = pd.DataFrame(columns=generated.columns)
    if path.is_file():
        existing = pd.read_csv(path, dtype=str)
        if "source_id" in existing.columns:
            preserved = existing.loc[
                existing["source_id"].isin(PRODUCT_SOURCE_IDS)
            ].copy()
            preserved = preserved.reindex(columns=generated.columns, fill_value="")
    pd.concat([generated, preserved], ignore_index=True).to_csv(path, index=False)


def build_catalog_outputs(config: ProjectConfig) -> dict[str, int]:
    """Write audit/catalog outputs for the Word checklist and active repository."""
    metadata_dir = config.root / "metadata"
    qa_dir = config.root / "output" / "qa"
    metadata_dir.mkdir(parents=True, exist_ok=True)
    qa_dir.mkdir(parents=True, exist_ok=True)
    fact = _load_fact(config)
    audit = _docx_audit(config, fact)
    catalog = _data_catalog(config, fact)
    raw_completeness = _raw_repository_completeness(config)
    assumptions_frame = _proxy_methodology_assumptions()
    gap = audit.loc[
        ~audit["current_status"].isin(
            {"Implemented", "Scenario implemented", "Proxy implemented"}
        )
    ].copy()
    requirements = pd.DataFrame(DOCX_VARIABLES)
    requirements.to_csv(metadata_dir / "docx_variable_requirements.csv", index=False)
    assumptions_frame.to_csv(
        metadata_dir / "proxy_methodology_assumptions.csv", index=False
    )
    audit.to_csv(metadata_dir / "docx_variable_audit.csv", index=False)
    _write_catalog_preserving_monthly_finance_products(
        catalog,
        metadata_dir / "data_catalog.csv",
    )
    gap.to_csv(metadata_dir / "data_gap_analysis.csv", index=False)
    catalog.to_csv(metadata_dir / "source_discovery.csv", index=False)
    raw_completeness.to_csv(qa_dir / "raw_repository_completeness.csv", index=False)
    register_monthly_finance_contracts(root=config.root)
    catalog = pd.read_csv(metadata_dir / "data_catalog.csv")
    dictionary = pd.read_csv(metadata_dir / "data_dictionary.csv")
    _write_workbooks(
        config,
        catalog,
        audit,
        gap,
        raw_completeness,
        assumptions_frame,
        dictionary=dictionary,
    )
    _sync_raw_completeness_to_duckdb(config, raw_completeness)
    inventory_counts = build_system_inventory(config)
    return {
        "catalog_rows": len(catalog),
        "audit_rows": len(audit),
        "gap_rows": len(gap),
        "raw_completeness_rows": len(raw_completeness),
        **inventory_counts,
    }


# Compatibility catalog retained for the structural-variable pipeline.
CATALOG_COLUMNS = [
    "Source",
    "Frequency",
    "Geography",
    "Key variables",
    "source_id",
    "provider",
    "product",
    "official_source_page",
    "exact_download_url_or_api",
    "access_method",
    "active_status",
    "provenance_status",
    "observed_constructed_proxy_or_scenario",
    "manual_download_required",
    "raw_file_format",
    "raw_file_path",
    "source_manifest_path",
    "original_filename",
    "retrieval_date",
    "provider_vintage",
    "checksum_sha256",
    "native_frequency",
    "analytical_frequency",
    "geography_level",
    "geography_name",
    "geography_id",
    "start_date",
    "end_date",
    "row_count",
    "variable_ids",
    "seasonal_adjustment",
    "smoothing",
    "units",
    "destination_clean_table",
    "destination_analytical_table",
    "update_command",
    "missing_value_policy",
    "revision_policy",
    "known_limitations",
    "license_or_use_notes",
    "last_verified",
    "notes",
]

UNAVAILABLE_ROWS = [
    {
        "Source": "Rental listing inventory",
        "Frequency": "Unknown",
        "Geography": "Los Angeles County or ZIP",
        "Key variables": "rental_inventory",
        "source_id": "unavailable_rental_inventory",
        "provider": "No approved source named in the governing methodology",
        "product": "True rental-listing inventory",
        "active_status": "Inactive",
        "provenance_status": "unavailable_methodology_source",
        "known_limitations": (
            "True rental-listing inventory remains unavailable. ZORI is a rent "
            "index, not inventory. The active HUD/USPS measure is a quarterly "
            "all-residential address-vacancy proxy, not a rental-listing stock. "
            "Its modeled postal ZIP county-parts remain distinct from whole "
            "postal ZIPs and Census ZCTAs. A documented listing universe, stable "
            "geography, history, and revision policy would be required after the "
            "methodology names an approved qualifying source."
        ),
    },
    {
        "Source": "Income-based affordability ratios",
        "Frequency": "Annual or mixed",
        "Geography": "Los Angeles County or ZIP",
        "Key variables": "price_to_income|rent_to_income",
        "source_id": "deferred_income_affordability_ratios",
        "provider": "Requires explicit income denominator decision",
        "product": "Price-to-income and rent-to-income ratios",
        "active_status": "Deferred",
        "known_limitations": (
            "The current active stage uses rent/payment affordability metrics. "
            "Income-denominator ratios are later-stage work and are not exposed "
            "as active variables."
        ),
    },
    {
        "Source": "Treasury rate and mortgage spread",
        "Frequency": "Daily or weekly",
        "Geography": "United States",
        "Key variables": "treasury_rate|mortgage_spread",
        "source_id": "deferred_treasury_mortgage_spread",
        "provider": "U.S. Treasury or Federal Reserve",
        "product": "Treasury benchmark and constructed mortgage spread",
        "active_status": "Deferred",
        "known_limitations": "Optional control outside the current requested core.",
    },
]


def _manifest(config: ProjectConfig) -> pd.DataFrame:
    path = config.root / "metadata" / "raw_file_manifest.csv"
    frame = pd.read_csv(path, dtype=str) if path.exists() else pd.DataFrame()
    if not frame.empty:
        active_ids = {source.source_id for source in config.sources}
        frame = frame.loc[frame["source_id"].isin(active_ids)].copy()
    if not frame.empty:
        frame["_retrieved"] = pd.to_datetime(
            frame["retrieved_at"], utc=True, errors="coerce", format="mixed"
        )
    return frame


def build_raw_repository_completeness(
    config: ProjectConfig,
) -> pd.DataFrame:
    manifest = _manifest(config)
    database_path = config.root / config.pipeline["database_path"]
    registry_ids: set[str] = set()
    if database_path.exists():
        connection = duckdb.connect(str(database_path), read_only=True)
        registry_ids = {
            row[0]
            for row in connection.execute(
                "SELECT DISTINCT source_id FROM metadata.source_registry"
            ).fetchall()
        }
        connection.close()
    rows: list[dict] = []
    for source in config.sources:
        matches = (
            manifest.loc[manifest["source_id"].eq(source.source_id)]
            if not manifest.empty
            else pd.DataFrame()
        )
        paths = (
            [
                config.root / value
                for value in matches.get(
                    "relative_path", pd.Series(dtype=str)
                ).dropna()
            ]
            if not matches.empty
            else []
        )
        checks = {
            "raw_snapshot_present": bool(paths)
            and any(path.is_file() for path in paths),
            "manifest_present": not matches.empty,
            "registry_present": source.source_id in registry_ids,
            "checksum_present": not matches.empty
            and matches["checksum_sha256"].fillna("").str.len().eq(64).any(),
            "retrieval_timestamp_present": not matches.empty
            and pd.to_datetime(
                matches["retrieved_at"],
                utc=True,
                errors="coerce",
                format="mixed",
            )
            .notna()
            .any(),
            "source_url_present": bool(
                source.source_page_url and source.download_url
            ),
            "frequency_present": bool(source.native_frequency),
            "geography_present": bool(source.geography_levels),
            "key_variables_present": bool(source.key_variables),
        }
        if source.source_id == HUD_USPS_SOURCE_ID:
            checks.update(_hud_usps_manifest_checks(config))
        issues = [name for name, passed in checks.items() if not passed]
        lifecycle_status = _source_lifecycle_status(source)
        if lifecycle_status == "Active":
            completeness_status = "PASS" if not issues else "FAIL"
        elif lifecycle_status == "Quarantined":
            completeness_status = "QUARANTINED"
        else:
            completeness_status = "INACTIVE"
        rows.append(
            {
                "source_id": source.source_id,
                "provider": source.provider,
                "active_status": lifecycle_status,
                "provenance_status": source.provenance_status,
                **checks,
                "completeness_status": completeness_status,
                "issue": "|".join(issues),
            }
        )
    result = pd.DataFrame(rows)
    output = config.root / "output" / "qa" / "raw_repository_completeness.csv"
    output.parent.mkdir(parents=True, exist_ok=True)
    result.to_csv(output, index=False)
    if database_path.exists():
        connection = duckdb.connect(str(database_path))
        connection.register("raw_completeness_frame", result)
        connection.execute(
            "CREATE OR REPLACE TABLE quality.raw_repository_completeness AS "
            "SELECT * FROM raw_completeness_frame"
        )
        connection.close()
    return result


def _coverage(config: ProjectConfig) -> pd.DataFrame:
    database_path = config.root / config.pipeline["database_path"]
    if not database_path.exists():
        return pd.DataFrame()
    connection = duckdb.connect(str(database_path), read_only=True)
    frame = connection.execute(
        """
        SELECT
            source_filename,
            min(observation_date) AS start_date,
            max(observation_date) AS end_date,
            count(*) AS row_count,
            string_agg(DISTINCT variable_id, '|' ORDER BY variable_id)
                AS variable_ids,
            string_agg(DISTINCT final_frequency, '|' ORDER BY final_frequency)
                AS analytical_frequency,
            string_agg(DISTINCT unit, '|' ORDER BY unit) AS units
        FROM analytics.fact_housing_observation
        GROUP BY source_filename
        """
    ).df()
    connection.close()
    return frame


def _latest_manifest_by_source(manifest: pd.DataFrame) -> pd.DataFrame:
    if manifest.empty:
        return manifest
    return (
        manifest.sort_values("_retrieved")
        .groupby("source_id", as_index=False)
        .tail(1)
        .set_index("source_id")
    )


def _variable_text(config: ProjectConfig, variable_ids: tuple[str, ...]) -> dict:
    variables = {
        str(item["variable_id"]): item for item in config.variables
    }
    selected = [
        variables[variable_id]
        for variable_id in variable_ids
        if variable_id in variables
    ]
    return {
        "seasonal_adjustment": "|".join(
            sorted(
                {
                    str(item.get("seasonal_adjustment", ""))
                    for item in selected
                    if item.get("seasonal_adjustment")
                }
            )
        ),
        "smoothing": "|".join(
            sorted(
                {
                    str(item.get("smoothing", ""))
                    for item in selected
                    if item.get("smoothing")
                }
            )
        ),
        "units": "|".join(
            sorted(
                {
                    str(item.get("unit", ""))
                    for item in selected
                    if item.get("unit")
                }
            )
        ),
        "missing_value_policy": "|".join(
            sorted(
                {
                    str(item.get("missing_value_policy", ""))
                    for item in selected
                    if item.get("missing_value_policy")
                }
            )
        ),
    }


def _geography_description(source_id: str, levels: tuple[str, ...]) -> tuple[str, str]:
    if source_id == HUD_USPS_SOURCE_ID:
        return (
            "58 California counties; Los Angeles County standardized 2020 "
            "tracts; modeled Los Angeles County postal-ZIP county-parts",
            "COUNTY_*|TRACT_*|ZIPCOUNTYPART_*",
        )
    if "zillow_zip" in source_id or source_id == "redfin_zip_monthly":
        return "Los Angeles County postal ZIPs", "POSTAL_ZIP"
    if "county" in levels:
        return "Los Angeles County", "06037"
    if "state" in levels:
        return "California", "06"
    if "metro" in levels:
        if "fhfa" in source_id:
            return "Los Angeles-Long Beach-Glendale Metropolitan Division", "31084"
        return "Los Angeles published metro/area", "31080|S49A"
    if "national" in levels:
        return "United States", "US"
    return "|".join(levels), ""


def _catalog_frame(config: ProjectConfig) -> pd.DataFrame:
    manifest = _manifest(config)
    latest = _latest_manifest_by_source(manifest)
    coverage = _coverage(config)
    rows: list[dict] = []
    for source in config.sources:
        raw = (
            latest.loc[source.source_id].to_dict()
            if not latest.empty and source.source_id in latest.index
            else {}
        )
        if source.source_id == HUD_USPS_SOURCE_ID:
            raw = _hud_usps_catalog_manifest_row(config)
        source_files = (
            set(
                manifest.loc[
                    manifest["source_id"].eq(source.source_id),
                    "source_filename",
                ].dropna()
            )
            if not manifest.empty
            else set()
        )
        source_coverage = (
            coverage.loc[coverage["source_filename"].isin(source_files)]
            if not coverage.empty
            else pd.DataFrame()
        )
        variable_ids = tuple(source.key_variables)
        variable_text = _variable_text(config, variable_ids)
        geography_name, geography_id = _geography_description(
            source.source_id, source.geography_levels
        )
        if source_coverage.empty:
            start_date = end_date = analytical_frequency = ""
            row_count = 0
            coverage_units = ""
        else:
            start_date = source_coverage["start_date"].min()
            end_date = source_coverage["end_date"].max()
            row_count = int(source_coverage["row_count"].sum())
            analytical_frequency = "|".join(
                sorted(
                    {
                        value
                        for text in source_coverage[
                            "analytical_frequency"
                        ].dropna()
                        for value in str(text).split("|")
                    }
                )
            )
            coverage_units = "|".join(
                sorted(
                    {
                        value
                        for text in source_coverage["units"].dropna()
                        for value in str(text).split("|")
                    }
                )
            )
        lifecycle_status = _source_lifecycle_status(source)
        command = (
            "BLOCKED: resolve the source's documented provenance or access "
            "issue before running its download workflow."
            if lifecycle_status == "Quarantined"
            else "python run_pipeline.py --stage download --source "
            f"{source.source_id} --force-download"
            if source.auto_download
            else "python run_pipeline.py --stage download --source "
            f"{source.source_id} --local-file PATH_TO_OFFICIAL_FILE"
        )
        if source.source_id == HUD_USPS_SOURCE_ID:
            provider_vintage = raw.get("provider_vintage", "")
        elif end_date:
            provider_vintage = f"data through {str(end_date)[:10]}"
        else:
            provider_vintage = raw.get("provider_vintage", "") or raw.get(
                "data_vintage", ""
            )
        catalog_file_row = raw
        rows.append(
            {
                "Source": source.dataset_name,
                "Frequency": source.native_frequency,
                "Geography": geography_name,
                "Key variables": "|".join(variable_ids),
                "source_id": source.source_id,
                "provider": source.provider,
                "product": source.dataset_name,
                "official_source_page": source.source_page_url,
                "exact_download_url_or_api": source.download_url,
                "access_method": source.access_method,
                "active_status": lifecycle_status,
                "provenance_status": source.provenance_status,
                "observed_constructed_proxy_or_scenario": (
                    "observed tract components|constructed county aggregates|"
                    "modeled postal ZIP county-part proxy"
                    if source.source_id == HUD_USPS_SOURCE_ID
                    else
                    "constructed proxy (quarantined)"
                    if lifecycle_status == "Quarantined"
                    else "constructed published-table sensitivity"
                    if source.source_id in SIPP_SOURCE_IDS
                    else "observed published estimates|constructed reported-case proxy"
                    if source.source_id == AHS_TABLE_CREATOR_SOURCE_ID
                    else "observed"
                ),
                "manual_download_required": not source.auto_download,
                "raw_file_format": (
                    "json|txt (75 immutable artifacts indexed by CSV manifest)"
                    if source.source_id == HUD_USPS_SOURCE_ID
                    else source.file_format
                ),
                "raw_file_path": catalog_file_row.get("relative_path", ""),
                "source_manifest_path": catalog_file_row.get(
                    "source_manifest_path", ""
                ),
                "original_filename": catalog_file_row.get("original_filename", "")
                or catalog_file_row.get("source_filename", ""),
                "retrieval_date": raw.get("retrieved_at", ""),
                "provider_vintage": provider_vintage,
                "checksum_sha256": catalog_file_row.get("checksum_sha256", "")
                or catalog_file_row.get("sha256", ""),
                "native_frequency": source.native_frequency,
                "analytical_frequency": (
                    analytical_frequency or source.native_frequency
                ),
                "geography_level": "|".join(source.geography_levels),
                "geography_name": geography_name,
                "geography_id": geography_id,
                "start_date": start_date,
                "end_date": end_date,
                "row_count": row_count,
                "variable_ids": "|".join(variable_ids),
                "seasonal_adjustment": variable_text["seasonal_adjustment"],
                "smoothing": variable_text["smoothing"],
                "units": coverage_units or variable_text["units"],
                "destination_clean_table": (
                    ""
                    if lifecycle_status == "Quarantined"
                    else "clean.hud_usps_ncwm_quarterly|"
                    "data/restricted/hud_usps/active/"
                    "long_observations_with_lineage.parquet"
                    if source.source_id == HUD_USPS_SOURCE_ID
                    else source.destination_clean_table
                    or f"clean.{source.source_id}"
                ),
                "destination_analytical_table": (
                    ""
                    if lifecycle_status == "Quarantined"
                    else "analytics.fact_housing_observation|"
                    + "|".join(HUD_USPS_RESTRICTED_TABLES[:3])
                    if source.source_id == HUD_USPS_SOURCE_ID
                    else source.destination_analytical_table
                ),
                "update_command": command,
                "missing_value_policy": variable_text[
                    "missing_value_policy"
                ],
                "revision_policy": source.revision_notes,
                "known_limitations": (
                    source.methodology_notes or source.revision_notes
                ),
                "license_or_use_notes": source.license_or_use_notes,
                "last_verified": config.verified_at,
                "notes": (
                    "Official 75-artifact manifest verified at "
                    f"{HUD_USPS_SOURCE_MANIFEST_PATH}; active licensed detail "
                    "is restricted. The former package file was deleted; its "
                    "historical record remains only in "
                    f"{HUD_USPS_HISTORICAL_QUARANTINE_MANIFEST_PATH} as a "
                    "historical, superseded record."
                    if source.source_id == HUD_USPS_SOURCE_ID
                    else "Intentional quarantine; excluded from active ingestion."
                    if lifecycle_status == "Quarantined"
                    else "Identifiers remain strings; raw vintages are immutable."
                ),
            }
        )
    existing_ids = {str(row["source_id"]) for row in rows}
    for source in _version_a_catalog_records(config):
        if source["source_id"] in existing_ids:
            continue
        variable_ids = tuple(source["variable_ids"])
        rows.append(
            {
                "Source": source["dataset_name"],
                "Frequency": source["native_frequency"],
                "Geography": source["geography"],
                "Key variables": "|".join(variable_ids),
                "source_id": source["source_id"],
                "provider": source["provider"],
                "product": source["dataset_name"],
                "official_source_page": source["source_page_url"],
                "exact_download_url_or_api": source["download_url"],
                "access_method": source["access_method"],
                "active_status": source["active_status"],
                "provenance_status": source["provenance_status"],
                "observed_constructed_proxy_or_scenario": source[
                    "observed_constructed_proxy_or_scenario"
                ],
                "manual_download_required": source[
                    "manual_download_required"
                ],
                "raw_file_format": source["file_format"],
                "raw_file_path": source["raw_file_path"],
                "original_filename": source["original_filename"],
                "retrieval_date": source["retrieved_at"],
                "provider_vintage": source["provider_vintage"],
                "checksum_sha256": source["checksum_sha256"],
                "native_frequency": source["native_frequency"],
                "analytical_frequency": source["native_frequency"],
                "geography_level": source["geography"],
                "geography_name": source["geography"],
                "geography_id": "",
                "start_date": source["start_date"],
                "end_date": source["end_date"],
                "row_count": source["row_count"],
                "variable_ids": "|".join(variable_ids),
                "seasonal_adjustment": "not applicable",
                "smoothing": "none",
                "units": "source-specific",
                "destination_clean_table": source[
                    "destination_clean_table"
                ],
                "destination_analytical_table": source[
                    "destination_analytical_table"
                ],
                "update_command": source["update_command"],
                "missing_value_policy": "preserve null",
                "revision_policy": "Archive every immutable source vintage.",
                "known_limitations": source["known_limitations"],
                "license_or_use_notes": source["license_or_use_notes"],
                "last_verified": config.verified_at,
                "notes": source["notes"],
            }
        )
        existing_ids.add(str(source["source_id"]))
    for unavailable in UNAVAILABLE_ROWS:
        row = {column: "" for column in CATALOG_COLUMNS}
        row.update(unavailable)
        row["last_verified"] = config.verified_at
        row["notes"] = "No placeholder observations are created."
        rows.append(row)
    return pd.DataFrame(rows).reindex(columns=CATALOG_COLUMNS)


def _gap_row(
    variable_name: str,
    market_segment: str,
    variable_role: str,
    requested_source: str,
    official_source_found: str,
    implementation_status: str,
    current_database_table: str,
    current_database_column: str,
    native_frequency: str,
    geography: str,
    available_start_date: str = "",
    available_end_date: str = "",
    limitation: str = "",
    recommended_action: str = "",
) -> dict[str, str]:
    return {
        "variable_name": variable_name,
        "market_segment": market_segment,
        "variable_role": variable_role,
        "requested_source": requested_source,
        "official_source_found": official_source_found,
        "implementation_status": implementation_status,
        "current_database_table": current_database_table,
        "current_database_column": current_database_column,
        "native_frequency": native_frequency,
        "geography": geography,
        "available_start_date": available_start_date,
        "available_end_date": available_end_date,
        "limitation": limitation,
        "recommended_action": recommended_action,
    }


def _refresh_gap_analysis(config: ProjectConfig) -> pd.DataFrame:
    rows = [
        _gap_row(
            "aggregate mortgage debt",
            "Homeowner",
            "Balance-sheet benchmark",
            "ACS PUMS and FHFA NMDB",
            "Yes, as a constructed benchmark",
            "Raw 20/25/30-year scenarios implemented; calibrated target blocked",
            "analytics.housing_finance_benchmark",
            "acs_implied_first_mortgage_balance_raw_pmms_20y_sensitivity_benchmark|acs_implied_first_mortgage_balance_raw_pmms_25y_sensitivity_benchmark|acs_implied_first_mortgage_balance_raw_pmms_30y_sensitivity_benchmark|acs_implied_first_mortgage_balance_raw_benchmark|acs_implied_first_mortgage_balance_cal_benchmark",
            "ACS 5-year plus exact-date parameter period",
            "California counties",
            limitation="All 174 explicit term-scenario debt values and the 58-row raw composite envelope are available. The calibrated value and normalized ratios remain null without an exact-date, owner-occupied first-lien California UPB control. The B25106 midpoint method remains in facts only as a labeled legacy scenario.",
            recommended_action="Use the explicit Version A raw scenario variables for sensitivity analysis; obtain or approve a compatible owner-occupied California UPB control before using calibrated fields.",
        ),
        _gap_row(
            "aggregate home equity",
            "Homeowner",
            "Balance-sheet benchmark",
            "ACS and FHFA NMDB",
            "Yes, as a constructed benchmark",
            "Raw 20/25/30-year scenarios implemented; calibrated target blocked",
            "analytics.housing_finance_benchmark",
            "acs_implied_owner_occupied_home_equity_mortgaged_pmms_20y_scenario_benchmark|acs_implied_owner_occupied_home_equity_mortgaged_pmms_25y_scenario_benchmark|acs_implied_owner_occupied_home_equity_mortgaged_pmms_30y_scenario_benchmark|acs_implied_owner_occupied_home_equity_all_pmms_20y_scenario_benchmark|acs_implied_owner_occupied_home_equity_all_pmms_25y_scenario_benchmark|acs_implied_owner_occupied_home_equity_all_pmms_30y_scenario_benchmark|acs_implied_owner_occupied_home_equity_mortgaged_benchmark|acs_implied_owner_occupied_home_equity_all_benchmark",
            "ACS 5-year plus exact-date parameter period",
            "California counties",
            limitation="All 348 matching PMMS term-scenario equity values are available. Preferred calibrated nominal and real equity remain null because they depend on a compatible calibrated first-mortgage balance.",
            recommended_action="Use the explicit scenario-equity variables for sensitivity analysis; retain published mortgage-free owner value in full and do not relabel scenarios as calibrated equity.",
        ),
        _gap_row("ZHVI", "Homeowner", "Outcome", "Zillow Research", "Yes", "Implemented", "analytics.county_monthly_panel", "zhvi", "monthly", "Los Angeles County", "2000-01-31", "2026-05-31", "Provider modeled index; histories can be revised.", "Keep as the local home-value index."),
        _gap_row("FHFA HPI", "Homeowner", "Benchmark", "FHFA", "Yes", "Implemented", "analytics.state_quarterly_panel|analytics.metro_quarterly_panel|analytics.county_annual_controls", "FHFA HPI columns", "quarterly and annual", "California; Los Angeles metro; Los Angeles County", "1975-12-31", "2026-03-31", "Products differ by geography, frequency, method, and seasonal adjustment.", "Retain at native frequency."),
        _gap_row("median sale price", "Homeowner", "Outcome", "Redfin Data Center", "Yes", "Implemented", "analytics.county_monthly_panel", "median_sale_price", "monthly", "Los Angeles County", "2012-01-31", "2026-05-31", "Transaction mix can change the median.", "Retain as a transaction-price indicator."),
        _gap_row("inventory", "Homeowner", "Supply", "Redfin Data Center", "Yes", "Implemented", "analytics.county_monthly_panel", "active_inventory", "monthly", "Los Angeles County", "2012-01-31", "2026-05-31", "For-sale inventory only; not rental inventory.", "Keep the label explicit."),
        _gap_row("active listings", "Homeowner", "Supply", "Redfin Data Center", "Yes", "Implemented", "analytics.zip_monthly_panel", "active_listings", "rolling 3 months", "Los Angeles County postal ZIPs", "2019-06-30", "2026-05-31", "ZIP product uses overlapping rolling-three-month windows.", "Keep separate from county inventory."),
        _gap_row("months of supply", "Homeowner", "Supply", "Redfin Data Center", "Yes", "Implemented", "analytics.county_monthly_panel", "months_of_supply", "monthly", "Los Angeles County", "2012-01-31", "2026-05-31", "Provider ratio; not additive.", "Retain provider value."),
        _gap_row("new listings", "Homeowner", "Leading indicator", "Redfin Data Center", "Yes", "Implemented", "analytics.county_monthly_panel|analytics.zip_monthly_panel", "new_listings", "monthly and rolling 3 months", "Los Angeles County and postal ZIPs", "2012-01-31", "2026-05-31", "ZIP observations are rolling windows.", "Use county monthly for time-series models."),
        _gap_row("homes sold", "Homeowner", "Demand", "Redfin Data Center", "Yes", "Implemented", "analytics.county_monthly_panel|analytics.zip_monthly_panel", "homes_sold", "monthly and rolling 3 months", "Los Angeles County and postal ZIPs", "2012-01-31", "2026-05-31", "Closed-sales flow; ZIP observations are rolling windows.", "Use with inventory ratios."),
        _gap_row("days on market", "Overall housing market", "Market time", "Redfin Data Center", "Yes", "Implemented", "analytics.county_monthly_panel|analytics.zip_monthly_panel", "days_on_market|median_days_on_market", "monthly and rolling 3 months", "Los Angeles County and postal ZIPs", "2012-01-31", "2026-05-31", "Median; not additive.", "Expose as market-time indicator."),
        _gap_row("share with price drops", "Overall housing market", "Leading indicator", "Redfin Data Center", "Yes", "Implemented", "analytics.county_monthly_panel", "price_drop_share", "monthly", "Los Angeles County", "2012-01-31", "2026-05-31", "Share, not count.", "Retain nulls and document coverage."),
        _gap_row("30-year mortgage rate", "Homeowner", "Control", "Freddie Mac PMMS", "Yes", "Implemented", "analytics.county_monthly_panel", "mortgage_rate_30y_monthly_average", "weekly aggregated to monthly", "United States", "1971-04-02", "2026-06-25", "National rate; not local.", "Keep U.S. geography explicit."),
        _gap_row("single-family permits", "Homeowner", "Supply", "Census BPS", "Yes", "Implemented", "analytics.county_monthly_panel", "permits_1_unit", "monthly", "Los Angeles County", "2000-01-31", "2026-05-31", "Permits are authorizations, not starts or completions.", "Retain as the single-family supply proxy."),
        _gap_row("monthly mortgage payment", "Homeowner", "Constructed indicator", "Redfin and Freddie Mac", "Yes", "Constructed variable", "analytics.county_monthly_panel", "mortgage_payment_10pct_down|mortgage_payment_20pct_down", "monthly", "Los Angeles County with U.S. rate", "2012-01-31", "2026-05-31", "Principal-and-interest scenarios exclude taxes, insurance, HOA, maintenance, and mortgage insurance.", "Keep scenario assumptions visible."),
        _gap_row("ZORI", "Renter", "Outcome", "Zillow Research", "Yes", "Implemented", "analytics.county_monthly_panel", "zori", "monthly", "Los Angeles County", "2015-01-31", "2026-05-31", "Smoothed seasonally adjusted rent index; not inventory.", "Use as primary rent index."),
        _gap_row("rent growth", "Renter", "Constructed indicator", "Zillow ZORI", "Yes", "Constructed variable", "analytics.county_monthly_panel", "zori_growth_mom|zori_growth_yoy", "monthly", "Los Angeles County", "2015-02-28", "2026-05-31", "Depends on the active SA ZORI vintage.", "Rebuild when ZORI refreshes."),
        _gap_row(
            "rental inventory",
            "Renter",
            "Supply",
            "No approved true-listing source named; ACS and HUD/USPS are separate proxies",
            "No; governing methodology/source decision required",
            RENTAL_INVENTORY_STATUS,
            "",
            "rental_listing_inventory_monthly",
            "monthly target; unavailable",
            "Los Angeles County or provider ZIP target",
            "",
            "",
            "The governing methodology does not name an approved, integrable "
            "historical active-rental-listing source. ACS annual availability and "
            "HUD/USPS quarterly all-residential vacancy remain separately labeled "
            "proxies and do not satisfy this target.",
            "Revise the methodology to name an approved qualifying source, or "
            "leave rental_listing_inventory_monthly missing. No credential or API "
            "activation resolves the current blocker.",
        ),
        _gap_row("rental vacancy", "Renter", "Supply", "Census HVS", "Yes", "Implemented", "analytics.state_quarterly_panel|analytics.metro_quarterly_panel", "hvs_rental_vacancy_rate_ca|hvs_rental_vacancy_rate_la_metro", "quarterly", "California and published metro geographies", "2005-03-31", "2026-03-31", "Los Angeles County is not an HVS published geography; sampling uncertainty is retained.", "Retain native geography/frequency."),
        _gap_row("multifamily permits", "Renter", "Supply", "Census BPS", "Yes", "Implemented", "analytics.county_monthly_panel", "permits_5_plus_units", "monthly", "Los Angeles County", "2000-01-31", "2026-05-31", "Permits are authorizations, not completed or occupied rental units.", "Retain as multifamily supply proxy."),
        _gap_row("price-to-rent", "Overall housing market", "Constructed indicator", "Zillow ZHVI and ZORI", "Yes", "Constructed variable", "analytics.county_monthly_panel", "price_to_rent_ratio", "monthly", "Los Angeles County", "2015-01-31", "2026-05-31", "Index-based valuation ratio.", "Keep explicit formula."),
        _gap_row("inventory-to-sales", "Overall housing market", "Constructed indicator", "Redfin Data Center", "Yes", "Constructed variable", "analytics.county_monthly_panel", "inventory_to_sales_ratio|sales_to_inventory_ratio", "monthly", "Los Angeles County", "2012-01-31", "2026-05-31", "Uses for-sale inventory and sales, not total housing stock.", "Keep both orientations clearly named."),
        _gap_row("market tightness", "Overall housing market", "Constructed indicator", "Redfin Data Center", "Yes", "Constructed variable", "analytics.county_monthly_panel", "market_tightness_index|market_tightness_components", "monthly", "Los Angeles County", "2012-01-31", "2026-05-31", "Composite, not official provider statistic.", "Retain component-count metadata."),
        _gap_row("affordability metrics", "Overall housing market", "Affordability", "Zillow, Redfin, Freddie Mac", "Yes", "Implemented", "analytics.county_monthly_panel", "price_to_rent_ratio|mortgage_payment_to_rent_10pct_down|mortgage_payment_to_rent_20pct_down", "monthly", "Los Angeles County", "2015-01-31", "2026-05-31", "Income-based ratios are deferred at this active stage.", "Do not expose ACS-derived affordability columns until requested."),
        _gap_row("LAUS unemployment rate", "Macro", "Control", "BLS LAUS", "Yes", "Implemented", "analytics.county_monthly_panel", "unemployment_rate", "monthly", "Los Angeles County", "1990-01-31", "2026-05-31", "County series is not seasonally adjusted and may be revised.", "Archive API responses."),
        _gap_row("LAUS labor force", "Macro", "Control", "BLS LAUS", "Yes", "Implemented", "analytics.county_monthly_panel", "labor_force", "monthly", "Los Angeles County", "1990-01-31", "2026-05-31", "County series is not seasonally adjusted and may be revised.", "Archive API responses."),
        _gap_row("CES payroll employment", "Macro", "Control", "BLS CES", "Yes", "Implemented", "analytics.metro_monthly_controls", "ces_total_nonfarm_employment_sa", "monthly", "Los Angeles metro", "2005-01-31", "2026-05-31", "Metro establishment series; not county LAUS employment.", "Retain geography label."),
        _gap_row("personal income", "Macro", "Control", "BEA", "Yes", "Implemented", "analytics.state_quarterly_panel|analytics.county_annual_controls", "bea_ca_personal_income_q_sa|bea_la_county_personal_income_annual", "quarterly and annual", "California and Los Angeles County", "1948-03-31", "2026-03-31", "State quarterly and county annual products differ.", "Retain native frequency."),
        _gap_row("CPI / inflation", "Macro", "Control", "BLS CPI", "Yes", "Implemented", "analytics.metro_monthly_controls", "cpi_u_all_items_la_nsa|cpi_inflation_yoy", "monthly with historical gaps", "Los Angeles published area", "1914-12-31", "2026-05-31", "Local CPI is not seasonally adjusted and has historical publication gaps.", "Preserve gaps."),
        _gap_row("price-to-income", "Homeowner", "Affordability", "Census ACS B25077/B19013", "Yes", "Partially implemented", "analytics.county_annual_controls", "county_annual_price_to_income_acs", "annual ACS five-year", "Los Angeles County", "2010-12-31", "2024-12-31", "Fifteen annual ratios are active; overlapping five-year estimates are not quarterly observations.", "Retain as annual or approve a separately labeled mixed-frequency method."),
    ]
    gap = pd.DataFrame(rows)
    path = config.root / "metadata" / "data_gap_analysis.csv"
    gap.to_csv(path, index=False)
    return gap


def _format_workbook(path: Path, build_version: str) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    status_colors = {
        "Active": "C6E0B4",
        "Implemented": "C6E0B4",
        "Constructed variable": "D9EAF7",
        "Partially implemented": "FFF2CC",
        RENTAL_INVENTORY_STATUS: "FFF2CC",
        "Missing — no verified public source": "F4CCCC",
        "Deferred": "E2F0D9",
        "Proxy only": "E4DFEC",
        "PASS": "C6E0B4",
        "FAIL": "F4CCCC",
    }
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        headers = {
            str(cell.value): cell.column
            for cell in sheet[1]
            if cell.value is not None
        }
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, str) and cell.value in status_colors:
                    cell.fill = PatternFill(
                        "solid", fgColor=status_colors[cell.value]
                    )
            for header, column in headers.items():
                if any(
                    token in header.lower()
                    for token in ("id", "fips", "zip", "checksum")
                ):
                    row[column - 1].number_format = "@"
        for column_index in range(1, sheet.max_column + 1):
            maximum = max(
                len(str(sheet.cell(row, column_index).value or ""))
                for row in range(1, min(sheet.max_row, 250) + 1)
            )
            sheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(maximum + 2, 12), 42
            )
    workbook.properties.creator = "LA housing database pipeline"
    workbook.properties.title = "LA Housing Database Data Catalog"
    workbook.properties.description = (
        "Generated from configuration, raw manifest, DuckDB metadata, coverage, "
        f"and gap analysis. Database build version: {build_version}."
    )
    workbook.properties.keywords = f"database_build_version={build_version}"
    workbook.save(path)


def build_data_catalog(config: ProjectConfig) -> dict[str, int]:
    generated_at = datetime.now(timezone.utc).isoformat()
    database_path = config.root / config.pipeline["database_path"]
    if database_path.exists():
        connection = duckdb.connect(str(database_path), read_only=True)
        fact_rows, latest_vintage = connection.execute(
            """
            SELECT count(*), max(source_vintage)
            FROM analytics.fact_housing_observation
            """
        ).fetchone()
        connection.close()
        build_version = f"{latest_vintage}-{fact_rows}"
    else:
        build_version = "database-not-built"
    completeness = build_raw_repository_completeness(config)
    catalog = _catalog_frame(config)
    gap = _refresh_gap_analysis(config)
    catalog["notes"] = catalog["notes"].map(
        lambda value: (
            f"{value} Catalog generated {generated_at}; "
            f"database build version {build_version}."
        )
    )
    csv_path = config.root / "metadata" / "data_catalog.csv"
    xlsx_path = config.root / "metadata" / "data_catalog.xlsx"
    gap_xlsx = config.root / "metadata" / "data_gap_analysis.xlsx"
    _write_catalog_preserving_monthly_finance_products(catalog, csv_path)
    register_monthly_finance_contracts(root=config.root)
    catalog = pd.read_csv(csv_path)
    manifest = _manifest(config).drop(columns="_retrieved", errors="ignore")
    variables = pd.read_csv(config.root / "metadata" / "data_dictionary.csv")
    matrix = catalog[
        [
            "source_id",
            "active_status",
            "native_frequency",
            "analytical_frequency",
            "geography_level",
            "geography_name",
        ]
    ].copy()
    coverage = catalog[
        [
            "source_id",
            "active_status",
            "start_date",
            "end_date",
            "row_count",
            "variable_ids",
        ]
    ].copy()
    unavailable = catalog.loc[
        ~catalog["active_status"].eq("Active")
    ].copy()
    commands = catalog.loc[
        catalog["active_status"].eq("Active"),
        ["source_id", "provider", "update_command"],
    ].copy()
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        catalog.to_excel(writer, sheet_name="Data Catalog", index=False)
        variables.to_excel(
            writer, sheet_name="Variable Dictionary", index=False
        )
        manifest.to_excel(
            writer, sheet_name="Raw File Inventory", index=False
        )
        matrix.to_excel(
            writer, sheet_name="Geography-Frequency Matrix", index=False
        )
        coverage.to_excel(
            writer, sheet_name="Coverage Summary", index=False
        )
        unavailable.to_excel(
            writer, sheet_name="Unavailable and Deferred", index=False
        )
        commands.to_excel(
            writer, sheet_name="Update Commands", index=False
        )
        gap.to_excel(writer, sheet_name="Gap Analysis", index=False)
    _format_workbook(xlsx_path, build_version)
    _write_gap_workbook_atomic(gap_xlsx, gap, build_version)
    if database_path.exists():
        connection = duckdb.connect(str(database_path))
        connection.register("data_catalog_frame", catalog)
        connection.register("gap_analysis_frame", gap)
        connection.execute(
            "CREATE OR REPLACE TABLE metadata.data_catalog AS "
            "SELECT * FROM data_catalog_frame"
        )
        connection.execute(
            "CREATE OR REPLACE TABLE metadata.data_gap_analysis AS "
            "SELECT * FROM gap_analysis_frame"
        )
        connection.close()
    active_failures = completeness.loc[
        completeness["active_status"].eq("Active")
        & completeness["completeness_status"].ne("PASS")
    ]
    failures = len(active_failures)
    if failures:
        failed = ", ".join(active_failures["source_id"])
        raise RuntimeError(
            f"Raw repository completeness failed for {failures} sources: {failed}"
        )
    return {
        "catalog_rows": len(catalog),
        "variable_rows": len(variables),
        "raw_inventory_rows": len(manifest),
        "gap_rows": len(gap),
    }
