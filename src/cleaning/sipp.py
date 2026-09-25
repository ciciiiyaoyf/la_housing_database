"""Clean published 2025 SIPP California state debt and wealth tables.

These published-table products provide deliberately broad state sensitivities.
They are not substitutes for the repository's exact, county-level FHFA/NMDB
calibration targets.
"""

from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.cleaning.common import finalize, metadata_fields


DEBT_SOURCE_ID = "census_sipp_2025_state_debt_2024"
WEALTH_SOURCE_ID = "census_sipp_2025_state_wealth_2024"
SOURCE_IDS = frozenset({DEBT_SOURCE_ID, WEALTH_SOURCE_ID})

REFERENCE_YEAR = 2024
EXPECTED_RELEASE_DATE = "2026-07-15"
RELEASE_PATTERN = re.compile(
    r"^Internet Release Date:\s*(\d{2})/(\d{2})/(\d{4})$"
)

DEBT_VARIABLE_ID = "sipp_ca_aggregate_primary_residence_debt_sensitivity"
EQUITY_VARIABLE_ID = "sipp_ca_aggregate_positive_home_equity_sensitivity"

DEBT_TRANSFORMATION_ID = "sipp_ca_primary_residence_debt_sensitivity_v1"
EQUITY_TRANSFORMATION_ID = "sipp_ca_positive_home_equity_sensitivity_v1"

DEBT_TABLE_2_TITLE = (
    "Table 2. Percent Holding Debt for Households, by Type of Debt and "
    "States: 2024"
)
DEBT_TABLE_4_TITLE = (
    "Table 4. Mean Value of Debt for Households, by Type of Debt and "
    "States: 2024"
)
WEALTH_TABLE_2_TITLE = (
    "Table 2. Percent Holding Assets for Households, by Type of Asset "
    "Owned and States: 2024"
)
WEALTH_TABLE_3_TITLE = (
    "Table 3. Ownership Rates for Assets at Financial Institutions, by "
    "States: 2024"
)
WEALTH_TABLE_5_TITLE = (
    "Table 5. Mean Value of Assets for Households, by Type of Asset Owned "
    "and States: 2024"
)


def _workbook(path: Path):
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"SIPP workbook is not a readable XLSX: {path.name}") from exc


def published_release_date(path: Path) -> str:
    """Return the unique provider release date embedded in a SIPP workbook."""

    workbook = _workbook(path)
    release_dates: set[str] = set()
    try:
        for worksheet in workbook.worksheets:
            for (value,) in worksheet.iter_rows(
                min_col=1,
                max_col=1,
                values_only=True,
            ):
                if not isinstance(value, str):
                    continue
                match = RELEASE_PATTERN.fullmatch(value.strip())
                if match:
                    month, day, year = match.groups()
                    release_dates.add(f"{year}-{month}-{day}")
    finally:
        workbook.close()
    if release_dates != {EXPECTED_RELEASE_DATE}:
        raise ValueError(
            "SIPP workbook release-date contract changed: "
            f"expected={[EXPECTED_RELEASE_DATE]}, actual={sorted(release_dates)}"
        )
    return EXPECTED_RELEASE_DATE


def _require_sheet(workbook, name: str):
    if name not in workbook.sheetnames:
        raise ValueError(f"SIPP workbook is missing required worksheet {name!r}")
    return workbook[name]


def _require_cell(worksheet, coordinate: str, expected: object) -> None:
    actual = worksheet[coordinate].value
    if actual != expected:
        raise ValueError(
            f"SIPP schema drift in {worksheet.title}!{coordinate}: "
            f"expected={expected!r}, actual={actual!r}"
        )


def _california_row(worksheet) -> int:
    rows = [
        row_number
        for row_number, (value,) in enumerate(
            worksheet.iter_rows(min_col=1, max_col=1, values_only=True),
            start=1,
        )
        if value == "California"
    ]
    if len(rows) != 1:
        raise ValueError(
            f"SIPP {worksheet.title} must contain exactly one California row; "
            f"found={rows}"
        )
    return rows[0]


def _positive_decimal(worksheet, coordinate: str, label: str) -> Decimal:
    value = worksheet[coordinate].value
    if isinstance(value, bool):
        raise ValueError(f"SIPP {label} is not numeric: {value!r}")
    try:
        number = Decimal(str(value))
    except Exception as exc:
        raise ValueError(f"SIPP {label} is not numeric: {value!r}") from exc
    if not number.is_finite() or number <= 0:
        raise ValueError(f"SIPP {label} must be positive and finite: {value!r}")
    return number


def _holding_share(worksheet, coordinate: str, label: str) -> Decimal:
    share = _positive_decimal(worksheet, coordinate, label)
    if share > 100:
        raise ValueError(f"SIPP {label} must be at most 100 percent: {share}")
    return share


def _validate_source_note(worksheet) -> None:
    source_lines = [
        value.strip()
        for (value,) in worksheet.iter_rows(
            min_col=1,
            max_col=1,
            values_only=True,
        )
        if isinstance(value, str) and value.strip().startswith("Source:")
    ]
    if len(source_lines) != 1 or "Survey Year 2025" not in source_lines[0]:
        raise ValueError(
            f"SIPP {worksheet.title} source note does not identify Survey Year 2025"
        )


def _aggregate(
    households_thousands: Decimal,
    holding_share_percent: Decimal,
    conditional_mean: Decimal,
) -> Decimal:
    return (
        households_thousands
        * Decimal("1000")
        * holding_share_percent
        / Decimal("100")
        * conditional_mean
    )


def _observation(
    *,
    variable_id: str,
    value: Decimal,
    transformation_id: str,
    quality_flag: str,
    record: dict[str, str],
) -> pd.DataFrame:
    observation_date = pd.Timestamp(f"{REFERENCE_YEAR}-12-31")
    frame = pd.DataFrame(
        {
            "observation_date": [observation_date],
            "source_observation_date": [observation_date],
            "period_start": [pd.Timestamp(f"{REFERENCE_YEAR}-01-01")],
            "period_end": [observation_date],
            "geo_id": ["STATE_06"],
            "state_fips": ["06"],
            "county_fips": [None],
            "variable_id": [variable_id],
            "value": [float(value)],
            "reported_value": [None],
            "unit": ["current_dollars"],
            "provider": ["U.S. Census Bureau"],
            "native_frequency": ["annual_published_table_product"],
            "final_frequency": ["annual"],
            "transformation_id": [transformation_id],
            "is_imputed": [False],
            "is_interpolated": [False],
            "is_seasonally_adjusted": [False],
            "quality_flag": [quality_flag],
        }
    )
    for key, metadata_value in metadata_fields(record).items():
        frame[key] = metadata_value
    return frame


def _clean_debt(path: Path, record: dict[str, str]) -> pd.DataFrame:
    workbook = _workbook(path)
    try:
        table_2 = _require_sheet(workbook, "Table 2")
        table_4 = _require_sheet(workbook, "Table 4")
        _require_cell(table_2, "A2", DEBT_TABLE_2_TITLE)
        _require_cell(table_2, "B3", "Number of Households (thousands)")
        _require_cell(table_2, "E4", "Home Debt")
        _require_cell(table_4, "A2", DEBT_TABLE_4_TITLE)
        _require_cell(table_4, "D4", "Home Debt")
        _validate_source_note(table_2)
        _validate_source_note(table_4)
        row_2 = _california_row(table_2)
        row_4 = _california_row(table_4)
        households = _positive_decimal(
            table_2, f"B{row_2}", "California households (thousands)"
        )
        share = _holding_share(
            table_2, f"E{row_2}", "California home-debt holding share"
        )
        conditional_mean = _positive_decimal(
            table_4, f"D{row_4}", "California conditional mean home debt"
        )
    finally:
        workbook.close()

    aggregate = _aggregate(households, share, conditional_mean)
    quality_flag = (
        "published_table_product;statewide_sipp_sensitivity_not_exact_nmdb_"
        "calibration;home_debt_broader_than_owner_occupied_first_lien;"
        "households_rounded_nearest_thousand;holding_share_published_rounded;"
        "conditional_mean_published_rounded;conditional_mean_total_debt_below_"
        "99th_percentile;aggregate_constructed_from_rounded_inputs;"
        "no_covariance_based_standard_error"
    )
    return _observation(
        variable_id=DEBT_VARIABLE_ID,
        value=aggregate,
        transformation_id=DEBT_TRANSFORMATION_ID,
        quality_flag=quality_flag,
        record=record,
    )


def _clean_wealth(path: Path, record: dict[str, str]) -> pd.DataFrame:
    workbook = _workbook(path)
    try:
        table_2 = _require_sheet(workbook, "Table 2")
        table_3 = _require_sheet(workbook, "Table 3")
        table_5 = _require_sheet(workbook, "Table 5")
        _require_cell(table_2, "A2", WEALTH_TABLE_2_TITLE)
        _require_cell(table_2, "E3", "Equity in Own Home")
        _require_cell(table_3, "A2", WEALTH_TABLE_3_TITLE)
        _require_cell(table_3, "B3", "Number of Households (thousands)")
        _require_cell(table_5, "A2", WEALTH_TABLE_5_TITLE)
        _require_cell(table_5, "G3", "Equity in Own Home")
        _validate_source_note(table_2)
        _validate_source_note(table_3)
        _validate_source_note(table_5)
        row_2 = _california_row(table_2)
        row_3 = _california_row(table_3)
        row_5 = _california_row(table_5)
        households = _positive_decimal(
            table_3, f"B{row_3}", "California households (thousands)"
        )
        share = _holding_share(
            table_2,
            f"E{row_2}",
            "California positive-home-equity holding share",
        )
        conditional_mean = _positive_decimal(
            table_5, f"G{row_5}", "California conditional mean home equity"
        )
    finally:
        workbook.close()

    aggregate = _aggregate(households, share, conditional_mean)
    quality_flag = (
        "published_table_product;statewide_sipp_sensitivity_not_exact_"
        "matched_universe_equity;positive_equity_holders_only;households_"
        "rounded_nearest_thousand;holding_share_published_rounded;conditional_"
        "mean_published_rounded;conditional_mean_net_worth_below_99th_"
        "percentile;aggregate_constructed_from_rounded_inputs;"
        "no_covariance_based_standard_error"
    )
    return _observation(
        variable_id=EQUITY_VARIABLE_ID,
        value=aggregate,
        transformation_id=EQUITY_TRANSFORMATION_ID,
        quality_flag=quality_flag,
        record=record,
    )


def clean_sipp_state_sensitivity(
    path: Path,
    record: dict[str, str],
    output_path: Path,
) -> pd.DataFrame:
    """Create one California 2024 sensitivity from one official SIPP file."""

    source_id = str(record.get("source_id") or "")
    if source_id not in SOURCE_IDS:
        raise ValueError(f"Unexpected SIPP source_id={source_id!r}")
    release_date = published_release_date(path)
    recorded_release = str(record.get("provider_last_updated") or "").strip()
    if recorded_release and recorded_release != release_date:
        raise ValueError(
            "SIPP manifest and workbook release dates disagree: "
            f"manifest={recorded_release}, workbook={release_date}"
        )
    record = {**record, "provider_last_updated": release_date}
    frame = (
        _clean_debt(path, record)
        if source_id == DEBT_SOURCE_ID
        else _clean_wealth(path, record)
    )
    return finalize(frame, output_path)
