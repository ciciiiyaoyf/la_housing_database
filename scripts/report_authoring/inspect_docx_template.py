from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


path = Path(sys.argv[1]).resolve()
document = Document(path)


def inches(value):
    return None if value is None else round(value.inches, 3)


paragraphs = []
for index, paragraph in enumerate(document.paragraphs):
    fmt = paragraph.paragraph_format
    runs = []
    for run in paragraph.runs:
        runs.append(
            {
                "text": run.text,
                "font": run.font.name,
                "size_pt": None if run.font.size is None else round(run.font.size.pt, 2),
                "bold": run.bold,
                "italic": run.italic,
                "underline": bool(run.underline) if run.underline is not None else None,
                "color": None if run.font.color.rgb is None else str(run.font.color.rgb),
            }
        )
    paragraphs.append(
        {
            "index": index,
            "style": paragraph.style.name,
            "text": paragraph.text,
            "alignment": None if paragraph.alignment is None else int(paragraph.alignment),
            "left_indent_in": inches(fmt.left_indent),
            "right_indent_in": inches(fmt.right_indent),
            "first_line_indent_in": inches(fmt.first_line_indent),
            "space_before_pt": None if fmt.space_before is None else round(fmt.space_before.pt, 2),
            "space_after_pt": None if fmt.space_after is None else round(fmt.space_after.pt, 2),
            "line_spacing": str(fmt.line_spacing),
            "runs": runs,
        }
    )

sections = []
for index, section in enumerate(document.sections):
    sections.append(
        {
            "index": index,
            "page_width_in": inches(section.page_width),
            "page_height_in": inches(section.page_height),
            "top_margin_in": inches(section.top_margin),
            "bottom_margin_in": inches(section.bottom_margin),
            "left_margin_in": inches(section.left_margin),
            "right_margin_in": inches(section.right_margin),
            "header_distance_in": inches(section.header_distance),
            "footer_distance_in": inches(section.footer_distance),
            "different_first_page": bool(section.different_first_page_header_footer),
            "header": [p.text for p in section.header.paragraphs],
            "footer": [p.text for p in section.footer.paragraphs],
        }
    )

with zipfile.ZipFile(path) as archive:
    package = []
    for info in sorted(archive.infolist(), key=lambda item: item.filename):
        data = archive.read(info.filename)
        package.append(
            {
                "path": info.filename,
                "size": len(data),
                "sha256": hashlib.sha256(data).hexdigest(),
            }
        )

result = {
    "path": str(path),
    "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    "paragraph_count": len(document.paragraphs),
    "table_count": len(document.tables),
    "inline_shape_count": len(document.inline_shapes),
    "sections": sections,
    "paragraphs": paragraphs,
    "package": package,
}

print(json.dumps(result, indent=2, ensure_ascii=False))
