from __future__ import annotations

import hashlib
import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REFERENCE_HASH = "d3006f9074e46823d103fe283be19901f55f75c19811558455d7a16c61797355"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def restore_package_part(reference: Path, target: Path, part_name: str) -> None:
    replacement_path = target.with_suffix(".package-fix.docx")
    with zipfile.ZipFile(reference) as reference_zip, zipfile.ZipFile(target) as target_zip:
        replacement = reference_zip.read(part_name)
        with zipfile.ZipFile(replacement_path, "w") as output_zip:
            for info in target_zip.infolist():
                data = replacement if info.filename == part_name else target_zip.read(info.filename)
                output_zip.writestr(info, data)
    replacement_path.replace(target)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_cell_margins(cell, top=95, start=110, bottom=95, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color="D9D9D9", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, width_in: float) -> None:
    width = Inches(width_in)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.emu / 635)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=11.0, bold=None, italic=None, underline=None, color="000000") -> None:
    run.font.name = "Garamond"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Garamond")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline
    run.font.color.rgb = RGBColor.from_string(color)


def style_body_paragraph(paragraph, *, indent=True, space_after=4, line=1.18) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line
    fmt.first_line_indent = Inches(0.35) if indent else None
    for run in paragraph.runs:
        set_run_font(run, 11.0)


def add_body(document, text: str, *, indent=True, bold_lead: str | None = None):
    paragraph = document.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        run = paragraph.add_run(bold_lead)
        set_run_font(run, 11.0, bold=True)
        run = paragraph.add_run(text[len(bold_lead):])
        set_run_font(run, 11.0)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, 11.0)
    style_body_paragraph(paragraph, indent=indent)
    return paragraph


def add_heading(document, text: str, level=1):
    paragraph = document.add_paragraph(style="Heading 1" if level == 1 else "Heading 2")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(10 if level == 1 else 7)
    fmt.space_after = Pt(5)
    fmt.keep_with_next = True
    fmt.keep_together = True
    run = paragraph.add_run(text)
    set_run_font(run, 12.5 if level == 1 else 11.5, bold=True, underline=(level == 1))
    return paragraph


def add_bullet(document, text: str, *, bold_lead: str | None = None):
    paragraph = document.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(0.35)
    fmt.first_line_indent = Inches(-0.18)
    fmt.space_after = Pt(3)
    fmt.line_spacing = 1.12
    bullet = paragraph.add_run("•\t")
    set_run_font(bullet, 11.0)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, 11.0, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest, 11.0)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, 11.0)
    return paragraph


def add_numbered(document, number: int, text: str):
    paragraph = document.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(0.48)
    fmt.first_line_indent = Inches(-0.26)
    fmt.space_after = Pt(3)
    fmt.line_spacing = 1.12
    run = paragraph.add_run(f"{number}.\t{text}")
    set_run_font(run, 11.0)
    return paragraph


def add_table(document, headers, rows, widths, *, font_size=9.5, center_columns=()):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, (cell, value, width) in enumerate(zip(header.cells, headers, widths)):
        set_cell_width(cell, width)
        set_cell_margins(cell, top=105, bottom=105)
        set_cell_shading(cell, "E7E6E6")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        run = paragraph.add_run(str(value))
        set_run_font(run, font_size, bold=True)

    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        for col_index, (cell, value, width) in enumerate(zip(row.cells, row_values, widths)):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, "F7F7F7")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if col_index in center_columns else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(str(value))
            set_run_font(run, font_size)

    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    return table


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4E79")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Garamond")
    fonts.set(qn("w:hAnsi"), "Garamond")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    run_properties.extend([fonts, color, underline, size])
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_source(document, name: str, url: str) -> None:
    paragraph = document.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.left_indent = Inches(0.35)
    fmt.first_line_indent = Inches(-0.18)
    fmt.space_after = Pt(1)
    fmt.line_spacing = 1.0
    bullet = paragraph.add_run("•\t")
    set_run_font(bullet, 9.8)
    label = paragraph.add_run(f"{name}: ")
    set_run_font(label, 9.8, bold=True)
    add_hyperlink(paragraph, url, url)


def add_metadata(document, label: str, value: str):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(1)
    fmt.line_spacing = 1.0
    fmt.tab_stops.add_tab_stop(Inches(1.05))
    label_run = paragraph.add_run(label)
    set_run_font(label_run, 11.5, bold=True)
    value_run = paragraph.add_run(f"\t{value}")
    set_run_font(value_run, 11.5)
    return paragraph


def add_horizontal_rule(document):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    return paragraph


def build(reference: Path, output: Path) -> None:
    if sha256(reference) != REFERENCE_HASH:
        raise RuntimeError("The retained template no longer matches the distilled reference hash.")

    output.parent.mkdir(parents=True, exist_ok=True)
    working = output.with_suffix(".working.docx")
    shutil.copy2(reference, working)
    document = Document(working)

    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.clear()
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header_paragraph.add_run("LOS ANGELES HOUSING DATABASE")
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string("B7B7B7")

    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(20)
    title.paragraph_format.space_after = Pt(14)
    title.paragraph_format.keep_with_next = True
    title_run = title.add_run("MSA Variables and More Granular Alternatives")
    set_run_font(title_run, 22.0, bold=True, underline=False)

    add_metadata(document, "Prepared for", "Project mentor")
    add_metadata(document, "Prepared by", "Los Angeles Housing Database Project")
    add_metadata(document, "Inventory date", "August 24, 2026")
    add_metadata(document, "Source review", "September 4, 2026")
    add_metadata(document, "Subject", "Metropolitan geography and recommended local alternatives")
    add_horizontal_rule(document)

    add_heading(document, "Executive Summary")
    add_body(
        document,
        "Main conclusion. The database snapshot contains 11 variable IDs with metropolitan-area involvement. Eight measure areas broader than Los Angeles County: three CES employment variables, two CPI variables, and three Housing Vacancy Survey variables. Three FHFA variables use Metropolitan Division 31084, a provider-defined geography whose footprint consists only of Los Angeles County. Those FHFA observations are county-equivalent in area, but the official metropolitan-division label and method should remain intact.",
        indent=False,
        bold_lead="Main conclusion.",
    )
    add_body(
        document,
        "The highest-value correction is to add the official BLS seasonally adjusted total nonfarm payroll series for the Los Angeles Long Beach Glendale Metropolitan Division, series SMS06310840000000001, together with separate month-over-month and year-over-year growth IDs. The existing two-county CES series should remain as a regional comparator. FHFA annual ZIP5 and tract indexes, BLS QCEW county measures, and ACS tract vacancy estimates can then extend local coverage without changing native geography or frequency.",
        indent=False,
    )
    add_body(
        document,
        "This is a review document, not an ingestion result. The counts and date ranges below describe the August 24, 2026 database inventory. Public-source availability was reviewed on September 4, 2026. No source file was downloaded, no schema was changed, and no new observation was represented as implemented during this review.",
        indent=False,
    )

    add_heading(document, "Inventory of Metropolitan Variables")
    add_body(
        document,
        "The following inventory reports variable IDs rather than conceptual indicators. A level series and its derived growth series are therefore counted separately because each has its own ID, coverage, and transformation contract.",
        indent=False,
    )
    inventory_rows = [
        ("CES", "ces_total_nonfarm_employment_sa", "LA Long Beach Anaheim MSA", "Monthly", "257 observations\nJan 2005 to May 2026"),
        ("CES", "ces_total_nonfarm_employment_growth_mom", "LA Long Beach Anaheim MSA", "Monthly", "256 observations\nFeb 2005 to May 2026"),
        ("CES", "ces_total_nonfarm_employment_growth_yoy", "LA Long Beach Anaheim MSA", "Monthly", "245 observations\nJan 2006 to May 2026"),
        ("CPI", "cpi_u_all_items_la_nsa", "LA Long Beach Anaheim CPI area", "Monthly schedule", "1,094 observations\nDec 1914 to May 2026 with gaps"),
        ("CPI", "cpi_inflation_yoy", "LA Long Beach Anaheim CPI area", "Monthly schedule", "1,082 observations\nMar 1922 to May 2026 with gaps"),
        ("HVS", "hvs_rental_vacancy_rate_la_metro", "Published Los Angeles metro", "Quarterly", "85 observations\n2005 Q1 to 2026 Q1"),
        ("HVS", "hvs_rental_vacancy_rate_la_metro_moe", "Published Los Angeles metro", "Quarterly", "85 observations\n2005 Q1 to 2026 Q1"),
        ("HVS", "hvs_rental_vacancy_change_yoy_pp", "State or published metro", "Quarterly", "81 LA metro observations\n2006 Q1 to 2026 Q1"),
        ("FHFA", "fhfa_hpi_la_metropolitan_division_purchase_only_nsa", "LA Long Beach Glendale MD 31084", "Quarterly", "141 observations\n1991 Q1 to 2026 Q1"),
        ("FHFA", "fhfa_hpi_la_metropolitan_division_purchase_only_sa", "LA Long Beach Glendale MD 31084", "Quarterly", "141 observations\n1991 Q1 to 2026 Q1"),
        ("FHFA", "fhfa_hpi_growth_yoy", "MD quarterly and county annual", "Mixed by geography", "137 MD observations\n1992 Q1 to 2026 Q1"),
    ]
    add_table(
        document,
        ("Family", "Variable ID", "Source geography", "Frequency", "Snapshot coverage"),
        inventory_rows,
        (0.75, 2.15, 1.35, 0.85, 1.85),
        font_size=9.2,
        center_columns=(0, 3),
    )

    add_heading(document, "Geographic Interpretation")
    add_body(
        document,
        "The metropolitan label does not have one meaning across sources. CES and CPI use the Los Angeles Long Beach Anaheim area, which includes Los Angeles and Orange counties. The HVS Los Angeles metro series is also broader than the county. FHFA is different: Metropolitan Division 31084 contains Los Angeles County only. The database should retain each provider's official name and code while adding an explicit county-equivalence flag where supported.",
        indent=False,
    )
    geography_rows = [
        ("CES", "3", "Los Angeles and Orange counties", "Broader regional employment control"),
        ("CPI", "2", "Los Angeles and Orange counties", "Broader regional inflation control"),
        ("HVS", "3", "Published Los Angeles metro", "Broader quarterly vacancy control"),
        ("FHFA", "3", "Metropolitan Division 31084 consists of Los Angeles County", "County-equivalent footprint; preserve provider label"),
    ]
    add_table(
        document,
        ("Source", "IDs", "Footprint", "Required interpretation"),
        geography_rows,
        (0.8, 0.45, 2.2, 3.5),
        font_size=9.7,
        center_columns=(0, 1),
    )
    add_body(
        document,
        "Postal ZIP codes and Census ZCTAs are separate geographies. Postal ZIPs are mail-delivery constructs maintained by USPS; ZCTAs are Census statistical approximations assembled from Census blocks. Matching five-digit codes do not prove identical land area or population coverage. The database must keep both geography types explicit, preserve leading zeros, and require a dated crosswalk plus an approved weighting method before any comparison or aggregation that assumes spatial equivalence.",
        indent=False,
        bold_lead="Postal ZIP codes and Census ZCTAs are separate geographies.",
    )

    add_heading(document, "More Granular Alternatives")
    add_heading(document, "Employment", level=2)
    add_body(
        document,
        "BLS publishes seasonally adjusted total nonfarm employment for the Los Angeles Long Beach Glendale Metropolitan Division as series SMS06310840000000001. Because that division consists of Los Angeles County, it is the closest official monthly replacement for the current two-county CES level while preserving the payroll-establishment concept. It should receive a new ID and new derived growth IDs; it should not overwrite the existing MSA history.",
        indent=False,
    )
    add_bullet(document, "Existing LAUS county variables provide monthly employment, unemployment, unemployment rate, and labor force. They are residence-based and are not direct substitutes for CES payroll jobs.")
    add_bullet(document, "BLS QCEW can add quarterly Los Angeles County employment, establishment counts, and wages by industry. It is employer-based and highly detailed, but it is quarterly and generally not seasonally adjusted.")

    add_heading(document, "Inflation", level=2)
    add_body(
        document,
        "No official BLS general consumer price index specifically for Los Angeles County, postal ZIPs, ZCTAs, or Census tracts was identified. The current CPI variables should remain labeled as Los Angeles Orange regional controls. Allocating the CPI by population, employment, or income would create a modeled estimate and should not be introduced without a separately approved method and ID.",
        indent=False,
    )

    add_heading(document, "Rental Vacancy", level=2)
    add_body(
        document,
        "The Housing Vacancy Survey publishes quarterly vacancy measures for states and large metropolitan areas, not an official quarterly Los Angeles County series. More local alternatives exist, but they differ in frequency or concept.",
        indent=False,
    )
    add_bullet(document, "ACS county and ZCTA variables already maintained include renter occupied units, corresponding margins of error, for rent vacant units, corresponding margins of error, and a county rental availability rate.")
    add_bullet(document, "ACS five-year tract estimates can add annual local detail. They must retain their annual five-year-estimate status and published margins of error, and they must not be interpolated to quarters or relabeled as postal ZIP data.")
    add_bullet(document, "HUD USPS vacancy measures already provide quarterly county, tract, and modeled Los Angeles County portions of postal ZIPs. They measure vacant residential addresses, not renter-specific vacancy, so they remain a proxy rather than an HVS or ACS replacement.")

    add_heading(document, "House Prices", level=2)
    add_body(
        document,
        "The database already maintains fhfa_hpi_la_county_all_transactions_nsa_annual. FHFA also publishes developmental annual HPI products for five-digit ZIP codes and Census tracts. These are strong local additions, but they should remain separate from the standard quarterly purchase-only metropolitan-division series.",
        indent=False,
    )
    add_bullet(document, "Preserve missing or suppressed local indexes when transaction volume is insufficient.")
    add_bullet(document, "Verify and document the provider's ZIP5 geography before joining to postal ZIP or Census ZCTA records.")
    add_bullet(document, "Do not splice annual developmental series into quarterly purchase-only series or imply a common method.")

    alternatives_rows = [
        ("CES payroll employment", "BLS CES MD 31084", "Monthly; LA County-equivalent", "Best same-concept local addition"),
        ("Employment detail", "BLS QCEW", "Quarterly; county and industry", "Add detail; do not replace monthly CES"),
        ("General inflation", "No official local replacement found", "Current CPI area only", "Keep regional control; do not allocate"),
        ("Rental vacancy", "ACS five-year", "Annual; county, ZCTA, tract", "More local but less frequent; retain MOEs"),
        ("Address vacancy", "HUD USPS", "Quarterly; county, tract, modeled postal ZIP parts", "Useful proxy; not renter-specific"),
        ("House prices", "FHFA developmental local HPI", "Annual; county, ZIP5, tract", "Add separately; preserve suppression"),
    ]
    add_table(
        document,
        ("Current need", "Official option", "Native geography and frequency", "Use and limitation"),
        alternatives_rows,
        (1.25, 1.55, 2.0, 2.15),
        font_size=9.6,
    )

    add_heading(document, "Implementation Priorities")
    add_numbered(document, 1, "Add BLS series SMS06310840000000001 and create separately identified level, month-over-month growth, and year-over-year growth variables. Make the county-equivalent family the primary local payroll control while retaining the existing MSA family as a regional comparator.")
    add_numbered(document, 2, "Add FHFA annual ZIP5 HPI and, if tract research is required, the annual tract HPI. Add BLS QCEW county employment, establishments, and wages as a separate quarterly family.")
    add_numbered(document, 3, "Evaluate an annual ACS tract rental-vacancy product. Retain existing county and ZCTA estimates at native annual frequency and HVS as a broader quarterly control.")
    add_numbered(document, 4, "Keep the Los Angeles-area CPI as a Los Angeles Orange regional control. Do not allocate or relabel it as county, postal ZIP, ZCTA, or tract data.")

    add_heading(document, "Automation and Validation Requirements")
    add_body(
        document,
        "Every refresh should discover and ingest the newest provider-published release available at run time rather than relying on a frozen local filename. Automation must still preserve an immutable raw vintage and enough metadata to reproduce every clean or derived observation.",
        indent=False,
    )
    requirements = [
        "Archive source vintage, retrieval timestamp, checksum, original source filename, endpoint, and provider geography.",
        "Use a separate variable ID whenever geography, frequency, seasonal adjustment, or method changes.",
        "Apply source-specific frequency conversion rules; do not use one aggregation rule for every variable.",
        "Do not interpolate missing housing observations or silently fill ZIP, ZCTA, or tract observations with county, MSA, state, or national values.",
        "Stop validation on structural failures such as invalid geography codes, duplicate keys, unexpected source truncation, missing required columns, or a source vintage that is older than the archived active vintage.",
        "Treat plausible economic outliers as warnings for review rather than deleting them automatically.",
        "Update code, configuration, the variable dictionary, the transformation log, documentation, and tests together.",
        "Run the complete pytest suite and python run_pipeline.py --stage validate before declaring an implementation complete.",
    ]
    for requirement in requirements:
        add_bullet(document, requirement)
    add_body(
        document,
        "Access status. The reviewed BLS, Census, and FHFA routes appear public, and no new login or credential requirement was identified during source discovery. Full acquisition was not run. Implementation should stop and request user intervention if a provider introduces authentication, blocks automated access, changes an endpoint, or presents an ambiguous new schema.",
        indent=False,
        bold_lead="Access status.",
    )

    add_heading(document, "Official Sources")
    sources = [
        ("BLS California CES publication structure", "https://www.bls.gov/sae/additional-resources/list-of-published-state-and-metropolitan-area-series/california.htm"),
        ("BLS CPI geographic sample", "https://www.bls.gov/cpi/additional-resources/geographic-sample.htm"),
        ("BLS Local Area Unemployment Statistics", "https://www.bls.gov/lau/lauavail.htm"),
        ("BLS Quarterly Census of Employment and Wages", "https://www.bls.gov/cew/home.htm"),
        ("Census Housing Vacancy Survey tables", "https://www.census.gov/housing/hvs/data/rates.html"),
        ("Census housing data geography guidance", "https://www.census.gov/topics/housing/guidance/topics.html"),
        ("FHFA HPI datasets", "https://www.fhfa.gov/house-price-index?tab=HPI+Datasets"),
        ("OMB metropolitan area delineations", "https://www.bls.gov/bls/omb-bulletin-18-03-revised-delineations-of-metropolitan-statistical-areas.pdf"),
    ]
    for name, url in sources:
        add_source(document, name, url)

    add_heading(document, "Conclusion")
    add_body(
        document,
        "Eight CES, CPI, and HVS IDs use geographies broader than Los Angeles County. Three FHFA IDs use Division 31084, whose footprint equals the county but whose provider label remains metropolitan. Add the county-equivalent monthly CES series as a new family, retain the regional comparisons, and add FHFA ZIP5 or tract, QCEW county, and ACS tract measures only as separate source-native series.",
        indent=False,
    )

    document.save(working)
    if output.exists():
        output.unlink()
    working.replace(output)
    restore_package_part(reference, output, "[Content_Types].xml")

    if sha256(reference) != REFERENCE_HASH:
        raise RuntimeError("The retained template changed during authoring.")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_msa_report.py REFERENCE_DOCX OUTPUT_DOCX")
    build(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())
