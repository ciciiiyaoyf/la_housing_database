from pathlib import Path
import re
import shlex

import duckdb
import pandas as pd
from openpyxl import load_workbook

from src.common.config import load_config
from src.housing_finance_monthly.metadata import PRODUCT_SOURCE_IDS
from src.pipeline import STAGES, parser as pipeline_parser
from src.reporting.catalog import (
    HUD_USPS_RESTRICTED_TABLES,
    HUD_USPS_VARIABLE_IDS,
)


ROOT = Path(__file__).resolve().parents[1]
DATABASE = ROOT / "database" / "housing.duckdb"


def test_catalog_csv_excel_and_active_source_contract():
    config = load_config()
    catalog = pd.read_csv(ROOT / "metadata" / "data_catalog.csv", dtype=str)
    assert catalog.columns[:4].tolist() == [
        "Source",
        "Frequency",
        "Geography",
        "Key variables",
    ]
    workbook = load_workbook(
        ROOT / "metadata" / "data_catalog.xlsx", read_only=True
    )
    assert {
        "Data Catalog",
        "Variable Dictionary",
        "Raw File Inventory",
        "Geography-Frequency Matrix",
        "Coverage Summary",
        "Update Commands",
    }.issubset(workbook.sheetnames)
    excel_ids = {
        str(row[0])
        for row in workbook["Data Catalog"].iter_rows(
            min_row=2, min_col=5, max_col=5, values_only=True
        )
    }
    assert excel_ids == set(catalog["source_id"])
    active_config = {
        source.source_id for source in config.sources if source.enabled
    }
    active_catalog = set(
        catalog.loc[catalog["active_status"].eq("Active"), "source_id"]
    )
    assert active_config | PRODUCT_SOURCE_IDS == active_catalog
    hmda = catalog.loc[
        catalog["source_id"].eq("cfpb_hmda_modified_lar")
    ].iloc[0]
    assert "--start-year YYYY" in hmda["update_command"]
    assert "--local-file PATH_TO_OFFICIAL_FILE" in hmda["update_command"]
    assert "--release-date YYYY-MM-DD" in hmda["update_command"]


def test_raw_repository_completeness_and_checksums():
    report = pd.read_csv(
        ROOT / "output" / "qa" / "raw_repository_completeness.csv",
        dtype=str,
    )
    assert not report.empty
    active = report.loc[report["active_status"].eq("Active")]
    assert active["completeness_status"].eq("PASS").all()
    assert active[
        [
            "raw_snapshot_present",
            "manifest_present",
            "registry_present",
            "checksum_present",
            "retrieval_timestamp_present",
            "source_url_present",
            "frequency_present",
            "geography_present",
            "key_variables_present",
        ]
    ].eq("True").all().all()
    hud_usps = report.loc[
        report["source_id"].eq("hud_usps_ncwm_quarterly")
    ].iloc[0]
    assert hud_usps["active_status"] == "Active"
    assert hud_usps["completeness_status"] == "PASS"
    assert hud_usps["provenance_status"] == "official_raw_archive_rebuilt"
    assert hud_usps[
        [
            "raw_snapshot_present",
            "manifest_present",
            "checksum_present",
            "retrieval_timestamp_present",
        ]
    ].eq("True").all()


def test_added_tables_views_fact_dictionary_and_keys():
    with duckdb.connect(str(DATABASE), read_only=True) as connection:
        objects = {
            (row[0], row[1])
            for row in connection.execute(
                """
                SELECT table_schema, table_name
                FROM information_schema.tables
                WHERE table_schema IN ('clean', 'analytics', 'quality')
                """
            ).fetchall()
        }
        expected = {
            ("clean", "fhfa_hpi"),
            ("clean", "bls_ces_la_metro"),
            ("clean", "bls_cpi_la_metro"),
            ("clean", "bea_state_personal_income"),
            ("clean", "bea_county_personal_income"),
            ("clean", "census_hvs_rental_vacancy"),
            ("clean", "redfin_county_monthly"),
            ("clean", "redfin_zip_monthly"),
            ("clean", "zillow_county_zhvi"),
            ("clean", "zillow_county_zori_smoothed_sa"),
            ("clean", "hud_usps_ncwm_quarterly"),
            ("analytics", "metro_monthly_controls"),
            ("analytics", "state_quarterly_panel"),
            ("analytics", "metro_quarterly_panel"),
            ("analytics", "county_annual_controls"),
            ("analytics", "homeowner_market_monthly"),
            ("analytics", "renter_market_monthly"),
            ("analytics", "overall_market_monthly"),
            ("analytics", "macro_controls_monthly"),
            ("quality", "raw_repository_completeness"),
        }
        assert expected.issubset(objects)
        assert not any(name.startswith("acs_") for schema, name in objects if schema == "clean")
        assert (
            connection.execute(
                """
                SELECT count(*)
                FROM (
                    SELECT observation_date, geo_id, variable_id
                    FROM analytics.fact_housing_observation
                    GROUP BY ALL HAVING count(*) > 1
                )
                """
            ).fetchone()[0]
            == 0
        )
        assert (
            connection.execute(
                """
                SELECT count(*)
                FROM (
                    SELECT DISTINCT variable_id
                    FROM analytics.fact_housing_observation
                    EXCEPT
                    SELECT variable_id
                    FROM metadata.variable_dictionary
                )
                """
            ).fetchone()[0]
            == 0
        )

        hud_variable_count, hud_frequency_count = connection.execute(
            """
            SELECT count(DISTINCT variable_id),
                   count(DISTINCT final_frequency)
            FROM analytics.fact_housing_observation
            WHERE variable_id IN (
                'usps_total_residential_addresses',
                'usps_short_term_vacant_residential_addresses',
                'usps_long_term_vacant_residential_addresses',
                'usps_vacant_residential_addresses',
                'usps_no_stat_residential_addresses',
                'usps_active_residential_addresses',
                'usps_residential_vacancy_rate',
                'usps_residential_no_stat_rate'
            )
            """
        ).fetchone()
        assert hud_variable_count == len(HUD_USPS_VARIABLE_IDS)
        assert hud_frequency_count == 1
        assert connection.execute(
            """
            SELECT bool_and(final_frequency = 'quarterly')
            FROM analytics.fact_housing_observation
            WHERE variable_id LIKE 'usps_%'
            """
        ).fetchone()[0]


def test_hud_usps_restricted_native_quarterly_tables() -> None:
    for relative_path in HUD_USPS_RESTRICTED_TABLES:
        assert (ROOT / relative_path).is_file()

    county = pd.read_parquet(
        ROOT / HUD_USPS_RESTRICTED_TABLES[0],
        columns=["quarter_end", "is_modeled_geography"],
    )
    tract = pd.read_parquet(
        ROOT / HUD_USPS_RESTRICTED_TABLES[1],
        columns=["quarter_end", "is_modeled_geography"],
    )
    postal_zip_part = pd.read_parquet(
        ROOT / HUD_USPS_RESTRICTED_TABLES[2],
        columns=[
            "quarter_end",
            "geography_scope",
            "crosswalk_vintage",
            "is_modeled_geography",
        ],
    )
    lineage = pd.read_parquet(
        ROOT / HUD_USPS_RESTRICTED_TABLES[3],
        columns=["variable_id", "final_frequency", "geo_type", "is_modeled"],
    )

    assert county["is_modeled_geography"].eq(False).all()
    assert tract["is_modeled_geography"].eq(False).all()
    assert postal_zip_part["is_modeled_geography"].eq(True).all()
    assert postal_zip_part["geography_scope"].eq(
        "Los Angeles County portion"
    ).all()
    assert postal_zip_part["crosswalk_vintage"].eq("2026Q1").all()
    assert set(lineage["variable_id"]) == set(HUD_USPS_VARIABLE_IDS)
    assert lineage["final_frequency"].eq("quarterly").all()
    assert lineage.loc[
        lineage["geo_type"].eq("postal_zip_county_part"), "is_modeled"
    ].eq(True).all()


def test_added_geographies_frequencies_and_no_silent_expansion():
    with duckdb.connect(str(DATABASE), read_only=True) as connection:
        assert (
            connection.execute(
                """
                SELECT count(*)
                FROM analytics.fact_housing_observation
                WHERE variable_id LIKE 'fhfa_hpi_la_metropolitan_division%'
                  AND (
                    geo_id != 'METRODIV_31084'
                    OR final_frequency != 'quarterly'
                  )
                """
            ).fetchone()[0]
            == 0
        )
        acs_rows, invalid_acs_rows = connection.execute(
            """
            SELECT count(*),
                   count(*) FILTER (
                       WHERE final_frequency != 'annual'
                          OR NOT (
                              geo_id LIKE 'COUNTY_%'
                              OR geo_id LIKE 'ZCTA_%'
                          )
                   )
            FROM analytics.fact_housing_observation
            WHERE variable_id LIKE 'acs_%'
            """
        ).fetchone()
        assert acs_rows > 0
        assert invalid_acs_rows == 0
        assert (
            connection.execute(
                """
                SELECT count(*)
                FROM analytics.fact_housing_observation
                WHERE variable_id LIKE 'bea_la_county_%'
                  AND (
                    geo_id != 'COUNTY_06037'
                    OR county_fips != '06037'
                    OR state_fips != '06'
                    OR final_frequency != 'annual'
                  )
                """
            ).fetchone()[0]
            == 0
        )
        assert (
            connection.execute(
                """
                SELECT count(*)
                FROM analytics.fact_housing_observation
                WHERE variable_id = 'zori_smoothed_sa'
                  AND (
                    variable_id LIKE '%inventory%'
                    OR unit LIKE '%inventory%'
                  )
                """
            ).fetchone()[0]
            == 0
        )


def test_new_panel_parquet_csv_and_duckdb_row_counts_agree():
    stems = [
        "metro_monthly_controls",
        "state_quarterly_panel",
        "metro_quarterly_panel",
        "county_annual_controls",
    ]
    with duckdb.connect(str(DATABASE), read_only=True) as connection:
        for stem in stems:
            parquet_rows = len(
                pd.read_parquet(ROOT / "data" / "final" / f"{stem}.parquet")
            )
            csv_rows = len(
                pd.read_csv(ROOT / "output" / "exports" / f"{stem}.csv")
            )
            database_rows = connection.execute(
                f"SELECT count(*) FROM analytics.{stem}"
            ).fetchone()[0]
            assert parquet_rows == csv_rows == database_rows


def test_maintenance_guide_refresh_commands_match_cli():
    readme = (ROOT / "docs/database_build_and_refresh.md").read_text(encoding="utf-8")
    assert "catalog" in STAGES
    commands = [
        pipeline_parser().parse_args(shlex.split(line.split("run_pipeline.py ", 1)[1]))
        for block in re.findall(r"(?ms)^```powershell\n(.*?)^```", readme)
        for line in block.splitlines()
        if "run_pipeline.py " in line and "--stage " in line
    ]
    assert any(command.stage == "catalog" for command in commands)
    refresh_sources = {
        command.source for command in commands
        if command.stage == "download" and command.force_download
    }
    for source_id in (
        "bls_laus_la_county",
        "bls_ces_la_metro",
        "bls_cpi_la_metro",
        "fhfa_hpi_state_purchase_only",
        "fhfa_hpi_metro_purchase_only",
        "fhfa_hpi_county_annual",
        "bea_state_personal_income",
        "bea_county_personal_income",
        "census_hvs_state_rental_vacancy",
    ):
        assert source_id in refresh_sources
