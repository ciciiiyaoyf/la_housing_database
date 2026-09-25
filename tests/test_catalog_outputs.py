from pathlib import Path

import duckdb
from openpyxl import load_workbook
import pandas as pd

from src.common.config import load_config
from src.reporting.catalog import (
    DOCX_VARIABLES,
    HUD_USPS_RESTRICTED_TABLES,
    IMPLEMENTATION,
    RENTAL_INVENTORY_STATUS,
    _proxy_methodology_assumptions,
    _write_catalog_preserving_monthly_finance_products,
)


ROOT = Path(__file__).resolve().parents[1]
DATABASE_PATH = ROOT / "database" / "housing.duckdb"


def test_catalog_regeneration_preserves_completed_monthly_finance_product_fields(
    tmp_path: Path,
) -> None:
    path = tmp_path / "data_catalog.csv"
    columns = [
        "source_id",
        "provider_vintage",
        "raw_file_path",
        "checksum_sha256",
        "row_count",
        "notes",
    ]
    pd.DataFrame(
        [
            {
                "source_id": "old_general_source",
                "provider_vintage": "old",
                "raw_file_path": "old.csv",
                "checksum_sha256": "1" * 64,
                "row_count": "1",
                "notes": "must be regenerated",
            },
            {
                "source_id": "housing_finance_version_c_observations",
                "provider_vintage": "information_cutoff_date=2026-08-17",
                "raw_file_path": (
                    "data/final/housing_finance_monthly_rt/"
                    "information_cutoff_date=2026-08-17/part.parquet"
                ),
                "checksum_sha256": "a" * 64,
                "row_count": "3364",
                "notes": "completed cutoff metadata",
            },
        ],
        columns=columns,
    ).to_csv(path, index=False)
    generated = pd.DataFrame(
        [
            {
                "source_id": "current_general_source",
                "provider_vintage": "2026",
                "raw_file_path": "current.csv",
                "checksum_sha256": "2" * 64,
                "row_count": "2",
                "notes": "regenerated",
            }
        ],
        columns=columns,
    )

    _write_catalog_preserving_monthly_finance_products(generated, path)

    result = pd.read_csv(path, dtype=str)
    assert set(result["source_id"]) == {
        "current_general_source",
        "housing_finance_version_c_observations",
    }
    product = result.loc[
        result["source_id"].eq("housing_finance_version_c_observations")
    ].iloc[0]
    assert product["provider_vintage"] == "information_cutoff_date=2026-08-17"
    assert product["checksum_sha256"] == "a" * 64
    assert product["row_count"] == "3364"
    assert product["notes"] == "completed cutoff metadata"


def test_hud_and_irs_catalog_metadata_use_current_source_semantics():
    hud = IMPLEMENTATION[
        "Sheltered vs Unsheltered Homelessness (high frequency)"
    ]
    assert hud["implemented_source_id"] == "hud_pit_coc_2007_2025"
    assert "2007-2025" in hud["limitation"]
    assert "row-level" in hud["limitation"]

    irs = IMPLEMENTATION["Net Immigration / Migration (county)"]
    assert "2022-2023" in irs["limitation"]
    assert "methodology break" in irs["limitation"]
    assert {
        "irs_soi_county_migration",
        "census_acs_5year_b05002",
        "census_acs_5year_b07001",
        "census_pep_county_population_v2019",
        "census_pep_county_population_v2025",
    } == set(irs["implemented_source_id"].split("|"))
    assert "foreign-born field is a population stock" in irs["limitation"]
    assert (
        "moved-from-abroad field is gross current-resident mobility"
        in irs["limitation"]
    )


def test_noncontrolled_catalog_reports_structural_bounds_and_exact_blocker():
    item = IMPLEMENTATION[
        "Rental price increase in non-rent-controlled locations"
    ]
    assert item["current_status"] == "Partially implemented"
    assert {
        "census_acs_5year_b25032",
        "census_acs_5year_b25127",
    } <= set(item["implemented_source_id"].split("|"))
    assert {
        "noncontrolled_exposure_share_structural_proxy_lower_bound",
        "noncontrolled_exposure_share_structural_proxy_upper_bound",
    } <= set(item["implemented_variable_id"].split("|"))
    assert "null before 2020" in item["limitation"]
    assert "not legal eligibility determinations" in item["limitation"]
    assert "Exact noncontrolled_exposure_share" in item["limitation"]
    assert "Census ZCTAs are not USPS postal ZIPs" in item["limitation"]

    assumptions = _proxy_methodology_assumptions()
    method = assumptions.loc[
        assumptions["methodology_id"].eq(
            "acs_noncontrolled_structural_exposure_bounds_v1"
        )
    ]
    assert len(method) == 5
    assert {
        "noncontrolled_exposure_share_structural_proxy_lower_bound",
        "noncontrolled_exposure_share_structural_proxy_upper_bound",
        "noncontrolled_exposure_share",
    } <= set(method["variable_id"])
    assert method.loc[
        method["variable_id"].eq("noncontrolled_exposure_share"),
        "methodology_status",
    ].eq("blocked").all()
    upper = method.loc[
        method["variable_id"].eq(
            "noncontrolled_exposure_share_structural_proxy_upper_bound"
        )
    ].iloc[0]
    assert "min(1" in upper["value"]
    assert "prevents double counting" in upper["rationale"]


def test_docx_variable_audit_covers_attached_variable_sheet():
    audit = pd.read_csv(ROOT / "metadata" / "docx_variable_audit.csv")
    requirements = pd.read_csv(ROOT / "metadata" / "docx_variable_requirements.csv")
    expected_names = {item["variable_name"] for item in DOCX_VARIABLES}

    assert len(audit) == len(DOCX_VARIABLES)
    assert expected_names == set(audit["variable_name"])
    assert len(requirements) == len(DOCX_VARIABLES)
    assert {
        "sheet_source",
        "sheet_frequency",
        "sheet_la_availability",
        "sheet_notes",
        "best_public_workaround",
    }.issubset(requirements.columns)

    rental_inventory = audit.loc[
        audit["variable_name"].eq("Rental Inventory")
    ].iloc[0]
    assert rental_inventory["expected_source_from_docx"] == (
        "Zillow Observed Rent Index / Listings"
    )
    assert rental_inventory["current_status"] == RENTAL_INVENTORY_STATUS
    assert pd.isna(rental_inventory["implemented_source_id"])
    assert pd.isna(rental_inventory["implemented_variable_id"])
    assert not bool(rental_inventory["panel_or_view_present"])
    assert "not true rental-listing inventory" in rental_inventory["limitation"]
    assert "native-quarterly only" in rental_inventory["limitation"]
    assert "no-stat remains a separate diagnostic" in rental_inventory[
        "limitation"
    ].lower()
    assert "postal-ZIP county-parts are modeled" in rental_inventory[
        "limitation"
    ]
    assert "access-controlled" in rental_inventory["limitation"]
    assert "ZORI is not inventory" in rental_inventory["limitation"]
    assert "Redfin for-sale active listings are not a valid substitute" in (
        rental_inventory["limitation"]
    )
    assert "governing methodology names no approved source" in rental_inventory[
        "limitation"
    ].lower()

    permits = audit.loc[
        audit["variable_name"].eq("Single-Family Housing Starts (permits)")
    ].iloc[0]
    assert permits["expected_source_from_docx"] == "Census BPS (1-unit permits)"
    assert "authorizations, not starts or completions" in permits["limitation"]

    occupied_change = audit.loc[
        audit["variable_name"].eq("Net Change in Occupied Apartments (quarterly)")
    ].iloc[0]
    assert set(occupied_change["implemented_variable_id"].split("|")) == {
        "net_change_occupied_apartments_annual",
        "zcta_renter_occupied_units_vintage_change_annual",
    }
    assert occupied_change["row_count"] == 14
    assert occupied_change["start_date"] == "2011-12-31"
    assert occupied_change["end_date"] == "2024-12-31"

    price_to_income = audit.loc[
        audit["variable_name"].eq("Price-to-Income Ratio (quarterly/local)")
    ].iloc[0]
    assert price_to_income["row_count"] == 15
    assert price_to_income["start_date"] == "2010-12-31"
    assert price_to_income["end_date"] == "2024-12-31"


def test_data_catalog_shape_and_active_sources_are_complete():
    config = load_config(ROOT)
    active_source_ids = {source.source_id for source in config.sources if source.enabled}
    catalog = pd.read_csv(ROOT / "metadata" / "data_catalog.csv")
    raw_completeness = pd.read_csv(
        ROOT / "output" / "qa" / "raw_repository_completeness.csv"
    )

    assert catalog.columns[:4].tolist() == [
        "Source",
        "Frequency",
        "Geography",
        "Key variables",
    ]
    assert active_source_ids.issubset(set(catalog["source_id"]))
    assert active_source_ids.issubset(set(raw_completeness["source_id"]))
    assert raw_completeness.loc[
        raw_completeness["source_id"].isin(active_source_ids),
        "completeness_status",
    ].eq("PASS").all()

    hud_usps = catalog.loc[
        catalog["source_id"].eq("hud_usps_ncwm_quarterly")
    ].iloc[0]
    assert hud_usps["active_status"] == "Active"
    assert hud_usps["provenance_status"] == "official_raw_archive_rebuilt"
    assert str(hud_usps["update_command"]).startswith(
        "python run_pipeline.py --stage download --source "
        "hud_usps_ncwm_quarterly"
    )
    assert hud_usps["source_manifest_path"] == (
        "metadata/hud_usps_source_manifest.csv"
    )
    assert "data/restricted/hud_usps/raw/" in hud_usps["raw_file_path"]
    assert "quarantine" not in hud_usps["raw_file_path"]
    assert "historical, superseded record" in hud_usps["notes"]
    assert "clean.hud_usps_ncwm_quarterly" in hud_usps[
        "destination_clean_table"
    ]
    assert "analytics.fact_housing_observation" in hud_usps[
        "destination_analytical_table"
    ]
    assert all(
        path in hud_usps["destination_analytical_table"]
        for path in HUD_USPS_RESTRICTED_TABLES[:3]
    )


def test_hud_usps_v2_methodology_and_gap_workbook() -> None:
    assumptions = pd.read_csv(
        ROOT / "metadata" / "proxy_methodology_assumptions.csv"
    )
    hud = assumptions.loc[
        assumptions["methodology_id"].eq(
            "hud_usps_residential_vacancy_proxy_v2"
        )
    ]
    assert len(hud) == 4
    assert hud["methodology_status"].eq("active").all()
    combined = " ".join(hud["value"].astype(str))
    assert "metadata/hud_usps_source_manifest.csv" in combined
    assert "No-stat / total remains a separate" in combined
    assert "postal-ZIP county-part values are modeled" in combined
    assert "native-quarterly" in combined
    assert "data/restricted/hud_usps/" in " ".join(
        hud["rationale"].astype(str)
    )

    acs = assumptions.loc[
        assumptions["methodology_id"].eq(
            "acs_rental_availability_proxy_v1"
        )
    ]
    assert len(acs) == 1
    assert acs.iloc[0]["variable_id"] == "acs_rental_availability_rate"
    assert "B25004" in acs.iloc[0]["value"]
    assert "not an active rental-listing stock" in acs.iloc[0]["rationale"]

    workbook = load_workbook(
        ROOT / "metadata" / "data_gap_analysis.xlsx", read_only=True
    )
    assert workbook.sheetnames == ["Gap Analysis"]


def test_source_coverage_preserves_geography_types():
    coverage = pd.read_csv(ROOT / "metadata" / "source_coverage.csv")

    zcta = coverage.loc[
        coverage["variable_id"].eq("acs_median_gross_rent")
        & coverage["geo_type"].eq("zcta")
    ].iloc[0]
    assert zcta["geography_count"] == 274
    assert set(
        coverage.loc[
            coverage["variable_id"].eq("acs_median_gross_rent"),
            "geo_type",
        ]
    ) == {"county", "zcta"}
    assert coverage.loc[
        coverage["variable_id"].eq(
            "fhfa_hpi_la_metropolitan_division_purchase_only_nsa"
        ),
        "geo_type",
    ].eq("metro_division").all()
    assert coverage.loc[
        coverage["variable_id"].eq("ces_total_nonfarm_employment_sa"),
        "geo_type",
    ].eq("metro").all()
    assert coverage.loc[
        coverage["variable_id"].eq("mortgage_rate_30y_weekly"),
        "geo_type",
    ].eq("national").all()
    assert not coverage["geo_type"].eq("unknown").any()


def test_source_tracker_workbook_is_human_readable():
    workbook_path = ROOT / "metadata" / "LA_Housing_Source_Tracker.xlsx"
    workbook = load_workbook(workbook_path)

    assert workbook.sheetnames[:2] == ["Overview", "Source Summary"]
    expected_sheets = {
        "Overview",
        "Source Summary",
        "DOCX Variable Audit",
        "Variable Dictionary",
        "Coverage Summary",
        "Data Catalog",
        "Raw File Inventory",
        "Raw Completeness",
        "Missing Deferred",
        "Proxy Assumptions",
        "Update Commands",
    }
    assert expected_sheets.issubset(set(workbook.sheetnames))

    for sheet_name in expected_sheets:
        worksheet = workbook[sheet_name]
        assert worksheet.sheet_view.showGridLines is False
        assert worksheet.tables

    summary_headers = [
        workbook["Source Summary"].cell(row=1, column=column).value
        for column in range(1, workbook["Source Summary"].max_column + 1)
    ]
    assert summary_headers[:6] == [
        "Source",
        "Provider",
        "Status",
        "Frequency",
        "Geography",
        "Key variables",
    ]

    def sheet_records(sheet_name: str) -> list[dict[str, object]]:
        worksheet = workbook[sheet_name]
        headers = [
            worksheet.cell(row=1, column=column).value
            for column in range(1, worksheet.max_column + 1)
        ]
        records = []
        for row in range(2, worksheet.max_row + 1):
            records.append(
                {
                    headers[column - 1]: worksheet.cell(row=row, column=column).value
                    for column in range(1, worksheet.max_column + 1)
                }
            )
        return records

    data_catalog = sheet_records("Data Catalog")
    assert any(
        row.get("source_id") == "eviction_lab_county_estimates"
        for row in data_catalog
    )

    variable_audit = sheet_records("DOCX Variable Audit")
    eviction_rows = [
        row
        for row in variable_audit
        if row.get("variable_name") == "Eviction trends (high-frequency, consistent)"
    ]
    assert eviction_rows
    assert eviction_rows[0]["current_status"] == "Partially implemented"
    assert eviction_rows[0]["implemented_source_id"] == "eviction_lab_county_estimates"

    variable_dictionary = sheet_records("Variable Dictionary")
    dictionary_ids = {row.get("variable_id") for row in variable_dictionary}
    assert "eviction_lab_filings_estimate_annual" in dictionary_ids
    assert "natural_population_change_quarterly_allocated" in dictionary_ids

    coverage = sheet_records("Coverage Summary")
    assert any(
        str(row.get("variable_id", "")).startswith("eviction_lab_")
        for row in coverage
    )
    assert any(
        row.get("variable_id") == "acs_median_gross_rent"
        and row.get("geo_type") == "zcta"
        and row.get("geography_count") == 274
        for row in coverage
    )
    assert any(
        row.get("variable_id")
        == "fhfa_hpi_la_metropolitan_division_purchase_only_nsa"
        and row.get("geo_type") == "metro_division"
        for row in coverage
    )


def test_raw_repository_completeness_is_available_in_duckdb():
    with duckdb.connect(str(DATABASE_PATH), read_only=True) as connection:
        row_count = connection.execute(
            "SELECT count(*) FROM quality.raw_repository_completeness"
        ).fetchone()[0]
        duplicate_count = connection.execute(
            """
            SELECT count(*)
            FROM (
                SELECT geo_id, observation_date, variable_id
                FROM analytics.fact_housing_observation
                GROUP BY ALL
                HAVING count(*) > 1
            )
            """
        ).fetchone()[0]

    assert row_count >= len(load_config(ROOT).sources)
    assert duplicate_count == 0


def test_version_a_finance_is_authoritative_and_b25106_legacy_is_retained():
    assumptions = pd.read_csv(ROOT / "metadata" / "proxy_methodology_assumptions.csv")
    assert "acs_implied_aggregate_mortgage_debt_proxy" in set(
        assumptions["variable_id"]
    )
    legacy_assumptions = assumptions.loc[
        assumptions["methodology_id"].eq(
            "acs_implied_mortgage_debt_proxy_legacy_v0"
        )
    ]
    assert not legacy_assumptions.empty
    assert legacy_assumptions["methodology_status"].eq("legacy").all()

    audit = pd.read_csv(ROOT / "metadata" / "docx_variable_audit.csv")
    mortgage_debt = audit.loc[
        audit["variable_name"].eq("Aggregate Mortgage Debt (outstanding balances)")
    ].iloc[0]
    assert mortgage_debt["current_status"] == "Partially implemented"
    assert "census_acs5_finance" in mortgage_debt["implemented_source_id"]
    assert "acs_implied_first_mortgage_balance_cal_benchmark" in mortgage_debt[
        "implemented_variable_id"
    ]
    assert "older B25106" in mortgage_debt["limitation"]
    assert "raw_pmms_20y_sensitivity" in mortgage_debt[
        "implemented_variable_id"
    ]

    home_equity = audit.loc[
        audit["variable_name"].eq("Aggregate Home Equity")
    ].iloc[0]
    assert home_equity["current_status"] == "Partially implemented"
    assert "home_equity_all_pmms_25y_scenario" in home_equity[
        "implemented_variable_id"
    ]
    assert "uncalibrated scenarios" in home_equity["limitation"]

    with duckdb.connect(str(DATABASE_PATH), read_only=True) as connection:
        clean_rows = connection.execute(
            "SELECT count(*) FROM clean.census_acs_5year_b25106"
        ).fetchone()[0]
        proxy = connection.execute(
            """
            SELECT acs_implied_aggregate_mortgage_debt_proxy,
                   acs_implied_average_mortgage_balance_proxy,
                   acs_mortgage_debt_proxy_rate_assumption,
                   acs_mortgage_debt_proxy_term_years_assumption,
                   acs_mortgage_debt_proxy_debt_service_share_assumption
            FROM analytics.county_annual_controls
            WHERE date = DATE '2024-12-31'
            """
        ).fetchone()

    assert clean_rows > 0
    assert proxy[0] > 0
    assert proxy[1] > 0
    assert proxy[2] > 0
    assert proxy[3] == 25.0
    assert proxy[4] == 0.80


def test_irs_soi_net_migration_is_documented_and_loaded():
    audit = pd.read_csv(ROOT / "metadata" / "docx_variable_audit.csv")
    migration = audit.loc[
        audit["variable_name"].eq("Net Immigration / Migration (county)")
    ].iloc[0]
    assert migration["current_status"] == "Implemented"
    assert {
        "irs_soi_county_migration",
        "census_acs_5year_b05002",
        "census_acs_5year_b07001",
        "census_pep_county_population_v2019",
        "census_pep_county_population_v2025",
    } == set(migration["implemented_source_id"].split("|"))
    assert "irs_soi_net_migration_exemptions" in migration["implemented_variable_id"]
    assert "acs_foreign_born_population" in migration["implemented_variable_id"]
    assert "acs_moved_from_abroad_past_year" in migration["implemented_variable_id"]

    with duckdb.connect(str(DATABASE_PATH), read_only=True) as connection:
        clean_coverage = connection.execute(
            """
            SELECT min(observation_date), max(observation_date),
                   count(DISTINCT observation_date), count(*)
            FROM clean.irs_soi_county_migration
            """
        ).fetchone()
        county_coverage = connection.execute(
            """
            SELECT min(date), max(date), count(*)
            FROM analytics.county_annual_controls
            WHERE irs_soi_net_migration_exemptions IS NOT NULL
            """
        ).fetchone()
        duplicate_count = connection.execute(
            """
            SELECT count(*)
            FROM (
                SELECT geo_id, observation_date, variable_id
                FROM clean.irs_soi_county_migration
                GROUP BY ALL
                HAVING count(*) > 1
            )
            """
        ).fetchone()[0]
        acs_supplement_coverage = connection.execute(
            """
            SELECT variable_id,
                   count(*) AS row_count,
                   count(value) AS nonnull_count,
                   min(observation_date) AS first_date,
                   max(observation_date) AS last_date,
                   count(DISTINCT geo_id) AS geography_count
            FROM analytics.fact_housing_observation
            WHERE variable_id IN (
                'acs_foreign_born_population',
                'acs_foreign_born_population_moe',
                'acs_moved_from_abroad_past_year',
                'acs_moved_from_abroad_past_year_moe'
            )
            GROUP BY variable_id
            ORDER BY variable_id
            """
        ).fetchall()

    assert clean_coverage[0].year == 2011
    assert clean_coverage[1].year == 2023
    assert clean_coverage[2] == 13
    assert clean_coverage[3] == 117
    assert county_coverage[0].year == 2011
    assert county_coverage[1].year == 2023
    assert county_coverage[2] == 13
    assert duplicate_count == 0
    assert len(acs_supplement_coverage) == 4
    for _, row_count, nonnull_count, first_date, last_date, geography_count in (
        acs_supplement_coverage
    ):
        assert row_count == 3851
        assert nonnull_count > 0
        assert first_date.year == 2010
        assert last_date.year == 2024
        assert geography_count == 275


def test_acs_migration_supplement_transformation_lineage_keeps_both_sources():
    log = pd.read_csv(ROOT / "metadata" / "transformation_log.csv", dtype=str)
    rows = log.loc[
        log["source_id"].isin(
            {
                "census_acs_5year_b05002",
                "census_acs_5year_b07001",
            }
        )
        & log["transformation_id"].eq("acs_5year_api_summary_v1")
    ]

    assert set(rows["source_id"]) == {
        "census_acs_5year_b05002",
        "census_acs_5year_b07001",
    }
    assert rows["output_rows"].astype(int).eq(7702).all()
    assert rows["output_file"].str.endswith(".parquet").all()


def test_cdc_wonder_natural_population_change_is_loaded_and_allocated():
    audit = pd.read_csv(ROOT / "metadata" / "docx_variable_audit.csv")
    natural = audit.loc[
        audit["variable_name"].eq("Natural Population Change (quarterly)")
    ].iloc[0]
    assert natural["current_status"] == "Implemented with quarterly allocation"
    assert "cdc_wonder_natality_county" in natural["implemented_source_id"]
    assert "natural_population_change_quarterly_allocated" in natural[
        "implemented_variable_id"
    ]

    with duckdb.connect(str(DATABASE_PATH), read_only=True) as connection:
        annual = connection.execute(
            """
            SELECT min(date), max(date), count(*)
            FROM analytics.county_annual_controls
            WHERE natural_population_change_annual IS NOT NULL
              AND county_fips = '06037'
            """
        ).fetchone()
        quarterly = connection.execute(
            """
            SELECT min(date), max(date), count(*),
                   bool_and(is_imputed), bool_and(is_interpolated)
            FROM analytics.county_quarterly_controls
            WHERE natural_population_change_quarterly_allocated IS NOT NULL
              AND county_fips = '06037'
            """
        ).fetchone()

    assert annual[0].year == 2007
    assert annual[1].year == 2024
    assert annual[2] == 18
    assert quarterly[0] == pd.Timestamp("2007-03-31")
    assert quarterly[1] == pd.Timestamp("2024-12-31")
    assert quarterly[2] == 72
    assert quarterly[3] is True
    assert quarterly[4] is True


def test_eviction_lab_county_estimates_are_documented_and_loaded():
    audit = pd.read_csv(ROOT / "metadata" / "docx_variable_audit.csv")
    evictions = audit.loc[
        audit["variable_name"].eq("Eviction trends (high-frequency, consistent)")
    ].iloc[0]
    assert evictions["current_status"] == "Partially implemented"
    assert evictions["implemented_source_id"] == "eviction_lab_county_estimates"
    assert "eviction_lab_filings_rate_estimate_annual" in evictions[
        "implemented_variable_id"
    ]
    assert "neither Los Angeles nor any California geography" in evictions[
        "limitation"
    ]

    with duckdb.connect(str(DATABASE_PATH), read_only=True) as connection:
        clean_coverage = connection.execute(
            """
            SELECT min(observation_date), max(observation_date),
                   count(DISTINCT observation_date), count(*),
                   count(DISTINCT county_fips)
            FROM clean.eviction_lab_county_estimates
            WHERE variable_id IN (
                'eviction_lab_filings_estimate_annual',
                'eviction_lab_households_threatened_estimate_annual'
            )
            """
        ).fetchone()
        county_coverage = connection.execute(
            """
            SELECT min(date), max(date), count(*)
            FROM analytics.county_annual_controls
            WHERE eviction_lab_filings_rate_estimate_annual IS NOT NULL
              AND county_fips = '06037'
            """
        ).fetchone()
        duplicate_count = connection.execute(
            """
            SELECT count(*)
            FROM (
                SELECT geo_id, observation_date, variable_id
                FROM clean.eviction_lab_county_estimates
                GROUP BY ALL
                HAVING count(*) > 1
            )
            """
        ).fetchone()[0]

    assert clean_coverage[0].year == 2000
    assert clean_coverage[1].year == 2018
    assert clean_coverage[2] == 19
    assert clean_coverage[3] == 38
    assert clean_coverage[4] == 1
    assert county_coverage[0].year == 2000
    assert county_coverage[1].year == 2018
    assert county_coverage[2] == 19
    assert duplicate_count == 0
