from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from src.cleaning.sipp import (
    DEBT_SOURCE_ID,
    DEBT_TABLE_2_TITLE,
    DEBT_TABLE_4_TITLE,
    DEBT_VARIABLE_ID,
    EQUITY_VARIABLE_ID,
    EXPECTED_RELEASE_DATE,
    WEALTH_SOURCE_ID,
    WEALTH_TABLE_2_TITLE,
    WEALTH_TABLE_3_TITLE,
    WEALTH_TABLE_5_TITLE,
    clean_sipp_state_sensitivity,
)
from src.common.config import load_config


SOURCE_NOTE = (
    "Source: U.S. Census Bureau, Survey of Income and Program Participation, "
    "Survey Year 2025, Public Use Data."
)
RELEASE_NOTE = "Internet Release Date: 07/15/2026"


def _add_common_notes(worksheet) -> None:
    worksheet["A20"] = SOURCE_NOTE
    worksheet["A21"] = RELEASE_NOTE


def _debt_workbook(path: Path, *, home_debt_header: str = "Home Debt") -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    table_2 = workbook.create_sheet("Table 2")
    table_2["A2"] = DEBT_TABLE_2_TITLE
    table_2["B3"] = "Number of Households (thousands)"
    table_2["E4"] = home_debt_header
    table_2["A7"] = "California"
    table_2["B7"] = 14420
    table_2["E7"] = 32.6
    _add_common_notes(table_2)
    table_4 = workbook.create_sheet("Table 4")
    table_4["A2"] = DEBT_TABLE_4_TITLE
    table_4["D4"] = "Home Debt"
    table_4["A7"] = "California"
    table_4["D7"] = 346200
    _add_common_notes(table_4)
    workbook.save(path)


def _wealth_workbook(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    table_2 = workbook.create_sheet("Table 2")
    table_2["A2"] = WEALTH_TABLE_2_TITLE
    table_2["E3"] = "Equity in Own Home"
    table_2["A7"] = "California"
    table_2["E7"] = 51.0
    _add_common_notes(table_2)
    table_3 = workbook.create_sheet("Table 3")
    table_3["A2"] = WEALTH_TABLE_3_TITLE
    table_3["B3"] = "Number of Households (thousands)"
    table_3["A7"] = "California"
    table_3["B7"] = 14420
    _add_common_notes(table_3)
    table_5 = workbook.create_sheet("Table 5")
    table_5["A2"] = WEALTH_TABLE_5_TITLE
    table_5["G3"] = "Equity in Own Home"
    table_5["A7"] = "California"
    table_5["G7"] = 668100
    _add_common_notes(table_5)
    workbook.save(path)


def _record(source_id: str) -> dict[str, str]:
    return {
        "source_id": source_id,
        "data_vintage": "2026-08-10",
        "retrieved_at": "2026-08-10T18:00:00+00:00",
        "source_filename": f"{source_id}_archive.xlsx",
        "checksum_sha256": "a" * 64,
        "provider_last_updated": EXPECTED_RELEASE_DATE,
    }


def test_sipp_debt_constructs_only_state_sensitivity(tmp_path: Path) -> None:
    source = tmp_path / "debt.xlsx"
    output = tmp_path / "debt.parquet"
    _debt_workbook(source)

    frame = clean_sipp_state_sensitivity(
        source, _record(DEBT_SOURCE_ID), output
    )

    assert output.exists()
    assert len(frame) == 1
    row = frame.iloc[0]
    assert row["geo_id"] == "STATE_06"
    assert row["variable_id"] == DEBT_VARIABLE_ID
    assert row["value"] == pytest.approx(1_627_458_504_000.0)
    assert pd.isna(row["reported_value"])
    assert row["observation_date"] == pd.Timestamp("2024-12-31")
    assert row["provider_last_updated"] == EXPECTED_RELEASE_DATE
    assert "not_exact_nmdb_calibration" in row["quality_flag"]
    assert "no_covariance_based_standard_error" in row["quality_flag"]


def test_sipp_wealth_constructs_only_positive_equity_sensitivity(
    tmp_path: Path,
) -> None:
    source = tmp_path / "wealth.xlsx"
    output = tmp_path / "wealth.parquet"
    _wealth_workbook(source)

    frame = clean_sipp_state_sensitivity(
        source, _record(WEALTH_SOURCE_ID), output
    )

    assert len(frame) == 1
    row = frame.iloc[0]
    assert row["variable_id"] == EQUITY_VARIABLE_ID
    assert row["value"] == pytest.approx(4_913_341_020_000.0)
    assert "positive_equity_holders_only" in row["quality_flag"]
    assert "net_worth_below_99th_percentile" in row["quality_flag"]
    assert not frame["variable_id"].str.contains("cal_benchmark").any()


def test_sipp_schema_drift_stops_cleaning(tmp_path: Path) -> None:
    source = tmp_path / "changed.xlsx"
    _debt_workbook(source, home_debt_header="Mortgage Debt")

    with pytest.raises(ValueError, match="schema drift"):
        clean_sipp_state_sensitivity(
            source,
            _record(DEBT_SOURCE_ID),
            tmp_path / "changed.parquet",
        )


def test_sipp_sources_and_sensitivities_are_configured() -> None:
    config = load_config()
    source_ids = {source.source_id for source in config.sources}
    variable_ids = {str(variable["variable_id"]) for variable in config.variables}
    assert {DEBT_SOURCE_ID, WEALTH_SOURCE_ID}.issubset(source_ids)
    assert {DEBT_VARIABLE_ID, EQUITY_VARIABLE_ID}.issubset(variable_ids)
